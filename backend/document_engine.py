from __future__ import annotations

import copy
import hashlib
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pymupdf

from backend.domain.schemas import ClaimAssessment, JobRecord, PdfBoundingBox, ReviewComment


def sha256_file(path: str | Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def create_born_digital_reference_pdf(
    pdf_path: str | Path,
    pages_spec: list[dict[str, Any]],
) -> dict[str, Any]:
    """
    Creates a real born-digital PDF file on disk using PyMuPDF (`pymupdf`).
    Each item in `pages_spec` can specify:
    - `header`: optional running header string (set `include_header: False` or `scanned_image_only: True` to omit)
    - `passages: list[str]`: single-column full-width passages
    - `left_column: list[str]` and `right_column: list[str]`: two-column clinical reprint layout
    - `scanned_image_only: bool`: creates a textless/scanned page with vector graphics only
    """
    out_path = Path(pdf_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = pymupdf.open()

    for idx, p_spec in enumerate(pages_spec, start=1):
        page = doc.new_page(width=612, height=792)
        if p_spec.get("scanned_image_only"):
            shape = page.new_shape()
            shape.draw_rect(pymupdf.Rect(54, 72, 558, 700))
            shape.finish(color=(0.75, 0.75, 0.75), fill=(0.96, 0.96, 0.96))
            shape.commit()
            continue

        if p_spec.get("include_header", True):
            header = p_spec.get("header", f"Authorized Clinical & Regulatory Reference — Page {idx}")
            if header:
                page.insert_text(pymupdf.Point(54, 54), header, fontsize=11)

        left_col = p_spec.get("left_column", [])
        right_col = p_spec.get("right_column", [])
        if left_col or right_col:
            y_left = 96.0
            for passage in left_col:
                rect = pymupdf.Rect(54, y_left, 294, y_left + 95)
                page.insert_textbox(rect, passage, fontsize=9.5)
                y_left += 110.0
            y_right = 96.0
            for passage in right_col:
                rect = pymupdf.Rect(318, y_right, 558, y_right + 95)
                page.insert_textbox(rect, passage, fontsize=9.5)
                y_right += 110.0
        else:
            y_cursor = 96.0
            for passage in p_spec.get("passages", []):
                rect = pymupdf.Rect(54, y_cursor, 558, y_cursor + 90)
                page.insert_textbox(rect, passage, fontsize=10)
                y_cursor += 105.0

    doc.save(str(out_path))
    doc.close()
    return {
        "pdf_path": str(out_path),
        "page_count": len(pages_spec),
        "sha256": sha256_file(out_path),
    }


_PDF_STOPWORDS = {
    "this",
    "that",
    "with",
    "from",
    "were",
    "have",
    "been",
    "across",
    "page",
    "reference",
    "clinical",
    "study",
    "report",
    "authorized",
}


def normalize_pdf_rect_to_box_2d(
    *,
    x0: float,
    y0: float,
    x1: float,
    y1: float,
    page_width: float = 612.0,
    page_height: float = 792.0,
) -> list[int]:
    """
    Converts PDF point coordinates `(x0, y0, x1, y1)` into Gemini multimodal
    `box_2d` `[ymin, xmin, ymax, xmax]` normalized on the `[0, 1000]` integer grid.
    """
    pw = page_width if page_width > 0 else 612.0
    ph = page_height if page_height > 0 else 792.0
    ymin = max(0, min(1000, int(round((y0 / ph) * 1000))))
    xmin = max(0, min(1000, int(round((x0 / pw) * 1000))))
    ymax = max(0, min(1000, int(round((y1 / ph) * 1000))))
    xmax = max(0, min(1000, int(round((x1 / pw) * 1000))))
    return [ymin, xmin, ymax, xmax]


def denormalize_box_2d_to_pdf_rect(
    box_2d: list[int | float],
    *,
    page_width: float = 612.0,
    page_height: float = 792.0,
) -> dict[str, float]:
    """
    Converts Gemini `box_2d` `[ymin, xmin, ymax, xmax]` (`0..1000`) back into PDF points.
    """
    if not isinstance(box_2d, (list, tuple)) or len(box_2d) < 4:
        raise ValueError("box_2d must contain 4 coordinates [ymin, xmin, ymax, xmax].")
    pw = page_width if page_width > 0 else 612.0
    ph = page_height if page_height > 0 else 792.0
    ymin, xmin, ymax, xmax = [max(0.0, min(1000.0, float(v))) for v in box_2d[:4]]
    return {
        "x0": round((xmin / 1000.0) * pw, 2),
        "y0": round((ymin / 1000.0) * ph, 2),
        "x1": round((xmax / 1000.0) * pw, 2),
        "y1": round((ymax / 1000.0) * ph, 2),
    }


def _extract_sorted_page_blocks(
    page: Any,
    page_num: int,
) -> tuple[list[dict[str, Any]], bool, bool]:
    """
    Extracts text blocks from a PyMuPDF `page` and sorts them in proper reading order:
    - Detects two-column layouts (e.g. clinical journal reprints with distinct left and right
      column blocks around the horizontal page midpoint) and sorts top spanning headers first,
      then left-column blocks top-to-bottom, then right-column blocks top-to-bottom, then
      bottom spanning footers.
    - Detects textless/scanned pages (`ocr_fallback_required = True`) when no text blocks exist.
    Returns `(sorted_blocks, is_two_column, ocr_fallback_required)`.
    """
    pw = float(page.rect.width)
    ph = float(page.rect.height)
    mid_x = pw / 2.0
    blocks_raw = page.get_text("blocks")
    parsed_blocks: list[dict[str, Any]] = []

    for b_idx, blk in enumerate(blocks_raw):
        x0, y0, x1, y1, raw_text, block_no, block_type = blk[:7]
        if block_type != 0:
            continue
        cleaned = " ".join(str(raw_text).strip().split())
        if not cleaned:
            continue
        rx0 = round(float(x0), 2)
        ry0 = round(float(y0), 2)
        rx1 = round(float(x1), 2)
        ry1 = round(float(y1), 2)
        bw = rx1 - rx0
        cx = (rx0 + rx1) / 2.0

        if rx1 <= mid_x + 18.0 or (bw <= 0.56 * pw and cx < mid_x - 10.0):
            col_tag = "left"
        elif rx0 >= mid_x - 18.0 or (bw <= 0.56 * pw and cx > mid_x + 10.0):
            col_tag = "right"
        else:
            col_tag = "full"

        box_2d = normalize_pdf_rect_to_box_2d(
            x0=rx0,
            y0=ry0,
            x1=rx1,
            y1=ry1,
            page_width=pw,
            page_height=ph,
        )
        parsed_blocks.append(
            {
                "block_index": int(block_no if block_no is not None else b_idx),
                "page_number": page_num,
                "text": cleaned,
                "x0": rx0,
                "y0": ry0,
                "x1": rx1,
                "y1": ry1,
                "ymin": ry0,
                "xmin": rx0,
                "ymax": ry1,
                "xmax": rx1,
                "box_2d": box_2d,
                "page_width": pw,
                "page_height": ph,
                "column": col_tag,
            }
        )

    if not parsed_blocks:
        return [], False, True

    left_blocks = [b for b in parsed_blocks if b["column"] == "left"]
    right_blocks = [b for b in parsed_blocks if b["column"] == "right"]
    is_two_column = bool(left_blocks and right_blocks)

    if is_two_column:
        min_col_y0 = min(b["y0"] for b in left_blocks + right_blocks)
        top_full = sorted(
            [b for b in parsed_blocks if b["column"] == "full" and b["y0"] <= min_col_y0],
            key=lambda b: (b["y0"], b["x0"]),
        )
        left_sorted = sorted(left_blocks, key=lambda b: (b["y0"], b["x0"]))
        right_sorted = sorted(right_blocks, key=lambda b: (b["y0"], b["x0"]))
        bottom_full = sorted(
            [b for b in parsed_blocks if b["column"] == "full" and b["y0"] > min_col_y0],
            key=lambda b: (b["y0"], b["x0"]),
        )
        ordered = top_full + left_sorted + right_sorted + bottom_full
    else:
        for b in parsed_blocks:
            b["column"] = "full"
        ordered = sorted(parsed_blocks, key=lambda b: (b["y0"], b["x0"]))

    for order_idx, blk_item in enumerate(ordered, start=1):
        blk_item["reading_order"] = order_idx
        blk_item["ocr_fallback_required"] = False

    return ordered, is_two_column, False


def locate_supporting_quote_in_pdf(
    pdf_path: str | Path,
    quote_or_query: str,
    *,
    source_id: str = "SRC-UPLOADED-PDF",
    version: str = "1.0",
    preferred_page: int | None = None,
    used_page_rects: set[tuple[int, float, float, float, float]] | None = None,
) -> dict[str, Any]:
    """
    Locates a supporting clinical quote in `pdf_path` while respecting two-column journal
    reprint reading order (left column top-to-bottom, then right column top-to-bottom) and
    gracefully flagging scanned/textless pages with `ocr_fallback_required=True`.
    """
    path = Path(pdf_path)
    if not path.exists():
        raise FileNotFoundError(f"Source PDF not found: {path}")

    doc = pymupdf.open(str(path))
    raw_query = str(quote_or_query or "").strip()
    normalized_query = " ".join(raw_query.split())
    search_probe = normalized_query[:65].strip()
    seen_rects = used_page_rects if used_page_rects is not None else set()

    found_bbox: dict[str, Any] | None = None
    fallback_bbox: dict[str, Any] | None = None
    matched_verbatim = normalized_query
    any_text_blocks = False
    preferred_page_textless = False

    try:
        page_indices = list(range(len(doc)))
        if isinstance(preferred_page, int) and 1 <= preferred_page <= len(doc):
            pref_idx = preferred_page - 1
            page_indices = [pref_idx] + [i for i in page_indices if i != pref_idx]

        for page_idx in page_indices:
            page = doc[page_idx]
            page_num = page_idx + 1
            pw = float(page.rect.width)
            ph = float(page.rect.height)
            sorted_blocks, is_two_col, page_ocr_needed = _extract_sorted_page_blocks(page, page_num)
            if not page_ocr_needed:
                any_text_blocks = True
            elif preferred_page == page_num:
                preferred_page_textless = True

            if not search_probe or page_ocr_needed:
                continue

            # 1a. Exact/full query match within a single block in column-first reading order
            probe_lower = search_probe.lower()
            query_lower = normalized_query.lower()
            exact_blocks: list[dict[str, Any]] = []
            probe_or_partial_blocks: list[dict[str, Any]] = []
            for blk in sorted_blocks:
                blk_lower = blk["text"].lower()
                if query_lower in blk_lower:
                    exact_blocks.append(blk)
                elif probe_lower in blk_lower or (
                    len(blk_lower) >= 24
                    and not (blk["column"] == "full" and blk["y0"] < 75.0)
                    and blk_lower in query_lower
                ):
                    probe_or_partial_blocks.append(blk)

            for blk in exact_blocks:
                clip_rect = pymupdf.Rect(
                    max(0.0, blk["x0"] - 2.0),
                    max(0.0, blk["y0"] - 2.0),
                    min(pw, blk["x1"] + 2.0),
                    min(ph, blk["y1"] + 2.0),
                )
                rects = page.search_for(search_probe, clip=clip_rect)
                if rects:
                    union_rect = rects[0]
                    for r in rects[1:]:
                        union_rect = union_rect | r
                    rx0 = round(float(union_rect.x0), 2)
                    ry0 = round(float(union_rect.y0), 2)
                    rx1 = round(float(union_rect.x1), 2)
                    ry1 = round(float(union_rect.y1), 2)
                else:
                    rx0, ry0, rx1, ry1 = blk["x0"], blk["y0"], blk["x1"], blk["y1"]

                cand_bbox: dict[str, Any] = {
                    "page": page_num,
                    "x0": rx0,
                    "y0": ry0,
                    "x1": rx1,
                    "y1": ry1,
                    "box_2d": normalize_pdf_rect_to_box_2d(
                        x0=rx0, y0=ry0, x1=rx1, y1=ry1, page_width=pw, page_height=ph
                    ),
                    "page_width": pw,
                    "page_height": ph,
                    "column": blk["column"],
                    "reading_order": blk["reading_order"],
                    "layout_mode": "two_column" if is_two_col else "single_column",
                    "spans_columns": False,
                    "ocr_fallback_required": False,
                }
                sig = (page_num, rx0, ry0, rx1, ry1)
                if fallback_bbox is None:
                    fallback_bbox = cand_bbox
                if sig not in seen_rects:
                    seen_rects.add(sig)
                    found_bbox = cand_bbox
                    if not matched_verbatim:
                        matched_verbatim = blk["text"]
                    break

            if found_bbox is not None:
                break

            # 1b. Check consecutive block pairs in reading order (e.g., quote wrapping from the
            #     bottom of the left column to the top of the right column on a two-column reprint).
            for b_idx in range(len(sorted_blocks) - 1):
                blk_a = sorted_blocks[b_idx]
                blk_b = sorted_blocks[b_idx + 1]
                clean_b = re.sub(
                    r"^(?:(?:left|right)-(?:top|bottom)|col(?:umn)?\s*\d+)\s*:\s*",
                    "",
                    blk_b["text"],
                    flags=re.IGNORECASE,
                )
                combined_text = f"{blk_a['text']} {clean_b}"
                combined_lower = combined_text.lower()
                if query_lower in combined_lower or (
                    probe_lower not in blk_a["text"].lower()
                    and probe_lower not in blk_b["text"].lower()
                    and probe_lower in combined_lower
                ):
                    spans_cols = blk_a["column"] != blk_b["column"]
                    if spans_cols:
                        # Keep primary highlight box column-safe on blk_a and attach continuation_bbox for blk_b
                        rx0, ry0, rx1, ry1 = blk_a["x0"], blk_a["y0"], blk_a["x1"], blk_a["y1"]
                    else:
                        rx0 = round(min(blk_a["x0"], blk_b["x0"]), 2)
                        ry0 = round(min(blk_a["y0"], blk_b["y0"]), 2)
                        rx1 = round(max(blk_a["x1"], blk_b["x1"]), 2)
                        ry1 = round(max(blk_a["y1"], blk_b["y1"]), 2)

                    cont_bbox = {
                        "page": page_num,
                        "x0": blk_b["x0"],
                        "y0": blk_b["y0"],
                        "x1": blk_b["x1"],
                        "y1": blk_b["y1"],
                        "box_2d": blk_b["box_2d"],
                        "column": blk_b["column"],
                        "reading_order": blk_b["reading_order"],
                    }
                    cand_bbox = {
                        "page": page_num,
                        "x0": rx0,
                        "y0": ry0,
                        "x1": rx1,
                        "y1": ry1,
                        "box_2d": normalize_pdf_rect_to_box_2d(
                            x0=rx0, y0=ry0, x1=rx1, y1=ry1, page_width=pw, page_height=ph
                        ),
                        "page_width": pw,
                        "page_height": ph,
                        "column": blk_a["column"],
                        "reading_order": blk_a["reading_order"],
                        "layout_mode": "two_column" if is_two_col else "single_column",
                        "spans_columns": spans_cols,
                        "continuation_bbox": cont_bbox,
                        "ocr_fallback_required": False,
                    }
                    sig = (page_num, rx0, ry0, rx1, ry1)
                    if fallback_bbox is None:
                        fallback_bbox = cand_bbox
                    if sig not in seen_rects:
                        seen_rects.add(sig)
                        found_bbox = cand_bbox
                        if not matched_verbatim:
                            matched_verbatim = combined_text
                        break

            if found_bbox is not None:
                break

            # 1c. Fallback to probe or partial block containment in a single block
            for blk in probe_or_partial_blocks:
                clip_rect = pymupdf.Rect(
                    max(0.0, blk["x0"] - 2.0),
                    max(0.0, blk["y0"] - 2.0),
                    min(pw, blk["x1"] + 2.0),
                    min(ph, blk["y1"] + 2.0),
                )
                rects = page.search_for(search_probe, clip=clip_rect)
                if rects:
                    union_rect = rects[0]
                    for r in rects[1:]:
                        union_rect = union_rect | r
                    rx0 = round(float(union_rect.x0), 2)
                    ry0 = round(float(union_rect.y0), 2)
                    rx1 = round(float(union_rect.x1), 2)
                    ry1 = round(float(union_rect.y1), 2)
                else:
                    rx0, ry0, rx1, ry1 = blk["x0"], blk["y0"], blk["x1"], blk["y1"]
                cand_bbox = {
                    "page": page_num,
                    "x0": rx0,
                    "y0": ry0,
                    "x1": rx1,
                    "y1": ry1,
                    "box_2d": normalize_pdf_rect_to_box_2d(
                        x0=rx0, y0=ry0, x1=rx1, y1=ry1, page_width=pw, page_height=ph
                    ),
                    "page_width": pw,
                    "page_height": ph,
                    "column": blk["column"],
                    "reading_order": blk["reading_order"],
                    "layout_mode": "two_column" if is_two_col else "single_column",
                    "spans_columns": False,
                    "ocr_fallback_required": False,
                }
                sig = (page_num, rx0, ry0, rx1, ry1)
                if fallback_bbox is None:
                    fallback_bbox = cand_bbox
                if sig not in seen_rects:
                    seen_rects.add(sig)
                    found_bbox = cand_bbox
                    break

            if found_bbox is not None:
                break

            # 2. Fallback to page-level `page.search_for` grouped by column if block text didn't match directly
            rects = page.search_for(search_probe)
            if rects:
                mid_x = pw / 2.0
                if is_two_col:
                    left_rects = [r for r in rects if float(r.x1) <= mid_x + 18.0]
                    right_rects = [r for r in rects if float(r.x0) >= mid_x - 18.0]
                    rect_groups = [
                        ("left", left_rects),
                        ("right", right_rects),
                        ("full", [r for r in rects if r not in left_rects and r not in right_rects]),
                    ]
                else:
                    rect_groups = [("full", rects)]

                for col_name, grp in rect_groups:
                    if not grp:
                        continue
                    union_rect = grp[0]
                    for r in grp[1:]:
                        union_rect = union_rect | r
                    rx0 = round(float(union_rect.x0), 2)
                    ry0 = round(float(union_rect.y0), 2)
                    rx1 = round(float(union_rect.x1), 2)
                    ry1 = round(float(union_rect.y1), 2)
                    cand_bbox = {
                        "page": page_num,
                        "x0": rx0,
                        "y0": ry0,
                        "x1": rx1,
                        "y1": ry1,
                        "box_2d": normalize_pdf_rect_to_box_2d(
                            x0=rx0, y0=ry0, x1=rx1, y1=ry1, page_width=pw, page_height=ph
                        ),
                        "page_width": pw,
                        "page_height": ph,
                        "column": col_name,
                        "reading_order": 1,
                        "layout_mode": "two_column" if is_two_col else "single_column",
                        "spans_columns": False,
                        "ocr_fallback_required": False,
                    }
                    sig = (page_num, rx0, ry0, rx1, ry1)
                    if fallback_bbox is None:
                        fallback_bbox = cand_bbox
                    if sig not in seen_rects:
                        seen_rects.add(sig)
                        found_bbox = cand_bbox
                        break

            if found_bbox is not None:
                break
    finally:
        doc.close()

    if found_bbox is None and fallback_bbox is not None:
        found_bbox = fallback_bbox

    ocr_needed = (not any_text_blocks) or (found_bbox is None and preferred_page_textless)
    if found_bbox is None:
        target_p = preferred_page if isinstance(preferred_page, int) and preferred_page >= 1 else 1
        found_bbox = {
            "page": target_p,
            "x0": -1.0,
            "y0": -1.0,
            "x1": -1.0,
            "y1": -1.0,
            "box_2d": [0, 0, 0, 0],
            "page_width": 612.0,
            "page_height": 792.0,
            "column": "full",
            "reading_order": 0,
            "layout_mode": "scanned_ocr_fallback" if ocr_needed else "single_column",
            "spans_columns": False,
            "ocr_fallback_required": ocr_needed,
        }
        is_found = False
    else:
        is_found = True

    quote_str = matched_verbatim or raw_query
    snippet_hash = hashlib.sha256(quote_str.encode("utf-8")).hexdigest() if quote_str else ""
    anchor_obj = PdfBoundingBox(
        page_number=int(found_bbox["page"]),
        x0=float(found_bbox["x0"]) if is_found else -1.0,
        y0=float(found_bbox["y0"]) if is_found else -1.0,
        x1=float(found_bbox["x1"]) if is_found else -1.0,
        y1=float(found_bbox["y1"]) if is_found else -1.0,
        page_width=float(found_bbox["page_width"]),
        page_height=float(found_bbox["page_height"]),
        source_document_id=source_id,
        source_version=version,
        verbatim_quote=quote_str,
        snippet_sha256=snippet_hash,
    )

    return {
        "found": is_found,
        "page_number": int(found_bbox["page"]),
        "page": int(found_bbox["page"]),
        "x0": float(found_bbox["x0"]),
        "y0": float(found_bbox["y0"]),
        "x1": float(found_bbox["x1"]),
        "y1": float(found_bbox["y1"]),
        "box_2d": list(found_bbox["box_2d"]),
        "page_width": float(found_bbox["page_width"]),
        "page_height": float(found_bbox["page_height"]),
        "column": str(found_bbox.get("column", "full")),
        "reading_order": int(found_bbox.get("reading_order", 0)),
        "layout_mode": str(found_bbox.get("layout_mode", "single_column")),
        "spans_columns": bool(found_bbox.get("spans_columns", False)),
        "continuation_bbox": found_bbox.get("continuation_bbox"),
        "ocr_fallback_required": bool(found_bbox.get("ocr_fallback_required", ocr_needed)),
        "scanned_page_numbers": [int(found_bbox["page"])] if ocr_needed else [],
        "verbatim_quote": quote_str,
        "snippet_sha256": snippet_hash,
        "source_document_id": source_id,
        "source_version": version,
        "bbox": found_bbox,
        "anchor": anchor_obj,
    }


def extract_pdf_source_pack(
    pdf_path: str | Path,
    source_id: str = "SRC-UPLOADED-PDF",
    version: str = "1.0",
    passage_rules: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    """
    Opens a real PDF file with PyMuPDF (`pymupdf`):
    - When `passage_rules` is provided, locates each passage's exact bounding-box coordinates
      via `locate_supporting_quote_in_pdf(...)`, respecting two-column reading order and
      flagging scanned/textless pages (`ocr_fallback_required`).
    - When `passage_rules` is `None` or empty, extracts all non-header text blocks in
      column-first reading order (left column top-to-bottom, then right column top-to-bottom
      on two-column clinical reprints), and emits an `ocr_fallback_required=True` entry for
      any scanned/textless page without crashing.
    """
    path = Path(pdf_path)
    if not path.exists():
        raise FileNotFoundError(f"Source PDF not found: {path}")

    file_digest = sha256_file(path)
    source_pack: list[dict[str, Any]] = []
    used_page_rects: set[tuple[int, float, float, float, float]] = set()

    if passage_rules:
        for rule in passage_rules:
            verbatim = str(rule.get("verbatim_text", "")).strip()
            search_probe = str(rule.get("search_probe") or verbatim).strip()
            preferred_page = rule.get("page")
            loc = locate_supporting_quote_in_pdf(
                path,
                search_probe,
                source_id=source_id,
                version=version,
                preferred_page=preferred_page if isinstance(preferred_page, int) else None,
                used_page_rects=used_page_rects,
            )
            found_bbox = loc["bbox"]
            source_pack.append(
                {
                    "source_id": source_id,
                    "version": version,
                    "file_sha256": file_digest,
                    "verbatim_text": verbatim,
                    "keywords": list(rule.get("keywords", [])),
                    "locked_numbers": list(rule.get("locked_numbers", [])),
                    "required_qualifiers": list(rule.get("required_qualifiers", [])),
                    "ruled_out_terms": list(rule.get("ruled_out_terms", [])),
                    "device_version": rule.get("device_version"),
                    "column": loc["column"],
                    "reading_order": loc["reading_order"],
                    "layout_mode": loc["layout_mode"],
                    "ocr_fallback_required": loc["ocr_fallback_required"],
                    "bbox": found_bbox,
                }
            )
        return source_pack

    doc = pymupdf.open(str(path))
    try:
        for page_idx in range(len(doc)):
            page = doc[page_idx]
            page_num = page_idx + 1
            pw = float(page.rect.width)
            ph = float(page.rect.height)
            sorted_blocks, is_two_col, page_ocr_needed = _extract_sorted_page_blocks(page, page_num)

            if page_ocr_needed:
                source_pack.append(
                    {
                        "source_id": source_id,
                        "version": version,
                        "file_sha256": file_digest,
                        "page_number": page_num,
                        "verbatim_text": "",
                        "keywords": [],
                        "locked_numbers": [],
                        "required_qualifiers": [],
                        "ruled_out_terms": [],
                        "device_version": None,
                        "column": "full",
                        "reading_order": 0,
                        "layout_mode": "scanned_ocr_fallback",
                        "ocr_fallback_required": True,
                        "bbox": {
                            "page": page_num,
                            "x0": 0.0,
                            "y0": 0.0,
                            "x1": pw,
                            "y1": ph,
                            "box_2d": [0, 0, 1000, 1000],
                            "page_width": pw,
                            "page_height": ph,
                            "column": "full",
                            "ocr_fallback_required": True,
                        },
                    }
                )
                continue

            for blk in sorted_blocks:
                # Skip top full-width running page header if there are content blocks below it
                if blk["column"] == "full" and blk["y0"] < 75.0 and len(sorted_blocks) > 1:
                    continue
                cleaned = blk["text"]
                words = [
                    w.lower()
                    for w in re.findall(r"[a-zA-Z0-9%.+-]+", cleaned)
                    if len(w) >= 4 and w.lower() not in _PDF_STOPWORDS
                ]
                nums = re.findall(r"(?<![\d.])\d+(?:\.\d+)?%?(?!\d)", cleaned)
                locked_nums = [n for n in nums if n.endswith("%") or "." in n]
                auto_quals = _extract_auto_required_qualifiers(cleaned)
                source_pack.append(
                    {
                        "source_id": source_id,
                        "version": version,
                        "file_sha256": file_digest,
                        "page_number": page_num,
                        "verbatim_text": cleaned,
                        "keywords": list(dict.fromkeys(words[:12] + nums[:6])),
                        "locked_numbers": locked_nums,
                        "required_qualifiers": auto_quals,
                        "ruled_out_terms": [],
                        "device_version": None,
                        "column": blk["column"],
                        "reading_order": blk["reading_order"],
                        "layout_mode": "two_column" if is_two_col else "single_column",
                        "ocr_fallback_required": False,
                        "bbox": {
                            "page": page_num,
                            "x0": blk["x0"],
                            "y0": blk["y0"],
                            "x1": blk["x1"],
                            "y1": blk["y1"],
                            "box_2d": blk["box_2d"],
                            "page_width": pw,
                            "page_height": ph,
                            "column": blk["column"],
                            "reading_order": blk["reading_order"],
                            "ocr_fallback_required": False,
                        },
                    }
                )
    finally:
        doc.close()

    return source_pack


_PDF_QUALIFIER_PATTERNS: list[re.Pattern[str]] = [
    re.compile(r"\b2L\+\s+EGFR-mutated(?:\s+metastatic)?\s+NSCLC\b", flags=re.IGNORECASE),
    re.compile(r"\b95%\s+CI\b", flags=re.IGNORECASE),
    re.compile(r"\b510\(k\)\s+cleared\b", flags=re.IGNORECASE),
    re.compile(r"\bparoxysmal\s+atrial\s+fibrillation\b", flags=re.IGNORECASE),
    re.compile(r"\bsevere\s+hepatic\s+impairment\b", flags=re.IGNORECASE),
]


def _extract_auto_required_qualifiers(text: str) -> list[str]:
    quals: list[str] = []
    for pat in _PDF_QUALIFIER_PATTERNS:
        m = pat.search(text)
        if m:
            found = m.group(0).strip()
            if found not in quals:
                quals.append(found)
    return quals


def inspect_pdf_document_and_extract_source_pack(
    pdf_path: str | Path,
    source_id: str = "SRC-UPLOADED-PDF",
    version: str = "1.0",
    passage_rules: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    """
    Opens a real PDF file with PyMuPDF (`pymupdf`), extracts page dimensions, text blocks
    sorted in column-first reading order with exact 2D bounding boxes (`[x0, y0, x1, y1]`,
    `[ymin, xmin, ymax, xmax]`, and normalized `[0..1000]` `box_2d`), scanned/textless page
    `ocr_fallback_required` flags, existing PDF annotations, and builds a coordinate-bound `source_pack`.
    """
    path = Path(pdf_path)
    if not path.exists():
        raise FileNotFoundError(f"PDF file not found: {path}")

    file_digest = sha256_file(path)
    doc = pymupdf.open(str(path))
    pages_info: list[dict[str, Any]] = []
    auto_source_pack: list[dict[str, Any]] = []

    try:
        for page_idx in range(len(doc)):
            page = doc[page_idx]
            page_num = page_idx + 1
            pw = float(page.rect.width)
            ph = float(page.rect.height)
            text_blocks, is_two_col, page_ocr_needed = _extract_sorted_page_blocks(page, page_num)

            for blk in text_blocks:
                if blk["column"] == "full" and blk["y0"] < 75.0 and len(text_blocks) > 1:
                    continue
                cleaned = blk["text"]
                words = [
                    w.lower()
                    for w in re.findall(r"[a-zA-Z0-9%.+-]+", cleaned)
                    if len(w) >= 4 and w.lower() not in _PDF_STOPWORDS
                ]
                nums = re.findall(r"(?<![\d.])\d+(?:\.\d+)?%?(?!\d)", cleaned)
                locked_nums = [n for n in nums if n.endswith("%") or "." in n]
                auto_quals = _extract_auto_required_qualifiers(cleaned)
                auto_source_pack.append(
                    {
                        "source_id": source_id,
                        "version": version,
                        "file_sha256": file_digest,
                        "verbatim_text": cleaned,
                        "keywords": list(dict.fromkeys(words[:12] + nums[:6])),
                        "locked_numbers": locked_nums,
                        "required_qualifiers": auto_quals,
                        "ruled_out_terms": [],
                        "device_version": None,
                        "column": blk["column"],
                        "reading_order": blk["reading_order"],
                        "layout_mode": "two_column" if is_two_col else "single_column",
                        "ocr_fallback_required": False,
                        "bbox": {
                            "page": page_num,
                            "x0": blk["x0"],
                            "y0": blk["y0"],
                            "x1": blk["x1"],
                            "y1": blk["y1"],
                            "box_2d": blk["box_2d"],
                            "page_width": pw,
                            "page_height": ph,
                            "column": blk["column"],
                            "reading_order": blk["reading_order"],
                            "ocr_fallback_required": False,
                        },
                    }
                )

            embedded_annots: list[dict[str, Any]] = []
            for annot in page.annots() or []:
                arect = annot.rect
                info = annot.info or {}
                embedded_annots.append(
                    {
                        "page_number": page_num,
                        "type": annot.type[1] if isinstance(annot.type, tuple) and len(annot.type) > 1 else str(annot.type),
                        "x0": round(float(arect.x0), 2),
                        "y0": round(float(arect.y0), 2),
                        "x1": round(float(arect.x1), 2),
                        "y1": round(float(arect.y1), 2),
                        "ymin": round(float(arect.y0), 2),
                        "xmin": round(float(arect.x0), 2),
                        "ymax": round(float(arect.y1), 2),
                        "xmax": round(float(arect.x1), 2),
                        "title": str(info.get("title", "")),
                        "content": str(info.get("content", "")),
                    }
                )

            pages_info.append(
                {
                    "page_number": page_num,
                    "page_width": pw,
                    "page_height": ph,
                    "is_two_column": is_two_col,
                    "two_column_layout_detected": is_two_col,
                    "has_extractable_text": not page_ocr_needed,
                    "layout_mode": (
                        "scanned_ocr_fallback"
                        if page_ocr_needed
                        else ("two_column" if is_two_col else "single_column")
                    ),
                    "ocr_fallback_required": page_ocr_needed,
                    "text_blocks": text_blocks,
                    "embedded_annotations": embedded_annots,
                }
            )
    finally:
        doc.close()

    resolved_pack = (
        extract_pdf_source_pack(path, source_id=source_id, version=version, passage_rules=passage_rules)
        if passage_rules
        else auto_source_pack
    )
    ocr_fallback_pages = [p["page_number"] for p in pages_info if p.get("ocr_fallback_required")]
    return {
        "pdf_path": str(path),
        "file_sha256": file_digest,
        "page_count": len(pages_info),
        "ocr_fallback_required": len(ocr_fallback_pages) > 0,
        "ocr_fallback_pages": ocr_fallback_pages,
        "scanned_page_numbers": ocr_fallback_pages,
        "pages": pages_info,
        "source_pack": resolved_pack,
    }


def render_annotated_reference_pdf(
    source_pdf_path: str | Path,
    output_pdf_path: str | Path,
    claims: list[ClaimAssessment],
) -> dict[str, Any]:
    """
    Opens `source_pdf_path` with PyMuPDF, embeds real PDF highlight annotations
    and popup metadata for every valid `PdfBoundingBox` anchor across `claims`,
    saves the annotated PDF to `output_pdf_path`, and returns coordinate-verified
    annotation and page preview metadata.
    """
    src_path = Path(source_pdf_path)
    out_path = Path(output_pdf_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = pymupdf.open(str(src_path))
    annotation_count = 0
    annotations_detail: list[dict[str, Any]] = []

    try:
        for claim in claims:
            for anchor in claim.anchors:
                if not anchor.anchor_valid:
                    continue
                page_idx = anchor.page_number - 1
                if 0 <= page_idx < len(doc):
                    page = doc[page_idx]
                    rect = pymupdf.Rect(anchor.x0, anchor.y0, anchor.x1, anchor.y1)
                    annot = page.add_highlight_annot(rect)
                    popup_title = f"MLR Genie Anchor ({claim.claim_id})"
                    popup_content = (
                        f"Claim: {claim.sentence_text}\n"
                        f"Support: {claim.support_status.value}\n"
                        f"Quote: {anchor.verbatim_quote}"
                    )
                    annot.set_info(
                        title=popup_title,
                        content=popup_content,
                    )
                    annot.update()
                    annotation_count += 1
                    annotations_detail.append(
                        {
                            "annotation_id": f"annot_p{anchor.page_number}_{annotation_count}",
                            "claim_id": claim.claim_id,
                            "page_number": anchor.page_number,
                            "x0": round(float(anchor.x0), 2),
                            "y0": round(float(anchor.y0), 2),
                            "x1": round(float(anchor.x1), 2),
                            "y1": round(float(anchor.y1), 2),
                            "ymin": round(float(anchor.y0), 2),
                            "xmin": round(float(anchor.x0), 2),
                            "ymax": round(float(anchor.y1), 2),
                            "xmax": round(float(anchor.x1), 2),
                            "page_width": float(anchor.page_width),
                            "page_height": float(anchor.page_height),
                            "source_document_id": anchor.source_document_id,
                            "source_version": anchor.source_version,
                            "verbatim_quote": anchor.verbatim_quote,
                            "snippet_sha256": anchor.snippet_sha256,
                            "claim_sentence": claim.sentence_text,
                            "support_status": claim.support_status.value,
                            "anchor_valid": anchor.anchor_valid,
                            "pdf_annot_type": "Highlight",
                            "popup_title": popup_title,
                            "popup_content": popup_content,
                        }
                    )
        doc.save(str(out_path))
    finally:
        doc.close()

    inspected = inspect_pdf_document_and_extract_source_pack(out_path)
    return {
        "annotated_pdf_path": str(out_path),
        "annotation_count": annotation_count,
        "sha256": sha256_file(out_path),
        "page_count": inspected["page_count"],
        "annotations": annotations_detail,
        "pages": inspected["pages"],
    }


def create_working_draft_docx(
    docx_path: str | Path,
    title: str,
    paragraphs: list[str],
) -> dict[str, Any]:
    """
    Creates a real `.docx` working document on disk using `python-docx`.
    """
    out_path = Path(docx_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    document.add_heading(title, level=1)
    for para in paragraphs:
        document.add_paragraph(para)
    document.save(str(out_path))
    return {
        "docx_path": str(out_path),
        "paragraph_count": len(paragraphs),
        "sha256": sha256_file(out_path),
    }


def _iter_section_header_footer_containers(document: Any) -> list[tuple[str, Any]]:
    """
    Returns active `(kind, container)` pairs (`kind` in `("header", "footer")`) across all
    document sections without materializing unlinked first-page or even-page XML parts when
    those settings are disabled.
    """
    containers: list[tuple[str, Any]] = []
    odd_even = bool(
        getattr(getattr(document, "settings", None), "odd_and_even_pages_header_footer", False)
    )
    for section in getattr(document, "sections", []):
        hf_specs: list[tuple[str, str]] = [("header", "header"), ("footer", "footer")]
        if getattr(section, "different_first_page_header_footer", False):
            hf_specs.extend([("header", "first_page_header"), ("footer", "first_page_footer")])
        if odd_even:
            hf_specs.extend([("header", "even_page_header"), ("footer", "even_page_footer")])
        for kind, attr_name in hf_specs:
            hf = getattr(section, attr_name, None)
            if hf is not None:
                containers.append((kind, hf))
    return containers


def _iter_all_docx_paragraphs(document: Any) -> list[Any]:
    """
    Returns all unique paragraph objects across:
    1. Top-level document body paragraphs (`document.paragraphs`)
    2. Word table cells (`document.tables`)
    3. Section headers and footers (`section.header`, `section.footer`, plus header/footer tables)
    Deduplicated by underlying XML `<w:p>` element.
    """
    seen_p_elems: set[Any] = set()
    all_paragraphs: list[Any] = []

    def _add_paragraph(p: Any) -> None:
        p_el = p._p
        if p_el not in seen_p_elems:
            seen_p_elems.add(p_el)
            all_paragraphs.append(p)

    def _add_tables(tables: Any) -> None:
        for tbl in tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        _add_paragraph(p)

    for p in document.paragraphs:
        _add_paragraph(p)
    _add_tables(document.tables)

    for _, hf in _iter_section_header_footer_containers(document):
        try:
            for p in hf.paragraphs:
                _add_paragraph(p)
            _add_tables(hf.tables)
        except Exception:
            continue

    return all_paragraphs


def parse_docx_working_document(docx_path: str | Path) -> dict[str, Any]:
    """
    Parses a real `.docx` working document from disk, returning its full body text,
    paragraph list, table cell texts, section header/footer texts, and file SHA-256 digest.
    """
    path = Path(docx_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {path}")

    document = docx.Document(str(path))
    paragraphs: list[str] = []
    for p in document.paragraphs:
        text = p.text.strip()
        if text and not p.style.name.lower().startswith("heading"):
            paragraphs.append(text)

    table_cells: list[str] = []
    seen_tc: set[Any] = set()
    for tbl in document.tables:
        for row in tbl.rows:
            for cell in row.cells:
                tc_el = cell._tc
                if tc_el in seen_tc:
                    continue
                seen_tc.add(tc_el)
                ctext = cell.text.strip()
                if ctext:
                    table_cells.append(ctext)

    headers: list[str] = []
    footers: list[str] = []
    seen_hf_p: set[Any] = set()
    for kind, hf in _iter_section_header_footer_containers(document):
        target_list = headers if kind == "header" else footers
        for p in hf.paragraphs:
            if p._p in seen_hf_p:
                continue
            seen_hf_p.add(p._p)
            htext = p.text.strip()
            if htext and htext not in target_list:
                target_list.append(htext)
        for tbl in hf.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p._p in seen_hf_p:
                            continue
                        seen_hf_p.add(p._p)
                        htext = p.text.strip()
                        if htext and htext not in target_list:
                            target_list.append(htext)

    combined_blocks = list(paragraphs)
    for extra_text in table_cells + headers + footers:
        if extra_text not in combined_blocks:
            combined_blocks.append(extra_text)
    body_text = "\n".join(combined_blocks)
    return {
        "docx_path": str(path),
        "paragraphs": paragraphs,
        "table_cells": table_cells,
        "headers": headers,
        "footers": footers,
        "body_text": body_text,
        "file_sha256": sha256_file(path),
    }


def _extract_paragraph_styled_chars(p: Any) -> list[tuple[str, Any]]:
    """
    Extracts a character-level stream `(char, rPr_element_or_None)` from a Word paragraph `p`
    so inline run formatting (`<w:rPr>` bold, italic, underline, superscripts `®`, `†`, `‡`,
    font/color) is preserved across edits.
    """
    styled: list[tuple[str, Any]] = []
    runs = list(p.runs)
    if runs:
        for r in runs:
            r_text = r.text or ""
            if not r_text:
                continue
            rpr_copy = copy.deepcopy(r._r.rPr) if r._r.rPr is not None else None
            for ch in r_text:
                styled.append((ch, rpr_copy))
    elif p.text:
        for ch in p.text:
            styled.append((ch, None))
    return styled


_SPECIAL_SUPERSCRIPT_GLYPHS = ("®", "†", "‡", "™", "§", "*", "℠")


def _rpr_signature(rpr: Any) -> str:
    return str(rpr.xml) if rpr is not None else ""


def _dominant_alnum_rpr(chars_slice: list[tuple[str, Any]]) -> tuple[Any, bool]:
    """
    Returns `(dominant_rpr, found_any_alnum)` for alphanumeric characters in `chars_slice`,
    treating `None` (unformatted run) as a first-class style so a single footnote superscript
    (`†`, `®`) or minor sub-run never smears across an otherwise unformatted replacement span.
    """
    counts: dict[str, int] = {}
    exemplar: dict[str, Any] = {}
    for ch, rpr in chars_slice:
        if ch.isalnum() and ch not in _SPECIAL_SUPERSCRIPT_GLYPHS:
            sig = _rpr_signature(rpr)
            counts[sig] = counts.get(sig, 0) + 1
            if sig not in exemplar:
                exemplar[sig] = rpr
    if not counts:
        return None, False
    best_sig = max(counts.keys(), key=lambda s: (counts[s], 1 if s != "" else 0))
    return exemplar[best_sig], True


def _build_replacement_styled_chars(
    matched_chars: list[tuple[str, Any]],
    target_text: str,
    replacement_text: str,
) -> list[tuple[str, Any]]:
    """
    Maps `replacement_text` onto the `<w:rPr>` formatting of `matched_chars`:
    - Preserves exact character `<w:rPr>` on shared prefix and suffix characters
    - Preserves superscript/symbol `<w:rPr>` on special clinical/trademark glyphs (`®`, `†`, `‡`, `™`, `§`, `*`)
      without smearing superscript onto normal text
    - Preserves interior formatted tokens (e.g. bold brand name `AURENIX` or italic `p < 0.0001`)
      that appear in both `target_text` and `replacement_text`
    - Applies the dominant alphanumeric `<w:rPr>` of the replaced span (including `None` when
      normal text dominates) to newly inserted characters.
    """
    if not replacement_text:
        return []
    if not matched_chars:
        return [(ch, None) for ch in replacement_text]

    symbol_rpr: dict[str, Any] = {}
    for ch, rpr in matched_chars:
        if ch in _SPECIAL_SUPERSCRIPT_GLYPHS and rpr is not None and ch not in symbol_rpr:
            symbol_rpr[ch] = rpr

    base_rpr, _ = _dominant_alnum_rpr(matched_chars)

    prefix_len = 0
    max_pref = min(len(target_text), len(replacement_text))
    while prefix_len < max_pref and target_text[prefix_len] == replacement_text[prefix_len]:
        prefix_len += 1

    suffix_len = 0
    max_suff = min(len(target_text) - prefix_len, len(replacement_text) - prefix_len)
    while (
        suffix_len < max_suff
        and target_text[len(target_text) - 1 - suffix_len]
        == replacement_text[len(replacement_text) - 1 - suffix_len]
    ):
        suffix_len += 1

    mid_end = len(matched_chars) - suffix_len if suffix_len > 0 else len(matched_chars)
    mid_slice = matched_chars[prefix_len:mid_end]
    mid_rpr, has_mid_alnum = _dominant_alnum_rpr(mid_slice)
    if not has_mid_alnum:
        mid_rpr = base_rpr

    # Detect contiguous formatted runs inside `matched_chars` (e.g., bold brand or italic p-value)
    # that also appear verbatim inside `replacement_text`.
    preserved_token_rpr: dict[int, Any] = {}
    for run_str, run_rpr in _coalesce_styled_chars(matched_chars):
        stripped_tok = run_str.strip()
        if (
            run_rpr is not None
            and len(stripped_tok) >= 2
            and any(c.isalnum() for c in stripped_tok)
            and stripped_tok != target_text.strip()
        ):
            pos = replacement_text.find(stripped_tok)
            if pos != -1:
                for offset in range(len(stripped_tok)):
                    preserved_token_rpr[pos + offset] = run_rpr

    result: list[tuple[str, Any]] = []
    for idx, ch in enumerate(replacement_text):
        if idx < prefix_len:
            chosen_rpr = matched_chars[idx][1]
        elif suffix_len > 0 and idx >= len(replacement_text) - suffix_len:
            offset_from_end = len(replacement_text) - 1 - idx
            chosen_rpr = matched_chars[len(matched_chars) - 1 - offset_from_end][1]
        elif ch in symbol_rpr:
            chosen_rpr = symbol_rpr[ch]
        elif idx in preserved_token_rpr:
            chosen_rpr = preserved_token_rpr[idx]
        else:
            chosen_rpr = mid_rpr
        result.append((ch, chosen_rpr))
    return result


def _coalesce_styled_chars(styled_chars: list[tuple[str, Any]]) -> list[tuple[str, Any]]:
    if not styled_chars:
        return []
    runs: list[tuple[str, Any]] = []
    curr_chars: list[str] = [styled_chars[0][0]]
    curr_rpr: Any = styled_chars[0][1]
    curr_sig: str = _rpr_signature(curr_rpr)
    for ch, rpr in styled_chars[1:]:
        sig = _rpr_signature(rpr)
        if rpr is curr_rpr or sig == curr_sig:
            curr_chars.append(ch)
        else:
            runs.append(("".join(curr_chars), curr_rpr))
            curr_chars = [ch]
            curr_rpr = rpr
            curr_sig = sig
    runs.append(("".join(curr_chars), curr_rpr))
    return runs


def _append_formatted_run(
    parent_elem: Any,
    text: str,
    rpr: Any = None,
    *,
    is_del: bool = False,
) -> None:
    if not text:
        return
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = OxmlElement("w:delText" if is_del else "w:t")
    if text.startswith(" ") or text.endswith(" "):
        t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    parent_elem.append(r)


def write_clean_revised_docx(
    source_docx_path: str | Path,
    output_docx_path: str | Path,
    accepted_comments: list[ReviewComment],
) -> dict[str, Any]:
    """
    Opens `source_docx_path`, applies all placed `accepted_comments` spans across
    document body paragraphs, table cells, and section headers/footers while preserving
    character-level `<w:rPr>` run formatting (bold, italic, underline, superscripts `®`/`†`/`‡`),
    and writes a clean revised `.docx` to `output_docx_path`.
    """
    src_path = Path(source_docx_path)
    out_path = Path(output_docx_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    document = docx.Document(str(src_path))
    all_paragraphs = _iter_all_docx_paragraphs(document)
    para_styled: list[list[tuple[str, Any]]] = [
        _extract_paragraph_styled_chars(p) for p in all_paragraphs
    ]
    para_modified: list[bool] = [False] * len(all_paragraphs)
    applied_span_count = 0

    for comment in accepted_comments:
        if comment.human_decision != "accepted" or comment.placement_status != "placed":
            continue
        for span in comment.proposed_spans:
            if not span.target_text:
                continue
            for p_idx, styled_chars in enumerate(para_styled):
                plain_text = "".join(ch for ch, _ in styled_chars)
                match_pos = plain_text.find(span.target_text)
                if match_pos != -1:
                    end_pos = match_pos + len(span.target_text)
                    matched_slice = styled_chars[match_pos:end_pos]
                    repl_styled = _build_replacement_styled_chars(
                        matched_slice,
                        span.target_text,
                        span.replacement_text,
                    )
                    para_styled[p_idx] = (
                        styled_chars[:match_pos] + repl_styled + styled_chars[end_pos:]
                    )
                    para_modified[p_idx] = True
                    applied_span_count += 1
                    break

    for p_idx, p in enumerate(all_paragraphs):
        if not para_modified[p_idx]:
            continue
        p_elem = p._p
        for child in list(p_elem):
            if child.tag.endswith("}r") or child.tag.endswith("}del") or child.tag.endswith("}ins"):
                p_elem.remove(child)
        for run_text, run_rpr in _coalesce_styled_chars(para_styled[p_idx]):
            _append_formatted_run(p_elem, run_text, run_rpr, is_del=False)

    document.save(str(out_path))
    return {
        "clean_docx_path": str(out_path),
        "applied_span_count": applied_span_count,
        "sha256": sha256_file(out_path),
    }


def _append_tracked_del_and_ins(
    p_elem: Any,
    del_styled_chars: list[tuple[str, Any]],
    ins_styled_chars: list[tuple[str, Any]],
    change_id_start: int,
    author: str,
    date_iso: str,
) -> int:
    del_el = OxmlElement("w:del")
    del_el.set(qn("w:id"), str(change_id_start))
    del_el.set(qn("w:author"), author)
    del_el.set(qn("w:date"), date_iso)
    for run_text, run_rpr in _coalesce_styled_chars(del_styled_chars):
        _append_formatted_run(del_el, run_text, run_rpr, is_del=True)
    p_elem.append(del_el)

    ins_el = OxmlElement("w:ins")
    ins_el.set(qn("w:id"), str(change_id_start + 1))
    ins_el.set(qn("w:author"), author)
    ins_el.set(qn("w:date"), date_iso)
    for run_text, run_rpr in _coalesce_styled_chars(ins_styled_chars):
        _append_formatted_run(ins_el, run_text, run_rpr, is_del=False)
    p_elem.append(ins_el)

    return change_id_start + 2


def write_tracked_changes_redline_docx(
    source_docx_path: str | Path,
    output_redline_docx_path: str | Path,
    accepted_comments: list[ReviewComment],
    date_iso: str | None = None,
) -> dict[str, Any]:
    """
    Generates a genuine OpenXML Word Track Changes redline `.docx` (`<w:del>` + `<w:ins>`)
    from `source_docx_path` across top-level paragraphs, Word table cells, and section
    headers/footers while preserving character-level `<w:rPr>` formatting on unaffected runs,
    deleted spans, and inserted spans.
    """
    src_path = Path(source_docx_path)
    out_path = Path(output_redline_docx_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    ts = date_iso or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    document = docx.Document(str(src_path))
    all_paragraphs = _iter_all_docx_paragraphs(document)
    para_segments: list[list[tuple[Any, ...]]] = [
        [("text", _extract_paragraph_styled_chars(p))] for p in all_paragraphs
    ]
    para_modified: list[bool] = [False] * len(all_paragraphs)

    change_id = 100
    tracked_replacements = 0

    for comment in accepted_comments:
        if comment.human_decision != "accepted" or comment.placement_status != "placed":
            continue
        author = comment.human_actor or comment.reviewer_name
        for span in comment.proposed_spans:
            if not span.target_text:
                continue
            placed_this_span = False
            for p_idx, segs in enumerate(para_segments):
                for s_idx, seg in enumerate(segs):
                    if seg[0] != "text":
                        continue
                    styled_chars: list[tuple[str, Any]] = seg[1]
                    plain_text = "".join(ch for ch, _ in styled_chars)
                    match_pos = plain_text.find(span.target_text)
                    if match_pos != -1:
                        end_pos = match_pos + len(span.target_text)
                        prefix_chars = styled_chars[:match_pos]
                        del_chars = styled_chars[match_pos:end_pos]
                        suffix_chars = styled_chars[end_pos:]
                        ins_chars = _build_replacement_styled_chars(
                            del_chars,
                            span.target_text,
                            span.replacement_text,
                        )
                        new_segs: list[tuple[Any, ...]] = []
                        if prefix_chars:
                            new_segs.append(("text", prefix_chars))
                        new_segs.append(("redline", del_chars, ins_chars, author))
                        if suffix_chars:
                            new_segs.append(("text", suffix_chars))
                        segs[s_idx : s_idx + 1] = new_segs
                        para_modified[p_idx] = True
                        tracked_replacements += 1
                        placed_this_span = True
                        break
                if placed_this_span:
                    break

    for p_idx, p in enumerate(all_paragraphs):
        if not para_modified[p_idx]:
            continue
        p_elem = p._p
        for child in list(p_elem):
            if child.tag.endswith("}r") or child.tag.endswith("}del") or child.tag.endswith("}ins"):
                p_elem.remove(child)
        for seg in para_segments[p_idx]:
            if seg[0] == "text":
                for run_text, run_rpr in _coalesce_styled_chars(seg[1]):
                    _append_formatted_run(p_elem, run_text, run_rpr, is_del=False)
            elif seg[0] == "redline":
                change_id = _append_tracked_del_and_ins(
                    p_elem=p_elem,
                    del_styled_chars=seg[1],
                    ins_styled_chars=seg[2],
                    change_id_start=change_id,
                    author=seg[3],
                    date_iso=ts,
                )

    document.save(str(out_path))
    return {
        "redline_docx_path": str(out_path),
        "tracked_replacements": tracked_replacements,
        "sha256": sha256_file(out_path),
    }


def build_resubmission_package_zip(
    *,
    job: JobRecord,
    original_docx_path: str | Path,
    original_pdf_path: str | Path,
    clean_revised_docx_path: str | Path,
    redline_docx_path: str | Path,
    annotated_pdf_path: str | Path,
    output_zip_path: str | Path,
) -> dict[str, Any]:
    """
    Creates a real, manifest-verified resubmission `.zip` bundle bound to `job`'s
    exact revision hash and review disposition log.
    """
    out_zip = Path(output_zip_path)
    out_zip.parent.mkdir(parents=True, exist_ok=True)

    latest_rev = job.revisions[-1] if job.revisions else None
    base_rev = job.revisions[0] if job.revisions else None

    bound_rev_sha = latest_rev.sha256_hash if latest_rev else None
    is_bound_to_current = all(
        c.bound_version == job.current_version
        and (not bound_rev_sha or c.bound_revision_sha256 == bound_rev_sha)
        and not c.stale
        for c in job.claims
    )

    evidence_matrix = {
        "job_id": job.job_id,
        "tenant_id": job.tenant_id,
        "current_version": job.current_version,
        "bound_version": job.current_version,
        "revision_sha256": bound_rev_sha,
        "bound_revision_sha256": bound_rev_sha,
        "evidence_bound_to_current_version": is_bound_to_current,
        "claims": [c.model_dump() for c in job.claims],
    }
    comment_response_log = {
        "job_id": job.job_id,
        "total_comments": len(job.comments),
        "all_accounted_for": all(
            c.human_decision in ("accepted", "rejected", "escalated") for c in job.comments
        ),
        "comments": [c.model_dump() for c in job.comments],
    }
    decisions_and_provenance = {
        "job_id": job.job_id,
        "state": job.state.value,
        "base_revision_id": base_rev.revision_id if base_rev else None,
        "base_sha256": base_rev.sha256_hash if base_rev else None,
        "latest_revision_id": latest_rev.revision_id if latest_rev else None,
        "latest_sha256": bound_rev_sha,
        "revisions": [r.model_dump() for r in job.revisions],
        "telemetry_receipts": job.metadata.get("telemetry_receipts", []),
        "rd_handoff_boilerplate": job.metadata.get("rd_handoff_boilerplate"),
    }

    file_entries = [
        ("original/working_draft.docx", Path(original_docx_path)),
        ("original/clinical_reference_pack.pdf", Path(original_pdf_path)),
        ("revised/clean_revised.docx", Path(clean_revised_docx_path)),
        ("revised/redline_tracked_changes.docx", Path(redline_docx_path)),
        ("references/annotated_reference_pack.pdf", Path(annotated_pdf_path)),
    ]

    manifest_files: dict[str, dict[str, Any]] = {}
    with zipfile.ZipFile(out_zip, "w", zipfile.ZIP_DEFLATED) as zf:
        for arc_name, disk_path in file_entries:
            data = disk_path.read_bytes()
            zf.writestr(arc_name, data)
            manifest_files[arc_name] = {
                "bytes": len(data),
                "sha256": hashlib.sha256(data).hexdigest(),
            }

        for arc_json_name, obj in [
            ("review/evidence_matrix.json", evidence_matrix),
            ("review/comment_response_log.json", comment_response_log),
            ("review/decisions_and_provenance.json", decisions_and_provenance),
        ]:
            raw_bytes = json.dumps(obj, indent=2).encode("utf-8")
            zf.writestr(arc_json_name, raw_bytes)
            manifest_files[arc_json_name] = {
                "bytes": len(raw_bytes),
                "sha256": hashlib.sha256(raw_bytes).hexdigest(),
            }

        manifest_payload = {
            "package_type": "mlr_genie_resubmission_bundle_v1",
            "job_id": job.job_id,
            "tenant_id": job.tenant_id,
            "bound_version": job.current_version,
            "bound_revision_sha256": bound_rev_sha,
            "evidence_bound_to_current_version": is_bound_to_current,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "files": manifest_files,
        }
        zf.writestr("manifest.json", json.dumps(manifest_payload, indent=2).encode("utf-8"))

    return {
        "zip_path": str(out_zip),
        "zip_bytes": out_zip.stat().st_size,
        "zip_sha256": sha256_file(out_zip),
        "bound_version": job.current_version,
        "bound_revision_sha256": bound_rev_sha,
        "evidence_bound_to_current_version": is_bound_to_current,
        "manifest": manifest_payload,
    }
