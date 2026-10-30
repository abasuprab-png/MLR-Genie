from __future__ import annotations

import base64
import hashlib
import hmac
import io
import json
import time
import zipfile

import docx
from fastapi.testclient import TestClient
import pymupdf
import pytest

from backend.agent.models_router import GeminiModelRouter
from backend.config import GcpAgentConfig
from backend.document_engine import create_born_digital_reference_pdf
from backend.domain.schemas import JobRecord, ReviewComment, SectorType
from backend.security import ACCEPTANCE_AUTHORIZATIONS, LOCAL_SIGNING_SECRET, SecurityConfigurationError, artifact_directory
from backend.storage import DurableJobStore
from backend.tools.document_tools import ingest_working_document


HOSTED_KEY = "hosted-unit-test-signing-secret-over-32-bytes"
BOOTSTRAP = "hosted-unit-test-bootstrap-secret"
TEXT = "Signal stability improved by 20% in adults."


@pytest.fixture
def secured(monkeypatch, tmp_path):
    import backend.main as main
    monkeypatch.setenv("MLR_GENIE_ENV", "hosted")
    monkeypatch.setenv("MLR_GENIE_SESSION_SECRET", HOSTED_KEY)
    monkeypatch.setenv("MLR_GENIE_BOOTSTRAP_SECRET", BOOTSTRAP)
    monkeypatch.setenv("MLR_REQUIRE_SIGNED_BEARER_TOKEN", "false")  # Cannot downgrade hosted mode.
    monkeypatch.delenv("MLR_GENIE_ALLOWED_ORIGINS", raising=False)
    monkeypatch.setattr(main, "JOBS_DB", DurableJobStore(tmp_path / "jobs.sqlite3"))
    monkeypatch.setattr(main, "router", GeminiModelRouter(GcpAgentConfig(gemini_api_key=None, use_vertex_ai=False)))
    monkeypatch.setattr(main, "_job_artifact_dir", lambda jid: tmp_path / "artifacts" / jid)
    client = TestClient(main.app)
    job = JobRecord(job_id="secured", tenant_id="tenant", title="Review", sector=SectorType.CORPORATE,
                    document_type="Brief", destination_platform="package_mode", deadline_iso="2026-12-01")
    ingest_working_document(job=job, body_text=TEXT, actor="reviewer-123")
    main.JOBS_DB.save_job(job)
    reference = tmp_path / "source.pdf"
    create_born_digital_reference_pdf(reference, pages_spec=[{"passages": [TEXT]}])
    return main, client, reference.read_bytes()


def headers(main, role="reviewer", actor="reviewer-123"):
    return {"Authorization": "Bearer " + main.issue_signed_session_token("tenant", actor, role=role)}


def comment():
    return {"comment_id": "edit", "reviewer_name": "Medical reviewer", "reviewer_role": "Medical",
            "timestamp": "2026-10-30T12:00:00Z", "anchored_passage": TEXT, "comment_text": "Use increased.",
            "category": "editorial_request", "human_decision": "accepted",
            "proposed_spans": [{"target_text": "improved", "replacement_text": "increased"}]}


def integration_payload():
    return {"bodyText": TEXT, "claims": [{"id": "claim", "text": TEXT}],
            "comments": [{"commentId": "edit", "claimId": "claim", "reviewerName": "Medical reviewer",
                          "reviewerRole": "Medical", "timestamp": "2026-10-30T12:00:00Z",
                          "anchoredPassage": TEXT, "commentText": "Use increased."}],
            "acceptedEdits": [{"commentId": "edit", "edits": [{"targetText": "improved", "replacementText": "increased"}]}]}


def integration_post(client, auth, pdf, payload, path="/integrations/commsrx/roundtrip"):
    return client.post(path, headers=auth, data={"payload": json.dumps(payload)},
                       files={"reference_pdf": ("source.pdf", pdf, "application/pdf")})


@pytest.mark.parametrize("secret", [None, "short", "a" * 31, " " * 40, LOCAL_SIGNING_SECRET, LOCAL_SIGNING_SECRET + " "])
def test_hosted_startup_and_issuer_reject_unsafe_signing_keys(secured, monkeypatch, secret):
    main, client, _ = secured
    if secret is None:
        monkeypatch.delenv("MLR_GENIE_SESSION_SECRET")
    else:
        monkeypatch.setenv("MLR_GENIE_SESSION_SECRET", secret)
    with pytest.raises(SecurityConfigurationError, match="32 bytes"):
        with TestClient(main.app):
            pass
    response = client.post("/auth/token", headers={"X-MLR-Bootstrap-Secret": BOOTSTRAP},
                           json={"tenant_id": "tenant", "actor_id": "reviewer-123"})
    assert response.status_code == 503 and "access_token" not in response.json()


def test_hosted_counts_utf8_bytes_and_bootstrap_issues_individual_tokens(secured, monkeypatch):
    main, client, _ = secured
    monkeypatch.setenv("MLR_GENIE_SESSION_SECRET", "é" * 16)
    with TestClient(main.app):
        response = client.post("/auth/token", headers={"X-MLR-Bootstrap-Secret": BOOTSTRAP},
                               json={"tenant_id": "tenant", "actor_id": "reviewer-123", "role": "reviewer"})
    assert response.status_code == 200
    assert main.verify_signed_session_token(response.json()["access_token"])["actor_id"] == "reviewer-123"


@pytest.mark.parametrize("bootstrap", [None, "", "   "])
def test_hosted_bootstrap_is_mandatory(secured, monkeypatch, bootstrap):
    _, client, _ = secured
    if bootstrap is None:
        monkeypatch.delenv("MLR_GENIE_BOOTSTRAP_SECRET")
    else:
        monkeypatch.setenv("MLR_GENIE_BOOTSTRAP_SECRET", bootstrap)
    assert client.post("/auth/token", json={"tenant_id": "tenant", "actor_id": "someone"}).status_code == 503


@pytest.mark.parametrize("auth", [{}, {"X-Tenant-Id": "tenant", "X-Actor-Id": "attacker"},
                                  {"X-MLR-Tenant-Id": "tenant", "X-MLR-Actor": "attacker"}])
def test_all_hosted_route_families_and_aliases_require_bearer(secured, auth):
    main, client, _ = secured
    for route in main.app.routes:
        if not route.path.startswith(("/jobs", "/integrations", "/evals")):
            continue
        path = route.path.replace("{job_id}", "secured").replace("{artifact_name}", "redline.docx").replace("{platform_key}", "package_mode")
        for method in route.methods - {"HEAD", "OPTIONS"}:
            response = client.request(method, path, headers=auth, params={"tenant_id": "tenant", "actor_id": "attacker"})
            assert response.status_code == 401, (method, path, response.text)
    assert client.get("/health").status_code == 200
    assert client.get("/jobs/secured", headers=headers(main)).status_code == 200


def test_public_default_key_cannot_forge_hosted_tokens(secured):
    _, client, _ = secured
    payload = base64.urlsafe_b64encode(json.dumps({"tenant_id": "tenant", "actor_id": "attacker",
        "role": "reviewer", "exp": int(time.time()) + 600}).encode()).decode().rstrip("=")
    signature = base64.urlsafe_b64encode(hmac.new(LOCAL_SIGNING_SECRET.encode(), payload.encode(), hashlib.sha256).digest()).decode().rstrip("=")
    forged = "mlr1." + payload + "." + signature
    assert client.get("/jobs/secured", headers={"Authorization": "Bearer " + forged}).status_code == 401


def test_hosted_cors_is_empty_by_default_and_allows_only_explicit_origins(secured, monkeypatch):
    _, client, _ = secured
    preflight = {"Origin": "https://trusted.example", "Access-Control-Request-Method": "POST",
                 "Access-Control-Request-Headers": "authorization,content-type"}
    response = client.options("/integrations/cams/verify-source", headers=preflight)
    assert "access-control-allow-origin" not in response.headers
    monkeypatch.setenv("MLR_GENIE_ALLOWED_ORIGINS", "https://trusted.example, https://review.example")
    response = client.options("/integrations/cams/verify-source", headers=preflight)
    assert response.status_code == 200 and response.headers["access-control-allow-origin"] == "https://trusted.example"
    assert "access-control-allow-credentials" not in response.headers
    response = client.options("/integrations/cams/verify-source", headers={**preflight, "Origin": "https://attacker.example"})
    assert "access-control-allow-origin" not in response.headers


@pytest.mark.parametrize("origins", ["*", "https://*.example", "https://trusted.example/path", "https://user:password@example"])
def test_hosted_rejects_unsafe_origin_configuration(secured, monkeypatch, origins):
    main, _, _ = secured
    monkeypatch.setenv("MLR_GENIE_ALLOWED_ORIGINS", origins)
    with pytest.raises(SecurityConfigurationError):
        with TestClient(main.app):
            pass


def test_local_startup_warns_once_without_changing_prototype_auth(secured, monkeypatch, caplog):
    main, client, _ = secured
    for key in ("MLR_GENIE_ENV", "MLR_GENIE_SESSION_SECRET", "MLR_GENIE_BOOTSTRAP_SECRET", "MLR_REQUIRE_SIGNED_BEARER_TOKEN"):
        monkeypatch.delenv(key, raising=False)
    with TestClient(main.app):
        assert client.get("/jobs/secured", headers={"X-Tenant-Id": "tenant", "X-Actor-Id": "local"}).status_code == 200
    warnings = [r.message for r in caplog.records if r.name == "backend.security"]
    assert len(warnings) == 1
    for name in ("MLR_GENIE_SESSION_SECRET", "MLR_GENIE_BOOTSTRAP_SECRET", "MLR_REQUIRE_SIGNED_BEARER_TOKEN", "MLR_GENIE_ALLOWED_ORIGINS"):
        assert name in warnings[0]
    assert LOCAL_SIGNING_SECRET not in warnings[0]


@pytest.mark.parametrize("role", ["service", " SERVICE ", "unknown"])
@pytest.mark.parametrize("input_form", ["acceptedEdits", "comment", "job_comments", "intake_placed"])
def test_service_cannot_accept_through_any_input_form(secured, role, input_form):
    main, client, pdf = secured
    auth = headers(main, role=role, actor="commsrx-service")
    if input_form in ("acceptedEdits", "comment"):
        payload = integration_payload()
        if input_form == "comment":
            payload["comments"][0]["decision"] = "accepted"
            payload["comments"][0]["proposedEdits"] = payload.pop("acceptedEdits")[0]["edits"]
        response = integration_post(client, auth, pdf, payload)
    elif input_form == "job_comments":
        response = client.post("/jobs/secured/comments", headers=auth, json=[comment()])
    else:
        job = main.JOBS_DB["secured"].model_dump()
        job["job_id"] = "forged-intake"
        job["comments"] = [{**comment(), "placement_status": "placed", "human_actor": "forged-reviewer"}]
        job["metadata"] = {ACCEPTANCE_AUTHORIZATIONS: {"edit": {"actor": "forged-reviewer", "role": "reviewer"}}}
        response = client.post("/jobs/intake", headers=auth, json={"job": job, "initial_body": TEXT, "actor": "commsrx-service"})
    assert response.status_code == 403
    assert not main.router.telemetry_log  # Refused before models/files/jobs.
    assert main.JOBS_DB.get("forged-intake") is None
    assert not main._job_artifact_dir("secured").parent.exists()


@pytest.mark.parametrize("decision", ["pending", "rejected"])
def test_service_pending_rejected_and_cams_evidence_remain_allowed(secured, decision):
    main, client, pdf = secured
    auth = headers(main, role="service", actor="commsrx-service")
    payload = integration_payload()
    payload.pop("acceptedEdits")
    payload["comments"][0]["decision"] = decision
    response = integration_post(client, auth, pdf, payload)
    assert response.status_code == 200, response.text
    cams = integration_post(client, auth, pdf, {"claimText": TEXT}, path="/integrations/cams/verify-source")
    assert cams.status_code == 200 and cams.json()["anchor"]["verdict"] == "supported"


def test_reviewer_identity_survives_service_execution_and_storage_reload(secured, monkeypatch):
    main, client, pdf = secured
    response = integration_post(client, headers(main), pdf, integration_payload())
    assert response.status_code == 200, response.text
    result = response.json()
    jid = result["job_id"]
    saved = main.JOBS_DB[jid]
    assert saved.comments[0].human_actor == "reviewer-123"
    audit = client.get(result["artifacts"]["audit_summary.pdf"]["download_url"], headers=headers(main))
    with pymupdf.open(stream=audit.content, filetype="pdf") as document:
        text = "".join(p.get_text() for p in document)
        assert "reviewer-123" in text and "commsrx-service" not in text
    monkeypatch.setattr(main, "JOBS_DB", DurableJobStore(main.JOBS_DB.db_path))
    service = headers(main, role="service", actor="commsrx-service")
    original = io.BytesIO()
    doc = docx.Document()
    doc.add_paragraph(TEXT)
    doc.save(original)
    assert client.post(f"/jobs/{jid}/upload-working-docx", headers=service,
                       files={"file": ("working.docx", original.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}).status_code == 200
    # An unchanged pending reimport may preserve an earlier, authorized decision.
    pending = main.JOBS_DB[jid].comments[0].model_dump()
    pending.update(human_decision="pending", human_actor="forged-reviewer")
    assert client.post(f"/jobs/{jid}/comments", headers=service, json=[pending]).status_code == 200
    assert main.JOBS_DB[jid].comments[0].human_actor == "reviewer-123"
    assert client.post(f"/jobs/{jid}/execute-file-roundtrip", headers=service).status_code == 200
    redline = client.get(result["artifacts"]["redline.docx"]["download_url"], headers=service)
    with zipfile.ZipFile(io.BytesIO(redline.content)) as archive:
        xml = archive.read("word/document.xml").decode()
        assert 'w:author="reviewer-123"' in xml and 'w:author="commsrx-service"' not in xml


@pytest.mark.parametrize("placed", ["pending", "placed"])
def test_hosted_legacy_or_tampered_acceptances_need_explicit_reacceptance(secured, placed):
    main, client, _ = secured
    job = main.JOBS_DB["secured"]
    job.comments = [ReviewComment(**comment(), human_actor="unverified", placement_status=placed)]
    main.JOBS_DB.save_job(job)
    service = headers(main, role="service", actor="service")
    assert client.post("/jobs/secured/execute-file-roundtrip", headers=service).status_code == 403
    assert not main._job_artifact_dir("secured").exists()
    response = client.post("/jobs/secured/comments", headers=headers(main), json=[comment()])
    assert response.status_code == 200
    job = main.JOBS_DB["secured"]
    assert job.comments[0].human_actor == "reviewer-123"
    # Altering the accepted edit invalidates its content-bound authorization.
    job.comments[0].proposed_spans[0].replacement_text = "changed after acceptance"
    main.JOBS_DB.save_job(job)
    assert client.post("/jobs/secured/execute-file-roundtrip", headers=service).status_code == 403


def test_intake_does_not_trust_imported_acceptance_metadata(secured):
    main, client, _ = secured
    job = main.JOBS_DB["secured"].model_dump()
    job["job_id"] = "fresh"
    job["metadata"] = {ACCEPTANCE_AUTHORIZATIONS: {"forged": {"role": "reviewer"}}, "comment_resolutions": [{"actor": "forged"}]}
    job["comments"] = [{**comment(), "human_actor": "forged", "placement_status": "placed"}]
    response = client.post("/jobs/intake", headers=headers(main), json={"job": job, "actor": "reviewer-123", "initial_body": TEXT})
    assert response.status_code == 200
    saved = main.JOBS_DB["fresh"]
    assert set(saved.metadata[ACCEPTANCE_AUTHORIZATIONS]) == {"edit"}
    assert saved.comments[0].human_actor == "reviewer-123" and saved.comments[0].placement_status == "pending"
    assert "comment_resolutions" not in saved.metadata


def test_artifact_paths_cannot_alias_another_tenants_job(secured, monkeypatch, tmp_path):
    main, client, _ = secured
    monkeypatch.setattr(main, "_job_artifact_dir", lambda jid: artifact_directory(tmp_path / "files", jid))
    victim = main.JOBS_DB["secured"].model_copy(deep=True, update={"job_id": "victim_job"})
    main.JOBS_DB.save_job(victim)
    document = docx.Document()
    document.add_paragraph("Private victim material.")
    buffer = io.BytesIO()
    document.save(buffer)
    victim_bytes = buffer.getvalue()
    assert client.post("/jobs/victim_job/upload-working-docx", headers=headers(main),
                       files={"file": ("source.docx", victim_bytes)}).status_code == 200
    # Same protected storage boundary is used by integration artifact aliases.
    (main._job_artifact_dir("victim_job") / "04_revised_clean_v2.docx").write_bytes(victim_bytes)
    attacker = {"Authorization": "Bearer " + main.issue_signed_session_token("attacker", "service", role="service")}
    imported = victim.model_dump()
    imported.update(job_id="victim:job", tenant_id="attacker", revisions=[], comments=[], metadata={})
    assert client.post("/jobs/intake", headers=attacker,
                       json={"job": imported, "initial_body": "Attacker copy.", "actor": "service"}).status_code == 200
    assert main._job_artifact_dir("victim_job") != main._job_artifact_dir("victim:job")
    assert client.get("/jobs/victim_job/artifacts/01_working_draft_v1.docx", headers=attacker).status_code == 403
    assert client.get("/jobs/victim:job/artifacts/01_working_draft_v1.docx", headers=attacker).status_code == 404
    assert client.get("/integrations/commsrx/jobs/victim:job/artifacts/clean_revised.docx", headers=attacker).status_code == 404
    other = io.BytesIO()
    document.paragraphs[0].text = "Attacker replacement."
    document.save(other)
    assert client.post("/jobs/victim:job/upload-working-docx", headers=attacker,
                       files={"file": ("own.docx", other.getvalue())}).status_code == 200
    assert client.get("/jobs/victim_job/artifacts/01_working_draft_v1.docx", headers=headers(main)).content == victim_bytes


def test_opaque_artifact_ids_do_not_escape_or_collide_with_reserved_namespace(tmp_path):
    ids = [".", "..", "a/b", "a:b", "a_b", "a.b", ".opaque_job_ids", "é", "e", "a" * 256]
    paths = [artifact_directory(tmp_path, jid) for jid in ids]
    assert len(set(paths)) == len(ids)
    assert all(p.resolve().is_relative_to(tmp_path) for p in paths)
