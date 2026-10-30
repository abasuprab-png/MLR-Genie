from __future__ import annotations

import json
from types import SimpleNamespace
from unittest.mock import MagicMock
import zipfile
import xml.etree.ElementTree as ET

import docx
import pytest
from fastapi.testclient import TestClient

from backend.agent.gemini_gateway import ModelAssessmentError
from backend.agent.models_router import GeminiModelRouter
from backend.config import GcpAgentConfig
from backend.document_engine import write_clean_revised_docx, write_tracked_changes_redline_docx
from backend.domain.schemas import JobRecord, ReviewComment, SectorType
from backend.storage import DurableJobStore
from backend.tools.document_tools import apply_accepted_changes, ingest_working_document, locate_and_anchor_evidence
from backend.tools.review_tools import (
    analyze_approval_day_label_delta, import_and_reconcile_comments, screen_pharmacovigilance_cv08,
)


def job():
    return JobRecord(job_id="wiring", tenant_id="tenant", title="Review", sector=SectorType.MEDTECH,
                     document_type="Brief", destination_platform="package_mode", deadline_iso="2026-12-01")


@pytest.fixture
def offline():
    return GeminiModelRouter(GcpAgentConfig(gemini_api_key=None, use_vertex_ai=False))


def structured_router(payload):
    router = GeminiModelRouter(GcpAgentConfig(gemini_api_key=None, use_vertex_ai=False))
    router._client = MagicMock()
    router._client.models.generate_content.return_value = SimpleNamespace(text=json.dumps(payload))
    return router


def comments():
    base = dict(timestamp="2026-10-30T12:00:00Z", anchored_passage="Signal is stable.",
                category="editorial_request")
    return [ReviewComment(**base, comment_id="m", reviewer_name="M", reviewer_role="Medical",
                          comment_text="Change stable to steady.",
                          proposed_spans=[{"target_text": "stable", "replacement_text": "steady"}]),
            ReviewComment(**base, comment_id="r", reviewer_name="R", reviewer_role="Regulatory",
                          comment_text="Change stable to variable.",
                          proposed_spans=[{"target_text": "stable", "replacement_text": "variable"}])]


def test_resolution_survives_reimport_restart_and_generates_real_redline(offline, tmp_path):
    record = job()
    ingest_working_document(job=record, body_text="Signal is stable.", actor="reviewer")
    initial = comments()
    assert import_and_reconcile_comments(job=record, raw_comments=initial, router=offline)["requires_human_resolution"]
    initial[0].human_decision, initial[1].human_decision = "accepted", "rejected"
    resolved = import_and_reconcile_comments(job=record, raw_comments=initial, router=offline)
    assert not resolved["requires_human_resolution"]
    apply_accepted_changes(job=record, comments=record.comments, human_actor="reviewer")
    assert record.revisions[-1].body == "Signal is steady."
    store = DurableJobStore(tmp_path / "jobs.db")
    store.save_job(record)
    record = DurableJobStore(tmp_path / "jobs.db").get_for_tenant(record.job_id, record.tenant_id)
    assert not import_and_reconcile_comments(job=record, raw_comments=comments(), router=offline)["requires_human_resolution"]
    assert [c.human_decision for c in record.comments] == ["accepted", "rejected"]
    count = len(record.revisions)
    apply_accepted_changes(job=record, comments=record.comments, human_actor="reviewer")
    assert len(record.revisions) == count
    original, clean, redline = [tmp_path / name for name in ("original.docx", "clean.docx", "redline.docx")]
    document = docx.Document()
    document.add_paragraph("Signal is stable.")
    document.save(original)
    write_clean_revised_docx(original, clean, record.comments)
    write_tracked_changes_redline_docx(original, redline, record.comments)
    assert docx.Document(clean).paragraphs[0].text == record.revisions[-1].body
    assert docx.Document(original).paragraphs[0].text == "Signal is stable."
    with zipfile.ZipFile(redline) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert "stable" in "".join(root.find(".//w:del", ns).itertext())
    assert "steady" in "".join(root.find(".//w:ins", ns).itertext())


@pytest.mark.parametrize("changed_existing", [False, True])
def test_new_conflict_reopens_only_after_resolution(offline, changed_existing):
    record = job()
    initial = comments()
    import_and_reconcile_comments(job=record, raw_comments=initial, router=offline)
    initial[0].human_decision = "accepted"
    assert not import_and_reconcile_comments(job=record, raw_comments=[initial[0]], router=offline)["requires_human_resolution"]
    newer = comments()[1]
    if not changed_existing:
        newer.comment_id = "new"
    newer.comment_text = "A new conflicting instruction."
    result = import_and_reconcile_comments(job=record, raw_comments=[newer], router=offline)
    assert result["requires_human_resolution"]
    assert record.comments[0].human_decision == "escalated"


def test_live_conflict_assessment_receives_content_but_cannot_override_human(offline):
    record = job()
    initial = comments()
    live = structured_router({"conflicts": [{"comment_ids": ["m", "r"], "rationale": "Incompatible edits."}]})
    result = import_and_reconcile_comments(job=record, raw_comments=initial, router=live)
    assert result["used_live_sdk"] and result["requires_human_resolution"]
    sent = json.loads(live._client.models.generate_content.call_args.kwargs["contents"])
    assert sent["comments"][0]["comment_text"] == initial[0].comment_text
    initial[0].human_decision = "accepted"
    assert not import_and_reconcile_comments(job=record, raw_comments=initial, router=live)["requires_human_resolution"]


def test_incompatible_acceptances_cannot_reuse_an_old_resolution(offline):
    record = job()
    initial = comments()
    initial[0].human_decision = "accepted"
    initial[1].human_decision = "rejected"
    import_and_reconcile_comments(job=record, raw_comments=initial, router=offline)
    initial[1].human_decision = "accepted"
    result = import_and_reconcile_comments(job=record, raw_comments=initial, router=offline)
    assert result["requires_human_resolution"]
    assert all(c.human_decision == "escalated" for c in record.comments)


def test_unknown_model_comment_reference_does_not_change_saved_decisions(offline):
    record = job()
    initial = comments()
    import_and_reconcile_comments(job=record, raw_comments=initial, router=offline)
    before = record.model_dump()
    router = structured_router({"conflicts": [{"comment_ids": ["m", "invented"], "rationale": "Unknown target"}]})
    with pytest.raises(ModelAssessmentError):
        import_and_reconcile_comments(job=record, raw_comments=initial, router=router)
    assert record.model_dump() == before


def label_inputs():
    return dict(prepared_scenarios=[{"scenario_id": "before", "mechanism": "The mechanism is established."}],
                final_approved_label={"mechanism": "The mechanism remains uncertain."},
                launch_kit_assets=[{"asset_id": "brief", "body": "The mechanism is established."}])


def label_payload():
    return {"fits_prepared_scenario": False, "rationale": "Mechanistic certainty changed.", "impacts": [{
        "section": "mechanism", "asset_id": "brief", "claim_quote": "The mechanism is established.",
        "label_quote": "The mechanism remains uncertain.", "impact": "revise",
        "rationale": "Certainty is overstated.", "proposed_wording": "The mechanism remains uncertain.",
    }]}


def test_label_uses_semantic_result_even_when_layer_one_misses_it():
    router = structured_router(label_payload())
    result = analyze_approval_day_label_delta(**label_inputs(), router=router)
    assert result["layer_one_affected_assets"] == []
    assert result["affected_assets"][0]["model_impacts"][0]["proposed_wording"] == "The mechanism remains uncertain."
    assert result["used_live_sdk"] and result["label_quotes_verified"]
    assert router._client.models.generate_content.call_count == 1
    inputs = json.loads(router._client.models.generate_content.call_args.kwargs["contents"])
    assert inputs["asset_claims"]["brief"] == ["The mechanism is established."]
    assert inputs["changed_sections"]["mechanism"]["approved_text"] == ["The mechanism remains uncertain."]


@pytest.mark.parametrize("invalid", ["label_quote", "claim_quote", "asset_id", "section", "missing_pair", "duplicate_pair", "empty_proposal"])
def test_label_rejects_hallucinations_and_incomplete_coverage(invalid):
    payload = label_payload()
    if invalid == "missing_pair":
        payload["impacts"] = []
    elif invalid == "duplicate_pair":
        payload["impacts"] *= 2
    elif invalid == "empty_proposal":
        payload["impacts"][0]["proposed_wording"] = ""
    else:
        payload["impacts"][0][invalid] = "Invented text"
    with pytest.raises(ModelAssessmentError):
        analyze_approval_day_label_delta(**label_inputs(), router=structured_router(payload))


def test_live_required_never_returns_local_only(offline):
    with pytest.raises(ModelAssessmentError, match="required but unavailable"):
        analyze_approval_day_label_delta(**label_inputs(), router=offline, require_live=True)
    router = structured_router(label_payload())
    router._client.models.generate_content.side_effect = RuntimeError("Provider unavailable")
    with pytest.raises(ModelAssessmentError, match="required but unavailable"):
        analyze_approval_day_label_delta(**label_inputs(), router=router)


@pytest.mark.parametrize("output", [None, "", "not JSON", "{}"])
def test_empty_or_invalid_live_response_cannot_be_replaced_by_offline_assessment(output):
    router = structured_router(label_payload())
    router._client.models.generate_content.return_value = SimpleNamespace(text=output)
    with pytest.raises(ModelAssessmentError):
        analyze_approval_day_label_delta(**label_inputs(), router=router)


def safety_payload(text):
    return {"claims": [], "safety_findings": [{"kind": "individual_adverse_event", "span": text,
        "rationale": "Individual adverse experience following exposure.", "patient_span": "My cousin",
        "product_span": "used the inhaler", "event_span": "could no longer stand"}]}


def test_paraphrased_safety_finding_is_verified_deduplicated_and_persisted(tmp_path):
    text = "My cousin used the inhaler and could no longer stand."
    record = job()
    router = structured_router(safety_payload(text))
    result = screen_pharmacovigilance_cv08(job=record, intake_text=text, router=router)
    assert result["cv08_flagged"] and result["used_live_sdk"]
    assert record.state.value == "hold" and record.pv_icrs_flagged
    assert record.revisions == [] and record.claims == []
    updated = safety_payload(text)
    updated["safety_findings"][0]["rationale"] = "The same event described with another rationale."
    router._client.models.generate_content.return_value = SimpleNamespace(text=json.dumps(updated))
    screen_pharmacovigilance_cv08(job=record, intake_text=text, router=router)
    assert len(record.safety_findings) == 1
    finding = record.safety_findings[0]
    assert text[finding.start:finding.end] == finding.span
    store = DurableJobStore(tmp_path / "pv.db")
    store.save_job(record)
    assert store.get_for_tenant(record.job_id, record.tenant_id).safety_findings == record.safety_findings


@pytest.mark.parametrize("bad_span", ["span", "patient_span", "product_span", "event_span", "claim"])
def test_invalid_safety_citations_do_not_mutate_job(bad_span):
    text = "My cousin used the inhaler and could no longer stand."
    payload = safety_payload(text)
    if bad_span == "claim":
        payload["claims"] = [{"span": "Invented claim", "rationale": "Invented"}]
    else:
        payload["safety_findings"][0][bad_span] = "Invented evidence"
    record = job()
    before = record.model_dump()
    with pytest.raises(ModelAssessmentError):
        screen_pharmacovigilance_cv08(job=record, intake_text=text, router=structured_router(payload))
    assert record.model_dump() == before


def test_complete_copy_and_extracted_safety_claim_enter_active_review(offline):
    text = "Important Safety Information: Avoid handling while wet."
    payload = {"claims": [{"span": text, "rationale": "Safety instruction."}], "safety_findings": [{
        "kind": "safety_statement", "span": text, "rationale": "Requires review.",
        "patient_span": "", "product_span": "", "event_span": "",
    }]}
    record = job()
    router = structured_router(payload)
    result = screen_pharmacovigilance_cv08(job=record, intake_text=text, router=router)
    assert not result["cv08_flagged"]
    ingest_working_document(job=record, body_text=text, actor="reviewer")
    locate_and_anchor_evidence(job=record, source_pack=[], router=offline)
    assert record.state.value == "in_review"
    assert [c.sentence_text for c in record.claims] == [text]
    assert record.claims[0].support_status.value == "unresolved"
    # A later unrelated revision cannot inherit the previous extracted span.
    ingest_working_document(job=record, body_text="Changed copy.", actor="reviewer")
    locate_and_anchor_evidence(job=record, source_pack=[], router=offline)
    assert all(c.sentence_text != text for c in record.claims)
    long_copy = "Context. " * 50 + text
    screen_pharmacovigilance_cv08(job=job(), intake_text=long_copy, router=router)
    assert json.loads(router._client.models.generate_content.call_args.kwargs["contents"])["submitted_copy"] == long_copy


def test_required_live_failures_are_explicit_http_errors(monkeypatch, tmp_path, offline):
    import backend.main as main
    monkeypatch.setattr(main, "JOBS_DB", DurableJobStore(tmp_path / "api.db"))
    monkeypatch.setattr(main, "router", offline)
    monkeypatch.setenv("MLR_REQUIRE_LIVE_GEMINI", "true")
    main.JOBS_DB.save_job(job())
    client = TestClient(main.app)
    headers = {"X-Tenant-Id": "tenant", "X-Actor-Id": "reviewer"}
    result = client.post("/jobs/wiring/diff-label", headers=headers, json=label_inputs())
    assert result.status_code == 502 and result.json()["assessment_failed"]
    assert "last_approval_day_label_delta" not in main.JOBS_DB["wiring"].metadata
    another = job().model_copy(update={"job_id": "intake_failure"})
    result = client.post("/jobs/intake", headers=headers, json={"job": another.model_dump(), "initial_body": "Copy", "actor": "reviewer"})
    assert result.status_code == 502 and main.JOBS_DB.get("intake_failure") is None
