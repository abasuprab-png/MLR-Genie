"""API restart/instance checks using real PostgreSQL and an injected GCS SDK fake.

Set MLR_TEST_POSTGRES_DSN to an isolated PostgreSQL test database with its
mlr_genie schema bootstrapped. These tests do not exercise the GCS cloud service.
"""
from __future__ import annotations

from copy import deepcopy
import hashlib
import io
import json
import os
from types import SimpleNamespace
from uuid import uuid4
import zipfile

import docx
from fastapi.testclient import TestClient
import pymupdf
import pytest

from backend.agent.models_router import GeminiModelRouter
from backend.config import GcpAgentConfig
from backend.document_engine import create_born_digital_reference_pdf
from backend.domain.schemas import JobRecord, SectorType


CLAIM = "Signal stability improved by 20% in adults."
DOCX_MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DOWNSTREAM = {
    "04_revised_clean_v2.docx", "05_tracked_changes_redline_v1_to_v2.docx",
    "06_resubmission_package.zip", "07_audit_summary.pdf",
}


def docx_bytes(text: str) -> bytes:
    document = docx.Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def roundtrip_payload() -> dict:
    return {
        "bodyText": CLAIM,
        "claims": [{"id": "claim-1", "text": CLAIM}],
        "comments": [{
            "commentId": "edit-1", "claimId": "claim-1", "reviewerName": "Individual reviewer",
            "reviewerRole": "Medical", "timestamp": "2026-10-30T12:00:00Z",
            "anchoredPassage": CLAIM, "commentText": "Use increased rather than improved.",
        }],
        "acceptedEdits": [{
            "commentId": "edit-1",
            "edits": [{"targetText": "improved", "replacementText": "increased"}],
        }],
    }


@pytest.fixture
def shared_api(monkeypatch, tmp_path):
    dsn = os.getenv("MLR_TEST_POSTGRES_DSN")
    if not dsn:
        pytest.skip("Set MLR_TEST_POSTGRES_DSN for real PostgreSQL API persistence checks.")

    # Configure a safe import before replacing the module's runtime dependencies.
    monkeypatch.setenv("MLR_GENIE_ENV", "hosted")
    monkeypatch.setenv("DATABASE_URL", "sqlite:///" + str(tmp_path / "import.sqlite3"))
    monkeypatch.setenv("MLR_GENIE_ARTIFACT_BACKEND", "local")
    monkeypatch.setenv("MLR_GENIE_SESSION_SECRET", "shared-api-test-signing-secret-at-least-32-bytes")
    monkeypatch.setenv("MLR_GENIE_BOOTSTRAP_SECRET", "shared-api-test-bootstrap-secret")
    monkeypatch.setenv("MLR_GENIE_REQUIRE_DURABLE_STORAGE", "true")
    monkeypatch.delenv("MLR_GENIE_MAX_TOKEN_TTL_SECONDS", raising=False)
    monkeypatch.delenv("MLR_GENIE_ALLOWED_ORIGINS", raising=False)
    monkeypatch.delenv("K_SERVICE", raising=False)
    monkeypatch.delenv("MLR_REQUIRE_LIVE_GEMINI", raising=False)
    monkeypatch.delenv("GOOGLE_GENAI_LIVE_VERTEX", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    import backend.main as main
    from backend.artifact_storage import GcsArtifactStorage
    from backend.postgres_storage import PostgresJobStore
    from test_artifact_storage import FakeGcsClient

    fake = FakeGcsClient()
    tenant = "shared-api-" + uuid4().hex
    other_tenant = tenant + "-other"
    actor = "individual-" + uuid4().hex[:12]
    monkeypatch.setattr(main, "router", GeminiModelRouter(GcpAgentConfig(gemini_api_key=None, use_vertex_ai=False)))

    def activate():
        """Replace both adapters; every HTTP request receives a new empty workspace."""
        store = PostgresJobStore(dsn, iam_auth=False)
        artifacts = GcsArtifactStorage("shared-api-test-bucket", client=fake)
        store.before_save = artifacts.prepare_save
        paths = []
        original_directory = artifacts.directory

        def observe_directory(job_id, job_store, **kwargs):
            path = original_directory(job_id, job_store, **kwargs)
            paths.append(path)
            return path

        monkeypatch.setattr(artifacts, "directory", observe_directory)
        monkeypatch.setattr(main, "JOBS_DB", store)
        monkeypatch.setattr(main, "ARTIFACT_STORAGE", artifacts)
        return SimpleNamespace(store=store, artifacts=artifacts, paths=paths,
                               client=TestClient(main.app))

    def auth(owner=tenant, role="reviewer", person=actor):
        return {"Authorization": "Bearer " + main.issue_signed_session_token(owner, person, role=role)}

    reference = tmp_path / "source.pdf"
    create_born_digital_reference_pdf(reference, pages_spec=[{"passages": [CLAIM]}])
    replacement = tmp_path / "replacement.pdf"
    create_born_digital_reference_pdf(replacement, pages_spec=[{
        "header": "Replacement authorized source, revision 2", "passages": [CLAIM],
    }])

    fixture = SimpleNamespace(main=main, activate=activate, auth=auth, fake=fake,
                              tenant=tenant, other_tenant=other_tenant, actor=actor,
                              pdf=reference.read_bytes(), replacement_pdf=replacement.read_bytes(),
                              raw_store=lambda: PostgresJobStore(dsn, iam_auth=False))
    yield fixture
    cleanup = fixture.raw_store()
    for owner in (tenant, other_tenant):
        for job in cleanup.jobs_for_tenant(owner):
            del cleanup[job.job_id]


def create_roundtrip(api, instance):
    with instance.client as client:
        response = client.post("/integrations/commsrx/roundtrip", headers=api.auth(),
                               data={"payload": json.dumps(roundtrip_payload())},
                               files={"reference_pdf": ("source.pdf", api.pdf, "application/pdf")})
    assert response.status_code == 200, response.text
    result = response.json()
    assert set(result) == {"job_id", "state", "requires_human_resolution", "claims", "artifacts", "telemetry_receipts"}
    assert set(result["artifacts"]) == {"redline.docx", "clean_revised.docx", "audit_summary.pdf", "annotated_source.pdf"}
    return result


def download_published(api, instance, result):
    files = {}
    with instance.client as client:
        for name, descriptor in result["artifacts"].items():
            response = client.get(descriptor["download_url"], headers=api.auth())
            assert response.status_code == 200, response.text
            assert len(response.content) == descriptor["bytes"]
            assert hashlib.sha256(response.content).hexdigest() == descriptor["sha256"]
            files[name] = response.content
    assert instance.paths
    assert all(not path.exists() for path in instance.paths), "Request workspaces must be removed after responses finish."
    return files


def test_roundtrip_survives_fresh_instance_with_individual_authorship(shared_api):
    api = shared_api
    first = api.activate()
    result = create_roundtrip(api, first)
    saved = first.store[result["job_id"]]
    assert saved.comments[0].human_actor == api.actor
    assert DOWNSTREAM <= saved._artifact_manifest.keys()
    first_paths = set(first.paths)
    assert first_paths and all(not path.exists() for path in first_paths)

    fresh = api.activate()
    assert fresh.store is not first.store and fresh.artifacts is not first.artifacts
    downloads_before = api.fake.download_count
    files = download_published(api, fresh, result)
    assert api.fake.download_count > downloads_before
    assert first_paths.isdisjoint(fresh.paths)
    assert docx.Document(io.BytesIO(files["clean_revised.docx"])).paragraphs[0].text == CLAIM.replace("improved", "increased")
    with zipfile.ZipFile(io.BytesIO(files["redline.docx"])) as archive:
        xml = archive.read("word/document.xml").decode()
        assert f'w:author="{api.actor}"' in xml
        assert "<w:del " in xml and "<w:ins " in xml
    with pymupdf.open(stream=files["audit_summary.pdf"], filetype="pdf") as audit:
        assert api.actor in "".join(page.get_text() for page in audit)

    with fresh.client as client:
        for descriptor in result["artifacts"].values():
            assert client.get(descriptor["download_url"], headers=api.auth(api.other_tenant)).status_code == 403
        assert client.get(f"/jobs/{result['job_id']}", headers=api.auth(api.other_tenant)).status_code == 403
        assert client.get(f"/jobs/{result['job_id']}/artifacts/not-an-artifact.bin", headers=api.auth()).status_code == 404
        assert client.get(f"/jobs/absent-{uuid4().hex}", headers=api.auth()).status_code == 404


def test_health_and_job_reads_preserve_another_instances_committed_update(shared_api):
    api = shared_api
    instance = api.activate()
    result = create_roundtrip(api, instance)
    job_id = result["job_id"]
    stale = instance.store[job_id]
    peer = api.raw_store()
    edited = peer[job_id]
    edited.title = "Saved by the other instance"
    edited.metadata["peer_saved"] = True
    peer.save_job(edited)

    with instance.client as client:
        assert client.get("/health").status_code == 200
        response = client.get(f"/jobs/{job_id}", headers=api.auth())
        assert response.status_code == 200, response.text
        assert response.json()["job"]["title"] == edited.title
        assert client.get("/health").status_code == 200

    current = peer[job_id]
    assert stale.title != current.title
    assert current.metadata["peer_saved"] is True
    assert current._storage_revision == edited._storage_revision
    assert current._artifact_manifest == edited._artifact_manifest


def test_job_read_pins_artifacts_to_its_authorized_snapshot(shared_api, monkeypatch):
    """A peer commit between authorization and file loading cannot mix revisions."""
    from backend.artifact_storage import GcsArtifactStorage, artifact_workspace

    api = shared_api
    instance = api.activate()
    result = create_roundtrip(api, instance)
    job_id = result["job_id"]
    original = instance.store[job_id]
    original_get = instance.store.get_for_tenant
    changed_document = docx_bytes(CLAIM + " Peer revision published during the read.")
    changed_hash = hashlib.sha256(changed_document).hexdigest()
    committed = []

    def authorize_then_publish_peer_revision(requested_id, tenant):
        snapshot = original_get(requested_id, tenant)
        if requested_id == job_id and not committed:
            # This separate backend and workspace represent another instance
            # committing immediately after the GET loaded its authorized record.
            peer = api.raw_store()
            artifacts = GcsArtifactStorage("shared-api-test-bucket", client=api.fake)
            peer.before_save = artifacts.prepare_save
            updated = peer[job_id]
            with artifact_workspace():
                out = artifacts.directory(job_id, peer, job=updated)
                (out / "01_working_draft_v1.docx").write_bytes(changed_document)
                for filename in DOWNSTREAM:
                    (out / filename).unlink(missing_ok=True)
                updated.title = "Peer revision committed during GET"
                updated.metadata.pop("commsrx_roundtrip", None)
                updated.metadata.pop("last_roundtrip_artifacts", None)
                peer.save_job(updated)
            committed.append(updated)
        return snapshot

    monkeypatch.setattr(instance.store, "get_for_tenant", authorize_then_publish_peer_revision)
    with instance.client as client:
        response = client.get(f"/jobs/{job_id}", headers=api.auth())
        assert response.status_code == 200, response.text
        old_view = response.json()
        assert committed
        assert old_view["job"]["title"] == original.title
        assert old_view["artifacts"]["01_working_draft_v1.docx"]["sha256"] == original._artifact_manifest["01_working_draft_v1.docx"]["sha256"]
        assert old_view["artifacts"]["04_revised_clean_v2.docx"]["exists"] is True

        subsequent = client.get(f"/jobs/{job_id}", headers=api.auth())
        assert subsequent.status_code == 200, subsequent.text
        new_view = subsequent.json()
        assert new_view["job"]["title"] == committed[0].title
        assert new_view["artifacts"]["01_working_draft_v1.docx"]["sha256"] == changed_hash
        assert new_view["artifacts"]["04_revised_clean_v2.docx"]["exists"] is False


@pytest.mark.parametrize("upload", ["working-docx", "reference-pdf"])
def test_fresh_upload_invalidates_published_downstream_files_on_all_instances(shared_api, upload):
    api = shared_api
    result = create_roundtrip(api, api.activate())
    fresh = api.activate()
    data, filename, media = (
        (docx_bytes(CLAIM + " Revised working copy."), "revision.docx", DOCX_MEDIA)
        if upload == "working-docx" else (api.replacement_pdf, "replacement.pdf", "application/pdf")
    )
    with fresh.client as client:
        response = client.post(f"/jobs/{result['job_id']}/upload-{upload}", headers=api.auth(),
                               files={"file": (filename, data, media)})
    assert response.status_code == 200, response.text

    restarted = api.activate()
    current = restarted.store[result["job_id"]]
    assert not DOWNSTREAM.intersection(current._artifact_manifest)
    assert "commsrx_roundtrip" not in current.metadata
    assert "last_roundtrip_artifacts" not in current.metadata
    with restarted.client as client:
        for name in ("redline.docx", "clean_revised.docx", "audit_summary.pdf"):
            assert client.get(result["artifacts"][name]["download_url"], headers=api.auth()).status_code == 404
        source = "01_working_draft_v1.docx" if upload == "working-docx" else "02_reference_source.pdf"
        download = client.get(f"/jobs/{result['job_id']}/artifacts/{source}", headers=api.auth())
        assert download.status_code == 200
        assert download.content == data


@pytest.mark.parametrize("failure", ["upload", "database_cas"])
def test_failed_publication_preserves_the_last_committed_artifact_set(shared_api, monkeypatch, failure):
    api = shared_api
    result = create_roundtrip(api, api.activate())
    instance = api.activate()
    job_id = result["job_id"]
    before = instance.store[job_id]
    manifest = deepcopy(before._artifact_manifest)
    uploads_before = api.fake.upload_count

    if failure == "upload":
        # Publish one changed immutable object, then fail the next upload.
        api.fake.fail_upload_after = uploads_before + 1
    else:
        original_prepare = instance.artifacts.prepare_save

        def prepare_then_lose_database_race(job):
            original_prepare(job)
            peer = api.raw_store()
            winning = peer[job_id]
            winning.metadata["peer_won_cas"] = True
            peer.save_job(winning)

        monkeypatch.setattr(instance.store, "before_save", prepare_then_lose_database_race)

    with instance.client as client:
        response = client.post(f"/jobs/{job_id}/upload-reference-pdf", headers=api.auth(),
                               files={"file": ("replacement.pdf", api.replacement_pdf, "application/pdf")})
    assert response.status_code == (503 if failure == "upload" else 409), response.text
    assert api.fake.upload_count > uploads_before
    api.fake.fail_upload_after = None

    fresh = api.activate()
    persisted = fresh.store[job_id]
    assert persisted._artifact_manifest == manifest
    assert persisted.metadata["uploaded_reference_pdf"] == before.metadata["uploaded_reference_pdf"]
    if failure == "database_cas":
        assert persisted.metadata["peer_won_cas"] is True
    else:
        assert persisted._storage_revision == before._storage_revision
    download_published(api, fresh, result)


def test_opaque_job_ids_and_tenants_cannot_alias_published_files(shared_api):
    api = shared_api
    instance = api.activate()
    stem = "shared_" + uuid4().hex
    cases = [(stem + "_item", api.tenant, CLAIM),
             (stem + ":item", api.other_tenant, "Signal stability improved by 30% in adults.")]
    expected = {}
    with instance.client as client:
        for job_id, tenant, text in cases:
            job = JobRecord(job_id=job_id, tenant_id=tenant, title="Isolated artifact",
                            sector=SectorType.CORPORATE, document_type="Brief",
                            destination_platform="package_mode", deadline_iso="2026-12-01")
            intake = client.post("/jobs/intake", headers=api.auth(tenant),
                                 json={"job": job.model_dump(mode="json"), "initial_body": text, "actor": api.actor})
            assert intake.status_code == 200, intake.text
            original = docx_bytes(text)
            upload = client.post(f"/jobs/{job_id}/upload-working-docx", headers=api.auth(tenant),
                                 files={"file": ("working.docx", original, DOCX_MEDIA)})
            assert upload.status_code == 200, upload.text
            expected[job_id] = original

    fresh = api.activate()
    descriptors = [fresh.store[job_id]._artifact_manifest["01_working_draft_v1.docx"] for job_id, _, _ in cases]
    assert descriptors[0]["object"] != descriptors[1]["object"]
    with fresh.client as client:
        for job_id, tenant, _ in cases:
            url = f"/jobs/{job_id}/artifacts/01_working_draft_v1.docx"
            own = client.get(url, headers=api.auth(tenant))
            assert own.status_code == 200 and own.content == expected[job_id]
            other = api.other_tenant if tenant == api.tenant else api.tenant
            assert client.get(url, headers=api.auth(other)).status_code == 403


def test_required_durable_storage_refuses_local_backends_with_valid_hosted_secrets(monkeypatch, tmp_path):
    """Security configuration alone cannot authorize an ephemeral production backend."""
    monkeypatch.setenv("MLR_GENIE_ENV", "hosted")
    monkeypatch.setenv("MLR_GENIE_SESSION_SECRET", "durable-guard-test-signing-secret-over-32-bytes")
    monkeypatch.setenv("MLR_GENIE_BOOTSTRAP_SECRET", "durable-guard-test-bootstrap-secret")
    monkeypatch.setenv("MLR_GENIE_MAX_TOKEN_TTL_SECONDS", "3600")
    monkeypatch.setenv("MLR_GENIE_ALLOWED_ORIGINS", "")
    monkeypatch.setenv("MLR_GENIE_REQUIRE_DURABLE_STORAGE", "true")
    monkeypatch.setenv("MLR_GENIE_ARTIFACT_BACKEND", "local")
    monkeypatch.setenv("DATABASE_URL", "sqlite:///" + str(tmp_path / "guard.sqlite3"))

    import backend.main as main
    from backend.security import SecurityConfigurationError
    from backend.storage import DurableJobStore

    monkeypatch.setattr(main, "JOBS_DB", DurableJobStore(tmp_path / "guard.sqlite3"))
    monkeypatch.setattr(main, "ARTIFACT_STORAGE", None)
    with pytest.raises(SecurityConfigurationError, match="Production requires PostgreSQL jobs and GCS artifacts"):
        with TestClient(main.app):
            pass
