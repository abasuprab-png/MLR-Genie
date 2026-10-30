from __future__ import annotations

import pytest

from backend.domain.schemas import (
    AssetRevisionSnapshot,
    CommentCategory,
    ConnectorActionType,
    JobRecord,
    ProvenanceEventType,
    ReviewComment,
    SectorType,
    WordingSpanEdit,
)
from backend.tools.document_tools import (
    apply_accepted_changes,
    apply_wording_span,
    extract_human_resolved_facts,
    ingest_working_document,
    keep_source,
)
from backend.tools.platform_adapters import CapabilityCheckedPlatformAdapter
from backend.tools.review_tools import screen_pharmacovigilance_cv08, validate_strategy
from evals.acceptance_suite import run_all_12_acceptance_cases
from evals.pharma_comms_scenario_battery import run_pharma_comms_scenario_battery


@pytest.fixture
def local_prototype_security(monkeypatch):
    """Legacy HTTP scenarios intentionally exercise the prototype header workflow."""
    monkeypatch.setenv("MLR_GENIE_ENV", "local")
    monkeypatch.delenv("K_SERVICE", raising=False)


def test_all_12_section_11_acceptance_cases_pass() -> None:
    results = run_all_12_acceptance_cases()
    assert len(results) == 12
    for case_name, outcome in results.items():
        assert outcome["passed"] is True, f"Acceptance case failed: {case_name} -> {outcome}"


def test_all_15_pharma_comms_real_world_scenarios_pass() -> None:
    battery = run_pharma_comms_scenario_battery()
    assert len(battery) == 15
    for scen_key, outcome in battery.items():
        assert outcome["passed"] is True, f"Scenario failed: {scen_key} -> {outcome}"



def test_zero_silent_mutation_atomic_all_or_nothing_rollback() -> None:
    """
    Mandatory Rule 3:
    - If any span of a multi-part change fails to place uniquely, roll back the entire
      comment's edit (`all-or-nothing`) and mark it `unplaced`.
    - Never silently strip periods or rewrite punctuation inside `apply_wording_span`.
    """
    job = JobRecord(
        job_id="job_atomic_test",
        tenant_id="tenant_a",
        title="Atomic Revision Test",
        sector=SectorType.PHARMA_APPROVED,
        document_type="Press Release",
        destination_platform="veeva_promomats",
        deadline_iso="2026-11-05T17:00:00Z",
    )
    original_text = (
        "Sentence one with exact period. "
        "Repeated phrase appears here. "
        "Repeated phrase appears here."
    )
    ingest_working_document(job=job, body_text=original_text, actor="sarah.chen@example.com")

    # Comment 1 has two spans: first span is unique, second span is ambiguous (appears twice)
    multi_span_comment = ReviewComment(
        comment_id="cmt_multi_rollback",
        reviewer_name="Dr. David Miller",
        reviewer_role="Regulatory",
        timestamp="2026-10-28T10:00:00Z",
        anchored_passage=original_text,
        comment_text="Update both sentences.",
        category=CommentCategory.FACTUAL_CORRECTION,
        human_decision="accepted",
        proposed_spans=[
            WordingSpanEdit(
                target_text="Sentence one with exact period.",
                replacement_text="Sentence one updated with exact period.",
            ),
            WordingSpanEdit(
                target_text="Repeated phrase appears here.",
                replacement_text="Ambiguous replacement.",
            ),
        ],
    )

    outcome = apply_accepted_changes(
        job=job,
        comments=[multi_span_comment],
        human_actor="sarah.chen@example.com",
    )
    assert outcome["placed_comment_Ids"] == []
    assert outcome["unplaced_comment_Ids"] == ["cmt_multi_rollback"]
    assert multi_span_comment.placement_status == "unplaced"
    # Entire comment edit rolled back — Sentence one was NOT mutated!
    assert job.revisions[-1].body == original_text


def test_apply_wording_span_preserves_exact_punctuation() -> None:
    source = "Aurenix reduced risk by 38% (HR 0.62; 95% CI, 0.49 to 0.78)."
    span = WordingSpanEdit(
        target_text="reduced risk by 38%",
        replacement_text="reduced the risk of progression by 38%",
    )
    updated, ok = apply_wording_span(source, span)
    assert ok is True
    assert updated == "Aurenix reduced the risk of progression by 38% (HR 0.62; 95% CI, 0.49 to 0.78)."

    # Keep source leaves document untouched
    job = JobRecord(
        job_id="job_keep",
        tenant_id="tenant_a",
        title="Keep Source Test",
        sector=SectorType.PHARMA_APPROVED,
        document_type="Press Release",
        destination_platform="veeva_promomats",
        deadline_iso="2026-11-05T17:00:00Z",
    )
    ingest_working_document(job=job, body_text=source, actor="sarah.chen@example.com")
    cmt = ReviewComment(
        comment_id="cmt_keep",
        reviewer_name="Reviewer",
        reviewer_role="Medical",
        timestamp="2026-10-28T10:00:00Z",
        anchored_passage=source,
        comment_text="Optional editorial tweak.",
        category=CommentCategory.EDITORIAL_REQUEST,
        proposed_spans=[span],
    )
    kept_body = keep_source(job=job, comment=cmt, human_actor="sarah.chen@example.com")
    assert kept_body == source
    assert cmt.placement_status == "kept_source"


def test_invariant_1_strict_human_vs_model_provenance() -> None:
    """
    Mandatory Rule 2 (Invariant 1):
    Never classify model-generated text (`AssetDrafted`) as customer-supplied fact (`AssetEditRecorded`).
    """
    initial_body = "Aurenix studied 410 patients in Phase 3."
    revisions = [
        AssetRevisionSnapshot(
            revision_id="rev_1",
            version_number=1,
            body=initial_body,
            sha256_hash="h1",
            event_type=ProvenanceEventType.ASSET_EDIT_RECORDED,
            actor="sarah.chen@example.com",
        ),
        AssetRevisionSnapshot(
            revision_id="rev_2_model",
            version_number=2,
            body="Aurenix studied 410 patients in Phase 3. Model redrafted unconfirmed sentence.",
            sha256_hash="h2",
            event_type=ProvenanceEventType.ASSET_DRAFTED,
            actor="gemini-3.8-flash",
        ),
        AssetRevisionSnapshot(
            revision_id="rev_3_human",
            version_number=3,
            body="Aurenix studied 410 patients in Phase 3. Median PFS was 14.2 months per CSR.",
            sha256_hash="h3",
            event_type=ProvenanceEventType.ASSET_EDIT_RECORDED,
            actor="sarah.chen@example.com",
        ),
    ]
    human_facts = extract_human_resolved_facts(initial_payload_body=initial_body, revisions=revisions)
    assert "Median PFS was 14.2 months per CSR." in human_facts
    assert "Model redrafted unconfirmed sentence." not in human_facts


def test_invariant_10_heuristics_advise_never_block() -> None:
    """
    Mandatory Rule 1 (Invariant 10):
    Heuristics in `validate_strategy` never throw or return `severity: "block"`; they return `severity: "note"`.
    """
    job = JobRecord(
        job_id="job_inv10",
        tenant_id="tenant_a",
        title="Heuristic Advisory Check",
        sector=SectorType.PHARMA_APPROVED,
        document_type="Social Post",
        destination_platform="veeva_promomats",
        deadline_iso="2026-11-05T17:00:00Z",
    )
    findings = validate_strategy(
        job=job,
        text_to_check="Our therapy offers best-in-class efficacy and guaranteed patient convenience.",
    )
    assert len(findings) == 2
    assert all(f.severity == "note" for f in findings)


def test_aprimo_dam_vs_productivity_capability_separation_and_cv08_gate() -> None:
    job = JobRecord(
        job_id="job_aprimo_cv08",
        tenant_id="tenant_a",
        title="Aprimo DAM Boundary Check",
        sector=SectorType.PHARMA_APPROVED,
        document_type="Press Release",
        destination_platform="aprimo",
        deadline_iso="2026-11-05T17:00:00Z",
    )
    dam_adapter = CapabilityCheckedPlatformAdapter("aprimo_dam")
    rcpt = dam_adapter.execute_action(
        job=job,
        action=ConnectorActionType.REQUEST_REVIEW,
        actor_human="sarah.chen@example.com",
        payload={},
        idempotency_key="idem_dam_review",
    )
    assert rcpt.status == "unsupported_limitation"

    pv_res = screen_pharmacovigilance_cv08(
        job=job,
        intake_text="A 64-year-old patient experienced severe nausea and was hospitalized after taking dose 2.",
    )
    assert pv_res["cv08_flagged"] is True
    assert job.pv_icrs_flagged is True


@pytest.mark.usefixtures("local_prototype_security")
def test_phase2_full_stack_upload_docx_pdf_roundtrip_and_6_artifact_downloads(tmp_path) -> None:
    """
    Phase 2 verification:
    1. POST /jobs/intake with X-Tenant-Id and X-Actor-Id headers
    2. POST /jobs/{job_id}/upload-working-docx with a real .docx file
    3. POST /jobs/{job_id}/upload-reference-pdf with a real .pdf file and PyMuPDF coordinate extraction
    4. POST /jobs/{job_id}/comments with atomic accepted span resolution
    5. POST /jobs/{job_id}/execute-file-roundtrip generating all 6 artifacts and PyMuPDF preview metadata
    6. GET /jobs/{job_id}/artifacts/{filename} downloading all 6 artifacts with SHA-256 verification
    7. POST /jobs/{job_id}/submit/{platform_key} with X-Tenant-Id and X-Actor-Id headers
    """
    import hashlib
    from fastapi.testclient import TestClient
    import backend.main as main_mod
    from backend.document_engine import (
        create_born_digital_reference_pdf,
        create_working_draft_docx,
    )

    client = TestClient(main_mod.app)
    job_id = "job_phase2_e2e_roundtrip"
    if job_id in main_mod.JOBS_DB:
        del main_mod.JOBS_DB[job_id]

    headers = {
        "X-Tenant-Id": "tenant_aurenix",
        "X-Actor-Id": "sarah.chen@aurenix.example",
    }

    draft_paragraphs = [
        "In 2L+ EGFR-mutated NSCLC, Aurenix reduced progression-free survival risk by 38% (HR 0.62; 95% CI, 0.49 to 0.78; p<0.001).",
        "Important Safety Information: Aurenix is contraindicated in patients with severe hepatic impairment.",
    ]

    intake_resp = client.post(
        "/jobs/intake",
        headers=headers,
        json={
            "job": JobRecord(
                job_id=job_id,
                tenant_id="tenant_aurenix",
                title="Phase 2 Full-Stack Roundtrip Verification",
                sector=SectorType.PHARMA_APPROVED,
                document_type="Press Release",
                destination_platform="veeva_promomats",
                deadline_iso="2026-11-15T17:00:00Z",
            ).model_dump(),
            "initial_body": "\n".join(draft_paragraphs),
            "actor": "sarah.chen@aurenix.example",
            "source_pack": [],
        },
    )
    assert intake_resp.status_code == 200

    # Create and upload a real .docx working draft
    sample_docx = tmp_path / "working_upload.docx"
    create_working_draft_docx(
        sample_docx,
        title="Phase 2 Uploaded Working Draft",
        paragraphs=draft_paragraphs,
    )
    with sample_docx.open("rb") as fh:
        up_docx_resp = client.post(
            f"/jobs/{job_id}/upload-working-docx",
            headers=headers,
            files={
                "file": (
                    "working_upload.docx",
                    fh,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
    assert up_docx_resp.status_code == 200
    up_docx_data = up_docx_resp.json()
    assert up_docx_data["artifact_filename"] == "01_working_draft_v1.docx"
    assert len(up_docx_data["parsed_docx"]["paragraphs"]) == 2

    # Create and upload a real born-digital .pdf reference pack
    sample_pdf = tmp_path / "reference_upload.pdf"
    create_born_digital_reference_pdf(
        sample_pdf,
        pages_spec=[
            {
                "header": "AURORA-301 Clinical Study Report — Page 1",
                "passages": draft_paragraphs,
            }
        ],
    )
    with sample_pdf.open("rb") as fh:
        up_pdf_resp = client.post(
            f"/jobs/{job_id}/upload-reference-pdf",
            headers=headers,
            files={"file": ("reference_upload.pdf", fh, "application/pdf")},
            data={"source_id": "SRC-UPLOADED-CSR", "version": "3.0"},
        )
    assert up_pdf_resp.status_code == 200
    up_pdf_data = up_pdf_resp.json()
    assert up_pdf_data["reference_pdf"]["page_count"] == 1
    assert up_pdf_data["annotated_pdf"]["annotation_count"] >= 1
    first_annot = up_pdf_data["annotated_pdf"]["annotations"][0]
    assert first_annot["anchor_valid"] is True
    assert 0 <= first_annot["x0"] < first_annot["x1"] <= 612.0
    assert 0 <= first_annot["y0"] < first_annot["y1"] <= 792.0

    # Reconcile and atomically apply an accepted reviewer comment via POST /jobs/{job_id}/comments
    cmt_resp = client.post(
        f"/jobs/{job_id}/comments",
        headers=headers,
        json=[
            ReviewComment(
                comment_id="cmt_phase2_1",
                reviewer_name="Dr. David Miller",
                reviewer_role="Regulatory",
                timestamp="2026-10-28T16:00:00Z",
                anchored_passage=draft_paragraphs[0],
                comment_text="Specify statistically significant risk reduction.",
                category=CommentCategory.FACTUAL_CORRECTION,
                human_decision="accepted",
                proposed_spans=[
                    WordingSpanEdit(
                        target_text="reduced progression-free survival risk by 38%",
                        replacement_text="achieved a statistically significant 38% reduction in progression-free survival risk",
                    )
                ],
            ).model_dump()
        ],
    )
    assert cmt_resp.status_code == 200
    cmt_data = cmt_resp.json()
    assert "cmt_phase2_1" in cmt_data["reconciliation"]["applied_changes"]["placed_comment_Ids"]

    # Execute full file roundtrip via POST /jobs/{job_id}/execute-file-roundtrip
    rt_resp = client.post(f"/jobs/{job_id}/execute-file-roundtrip", headers=headers)
    assert rt_resp.status_code == 200
    rt_data = rt_resp.json()
    assert rt_data["redline_docx"]["tracked_replacements"] == 1
    assert rt_data["annotated_pdf"]["annotation_count"] >= 1
    assert len(rt_data["annotated_pdf"]["pages"]) == 1

    expected_6_files = [
        "01_working_draft_v1.docx",
        "02_reference_source.pdf",
        "03_annotated_reference.pdf",
        "04_revised_clean_v2.docx",
        "05_tracked_changes_redline_v1_to_v2.docx",
        "06_resubmission_package.zip",
    ]
    for fname in expected_6_files:
        assert fname in rt_data["artifacts"]
        assert rt_data["artifacts"][fname]["exists"] is True
        assert rt_data["artifacts"][fname]["bytes"] > 0

        # Download via header auth
        dl_resp = client.get(f"/jobs/{job_id}/artifacts/{fname}", headers=headers)
        assert dl_resp.status_code == 200
        assert len(dl_resp.content) == rt_data["artifacts"][fname]["bytes"]
        assert hashlib.sha256(dl_resp.content).hexdigest() == rt_data["artifacts"][fname]["sha256"]

        # Download via one-click browser query auth (?tenant_id=...&actor_id=...)
        dl_query_resp = client.get(
            f"/jobs/{job_id}/artifacts/{fname}",
            params={"tenant_id": "tenant_aurenix", "actor_id": "sarah.chen@aurenix.example"},
        )
        assert dl_query_resp.status_code == 200
        assert dl_query_resp.content == dl_resp.content

    # Submit to platform via POST /jobs/{job_id}/submit/{platform_key}
    sub_resp = client.post(
        f"/jobs/{job_id}/submit/veeva_promomats",
        headers=headers,
        json={"actor_human": "sarah.chen@aurenix.example", "idempotency_key": "idem_phase2_v1"},
    )
    assert sub_resp.status_code == 200
    assert sub_resp.json()["receipt"]["status"] in ("succeeded", "reconciled_no_duplicate")
    assert sub_resp.json()["job_state"] == "uploaded"


@pytest.mark.usefixtures("local_prototype_security")
def test_phase2_upload_and_download_edge_cases_and_tenant_isolation(tmp_path) -> None:
    """
    Tests edge cases on Phase 2 endpoints:
    - Unauthenticated upload & download rejected (401)
    - Cross-tenant upload & download rejected (403)
    - Corrupt / invalid .docx and .pdf rejected (400)
    - Unknown artifact filename rejected (404)
    - CV-08 Pharmacovigilance stop on uploaded .docx places job on HOLD
    """
    from fastapi.testclient import TestClient
    import backend.main as main_mod
    from backend.document_engine import create_working_draft_docx

    client = TestClient(main_mod.app)
    job_id = "job_phase2_security_edge"
    if job_id in main_mod.JOBS_DB:
        del main_mod.JOBS_DB[job_id]

    headers_a = {"X-Tenant-Id": "tenant_a", "X-Actor-Id": "alice@tenant_a.com"}
    headers_b = {"X-Tenant-Id": "tenant_b", "X-Actor-Id": "mallory@tenant_b.com"}

    client.post(
        "/jobs/intake",
        headers=headers_a,
        json={
            "job": JobRecord(
                job_id=job_id,
                tenant_id="tenant_a",
                title="Edge Case Security Job",
                sector=SectorType.PHARMA_APPROVED,
                document_type="Press Release",
                destination_platform="veeva_promomats",
                deadline_iso="2026-11-15T17:00:00Z",
            ).model_dump(),
            "initial_body": "Aurenix clinical update.",
            "actor": "alice@tenant_a.com",
            "source_pack": [],
        },
    )

    # 1. Unauthenticated upload -> 401
    unauth_up = client.post(
        f"/jobs/{job_id}/upload-working-docx",
        files={"file": ("bad.docx", b"corrupt", "application/octet-stream")},
    )
    assert unauth_up.status_code == 401

    # 2. Cross-tenant upload -> 403
    cross_up = client.post(
        f"/jobs/{job_id}/upload-working-docx",
        headers=headers_b,
        files={"file": ("bad.docx", b"corrupt", "application/octet-stream")},
    )
    assert cross_up.status_code == 403

    # 3. Corrupt .docx -> 400
    corrupt_docx = client.post(
        f"/jobs/{job_id}/upload-working-docx",
        headers=headers_a,
        files={"file": ("corrupt.docx", b"not a valid zip or docx", "application/octet-stream")},
    )
    assert corrupt_docx.status_code == 400

    # 4. Corrupt .pdf -> 400
    corrupt_pdf = client.post(
        f"/jobs/{job_id}/upload-reference-pdf",
        headers=headers_a,
        files={"file": ("corrupt.pdf", b"not a valid pdf", "application/pdf")},
    )
    assert corrupt_pdf.status_code == 400

    # 5. Unauthenticated & cross-tenant download -> 401 / 403
    unauth_dl = client.get(f"/jobs/{job_id}/artifacts/01_working_draft_v1.docx")
    assert unauth_dl.status_code == 401
    cross_dl = client.get(
        f"/jobs/{job_id}/artifacts/01_working_draft_v1.docx",
        params={"tenant_id": "tenant_b", "actor_id": "mallory@tenant_b.com"},
    )
    assert cross_dl.status_code == 403

    # 6. Disallowed filename -> 404
    bad_fname_dl = client.get(
        f"/jobs/{job_id}/artifacts/arbitrary_file.txt",
        headers=headers_a,
    )
    assert bad_fname_dl.status_code == 404

    # 7. CV-08 Pharmacovigilance stop on uploaded .docx containing an adverse event report
    pv_docx_path = tmp_path / "pv_ae_draft.docx"
    create_working_draft_docx(
        pv_docx_path,
        title="Draft With Adverse Event",
        paragraphs=[
            "A 64-year-old patient experienced severe nausea and was hospitalized after taking Aurenix."
        ],
    )
    with pv_docx_path.open("rb") as fh:
        pv_up = client.post(
            f"/jobs/{job_id}/upload-working-docx",
            headers=headers_a,
            files={"file": ("pv_ae_draft.docx", fh, "application/octet-stream")},
        )
    assert pv_up.status_code == 200
    assert pv_up.json()["pv_screening"]["cv08_flagged"] is True
    assert pv_up.json()["job"]["state"] == "hold"


@pytest.mark.usefixtures("local_prototype_security")
def test_phase2_table_cell_claims_multi_span_redline_and_reupload_invalidation(tmp_path) -> None:
    """
    Regression test for:
    1. Claims inside Word table cells (.docx tables) parsed into body_text and rewritten in
       both 04_revised_clean_v2.docx and 05_tracked_changes_redline_v1_to_v2.docx.
    2. Multiple accepted redline spans inside the SAME paragraph / table cell preserving
       exact document order and all <w:del>/<w:ins> elements.
    3. Frontend optimistic placement_status="placed" still triggering backend apply_accepted_changes.
    4. Re-uploading a new .docx invalidating stale 04-06 artifacts.
    """
    from docx import Document
    from fastapi.testclient import TestClient
    import backend.main as main_mod

    client = TestClient(main_mod.app)
    job_id = "job_phase2_table_multispan"
    if job_id in main_mod.JOBS_DB:
        del main_mod.JOBS_DB[job_id]

    headers = {"X-Tenant-Id": "tenant_aurenix", "X-Actor-Id": "sarah.chen@aurenix.example"}

    client.post(
        "/jobs/intake",
        headers=headers,
        json={
            "job": JobRecord(
                job_id=job_id,
                tenant_id="tenant_aurenix",
                title="Table & Multi-Span Redline Job",
                sector=SectorType.PHARMA_APPROVED,
                document_type="Detail Aid",
                destination_platform="veeva_promomats",
                deadline_iso="2026-11-20T17:00:00Z",
            ).model_dump(),
            "initial_body": "Initial placeholder.",
            "actor": "sarah.chen@aurenix.example",
            "source_pack": [],
        },
    )

    # Build a .docx with two target spans in the same paragraph AND a claim inside a Word table cell
    docx_path = tmp_path / "table_and_multispan.docx"
    doc = Document()
    doc.add_heading("Aurenix Table & Multi-Span Draft", level=1)
    doc.add_paragraph(
        "Aurenix provides unmatched efficacy and zero discontinuation in adults."
    )
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = "Key Claim"
    tbl.cell(0, 1).text = "Table Claim: 100% cure rate in all cohorts."
    doc.save(str(docx_path))

    with docx_path.open("rb") as fh:
        up_resp = client.post(
            f"/jobs/{job_id}/upload-working-docx",
            headers=headers,
            files={"file": ("table_and_multispan.docx", fh, "application/octet-stream")},
        )
    assert up_resp.status_code == 200
    up_data = up_resp.json()
    assert len(up_data["parsed_docx"]["table_cells"]) == 2
    assert "Table Claim: 100% cure rate in all cohorts." in up_data["parsed_docx"]["body_text"]

    # Reconcile comments where frontend sends optimistic placement_status="placed"
    # and two spans target the same paragraph while one targets the Word table cell
    cmt_resp = client.post(
        f"/jobs/{job_id}/comments",
        headers=headers,
        json=[
            ReviewComment(
                comment_id="cmt_multispan_1",
                reviewer_name="Dr. Lin",
                reviewer_role="Medical",
                timestamp="2026-10-28T16:00:00Z",
                anchored_passage="Aurenix provides unmatched efficacy and zero discontinuation in adults.",
                comment_text="Replace unmatched efficacy and zero discontinuation, and fix table claim.",
                category=CommentCategory.FACTUAL_CORRECTION,
                human_decision="accepted",
                placement_status="placed",
                proposed_spans=[
                    WordingSpanEdit(
                        target_text="unmatched efficacy",
                        replacement_text="42% symptom reduction (p<0.001)",
                    ),
                    WordingSpanEdit(
                        target_text="zero discontinuation",
                        replacement_text="a 4.1% discontinuation rate",
                    ),
                    WordingSpanEdit(
                        target_text="100% cure rate in all cohorts.",
                        replacement_text="42% response rate in Phase 3 (N=612).",
                    ),
                ],
            ).model_dump()
        ],
    )
    assert cmt_resp.status_code == 200
    cmt_data = cmt_resp.json()
    assert "cmt_multispan_1" in cmt_data["reconciliation"]["applied_changes"]["placed_comment_Ids"]

    # Execute roundtrip and inspect both clean and redline .docx
    rt_resp = client.post(f"/jobs/{job_id}/execute-file-roundtrip", headers=headers)
    assert rt_resp.status_code == 200
    rt_data = rt_resp.json()
    assert rt_data["clean_revised_docx"]["applied_span_count"] == 3
    assert rt_data["redline_docx"]["tracked_replacements"] == 3

    # Verify redline .docx XML preserves both spans in the same paragraph in exact order
    import zipfile
    import io

    redline_dl = client.get(
        f"/jobs/{job_id}/artifacts/05_tracked_changes_redline_v1_to_v2.docx",
        headers=headers,
    )
    assert redline_dl.status_code == 200
    with zipfile.ZipFile(io.BytesIO(redline_dl.content), "r") as zf:
        doc_xml = zf.read("word/document.xml").decode("utf-8")
    assert doc_xml.count("<w:del ") == 3
    assert doc_xml.count("<w:ins ") == 3
    idx_span1_del = doc_xml.find("unmatched efficacy")
    idx_span1_ins = doc_xml.find("42% symptom reduction (p&lt;0.001)")
    idx_and = doc_xml.find(" and ", idx_span1_ins)
    idx_span2_del = doc_xml.find("zero discontinuation", idx_and)
    idx_span2_ins = doc_xml.find("a 4.1% discontinuation rate", idx_span2_del)
    assert 0 < idx_span1_del < idx_span1_ins < idx_and < idx_span2_del < idx_span2_ins

    # Verify clean .docx updated both the paragraph and the table cell
    clean_dl = client.get(f"/jobs/{job_id}/artifacts/04_revised_clean_v2.docx", headers=headers)
    assert clean_dl.status_code == 200
    clean_path = tmp_path / "downloaded_clean.docx"
    clean_path.write_bytes(clean_dl.content)
    clean_doc = Document(str(clean_path))
    para_texts = [p.text for p in clean_doc.paragraphs]
    assert any(
        "Aurenix provides 42% symptom reduction (p<0.001) and a 4.1% discontinuation rate in adults."
        in t
        for t in para_texts
    )
    assert "42% response rate in Phase 3 (N=612)." in clean_doc.tables[0].cell(0, 1).text

    # Re-upload a fresh .docx and confirm stale 04-06 artifacts are invalidated
    with docx_path.open("rb") as fh:
        reup_resp = client.post(
            f"/jobs/{job_id}/upload-working-docx",
            headers=headers,
            files={"file": ("table_and_multispan.docx", fh, "application/octet-stream")},
        )
    assert reup_resp.status_code == 200
    reup_artifacts = reup_resp.json()["artifacts"]
    assert reup_artifacts["01_working_draft_v1.docx"]["exists"] is True
    assert reup_artifacts["04_revised_clean_v2.docx"]["exists"] is False
    assert reup_artifacts["05_tracked_changes_redline_v1_to_v2.docx"]["exists"] is False
    assert reup_artifacts["06_resubmission_package.zip"]["exists"] is False


@pytest.mark.usefixtures("local_prototype_security")
def test_phase2_multipage_pdf_anchoring_and_duplicate_quote_disambiguation(tmp_path) -> None:
    """
    Regression test for:
    1. Multi-page PDF reference pack where multiple pages share a general keyword ("Aurenix")
       and where safety evidence resides on Page 2 while efficacy evidence resides on Page 1.
    2. Verifies locate_and_anchor_evidence anchors the efficacy claim to Page 1 and the safety
       claim to Page 2 with non-zero page_width and page_height.
    3. Verifies 02_reference_source.pdf and 03_annotated_reference.pdf have distinct SHA-256 hashes.
    """
    import pymupdf
    from fastapi.testclient import TestClient
    import backend.main as main_mod

    client = TestClient(main_mod.app)
    job_id = "job_phase2_multipage_pdf"
    if job_id in main_mod.JOBS_DB:
        del main_mod.JOBS_DB[job_id]

    headers = {"X-Tenant-Id": "tenant_aurenix", "X-Actor-Id": "sarah.chen@aurenix.example"}

    client.post(
        "/jobs/intake",
        headers=headers,
        json={
            "job": JobRecord(
                job_id=job_id,
                tenant_id="tenant_aurenix",
                title="Multi-Page PDF Anchoring Job",
                sector=SectorType.PHARMA_APPROVED,
                document_type="Press Release",
                destination_platform="veeva_promomats",
                deadline_iso="2026-11-22T17:00:00Z",
            ).model_dump(),
            "initial_body": (
                "Aurenix achieved 42% sustained reduction in primary endpoint symptom score at Week 12. "
                "See Boxed Warning for hepatotoxicity monitoring and contraindications in severe hepatic impairment."
            ),
            "actor": "sarah.chen@aurenix.example",
            "source_pack": [],
        },
    )

    # Create a 2-page PDF where Page 1 has Efficacy and Page 2 has Safety/Boxed Warning
    pdf_path = tmp_path / "multipage_uspi.pdf"
    doc = pymupdf.open()
    p1 = doc.new_page(width=612, height=792)
    p1.insert_text(
        pymupdf.Point(54, 100),
        "Aurenix achieved 42% sustained reduction in primary endpoint symptom score at Week 12.",
        fontsize=10,
    )
    p2 = doc.new_page(width=612, height=792)
    p2.insert_text(
        pymupdf.Point(54, 140),
        "Aurenix Boxed Warning for hepatotoxicity monitoring and contraindications in severe hepatic impairment.",
        fontsize=10,
    )
    doc.save(str(pdf_path))
    doc.close()

    with pdf_path.open("rb") as fh:
        up_pdf = client.post(
            f"/jobs/{job_id}/upload-reference-pdf",
            headers=headers,
            data={"source_id": "src_multipage_uspi", "version": "4.0"},
            files={"file": ("multipage_uspi.pdf", fh, "application/pdf")},
        )
    assert up_pdf.status_code == 200
    pdf_data = up_pdf.json()
    assert pdf_data["reference_pdf"]["page_count"] == 2
    assert pdf_data["reference_pdf"]["file_sha256"] != pdf_data["annotated_pdf"]["sha256"]

    claims = pdf_data["evidence_anchoring"]
    anchored_pages = {
        anchor["page_number"]
        for claim in claims
        for anchor in claim.get("anchors", [])
    }
    assert 1 in anchored_pages
    assert 2 in anchored_pages

    for claim in claims:
        for anchor in claim.get("anchors", []):
            assert anchor["page_width"] == 612.0
            assert anchor["page_height"] == 792.0


def test_gemini_dual_mode_sdk_transport_interactions_vs_generate_content() -> None:
    """
    P0 Red-Team Remediation 1:
    Verifies dual-mode SDK transport in `GeminiModelRouter.create_interaction`:
    - Vertex AI / Enterprise Agent Platform (`self._client.interactions.create`) with `store`, `service_tier`,
      `response_format`, `safety_settings`, and `cached_content`.
    - Gemini Developer API (`self._client.models.generate_content`) with `config` containing
      `system_instruction`, `response_mime_type="application/json"`, `response_json_schema`,
      `safety_settings`, and `cached_content` when in `api_key` mode or when `interactions` is absent.
    """
    import json
    from types import SimpleNamespace
    from unittest.mock import MagicMock
    from backend.agent.models_router import GeminiModelRouter
    from backend.config import GcpAgentConfig
    from backend.domain.schemas import ClaimVerificationVerdict

    schema = ClaimVerificationVerdict.model_json_schema()

    # Mode A: Vertex AI / Enterprise Agent Platform (`interactions.create`)
    vertex_router = GeminiModelRouter(GcpAgentConfig(use_vertex_ai=True, zero_data_retention=True))
    vertex_mock = MagicMock()
    vertex_mock.interactions.create.return_value = SimpleNamespace(
        id="int_vertex_live_001",
        output_text=json.dumps(
            {
                "status": "supported",
                "classification": "supported",
                "population_match": True,
                "endpoint_match": True,
                "numeric_values_matched": ["38%"],
                "unsupported_numeric_values": [],
                "missing_qualifiers": [],
                "plain_english_rationale": "Exact match to CSR Table 2.",
            }
        ),
    )
    vertex_router._client = vertex_mock

    out_vertex, tel_vertex = vertex_router.create_interaction(
        task_type="claim_verification",
        prompt="Verify claim: Aurenix reduced PFS risk by 38%.",
        system_instruction="Verify clinical numbers and qualifiers.",
        response_schema=schema,
        previous_interaction_id="int_prev_000",
        cached_content="cachedContents/mlr_csr_301",
    )
    assert tel_vertex.used_live_sdk is True
    assert tel_vertex.transport_mode == "interactions"
    assert tel_vertex.interaction_id == "int_vertex_live_001"
    assert tel_vertex.cached_content == "cachedContents/mlr_csr_301"
    assert out_vertex["status"] == "supported"

    vertex_kwargs = vertex_mock.interactions.create.call_args.kwargs
    assert vertex_kwargs["model"] == "gemini-4-argon"
    assert vertex_kwargs["store"] is False
    assert vertex_kwargs["service_tier"] == "standard"
    assert vertex_kwargs["previous_interaction_id"] == "int_prev_000"
    assert vertex_kwargs["cached_content"] == "cachedContents/mlr_csr_301"
    assert vertex_kwargs["response_format"] == {
        "type": "text",
        "mime_type": "application/json",
        "json_schema": schema,
    }
    assert len(vertex_kwargs["safety_settings"]) == 4

    # Mode B: Gemini Developer API (`api_key` mode -> `models.generate_content`)
    dev_router = GeminiModelRouter(
        GcpAgentConfig(use_vertex_ai=False, gemini_api_key="test-api-key-no-network")
    )
    dev_mock = MagicMock()
    dev_mock.models.generate_content.return_value = SimpleNamespace(
        id="gen_dev_live_002",
        text=json.dumps(
            {
                "status": "contradicted",
                "classification": "contradicted",
                "population_match": True,
                "endpoint_match": False,
                "numeric_values_matched": [],
                "unsupported_numeric_values": ["48%"],
                "missing_qualifiers": [],
                "plain_english_rationale": "Copy cites 48% vs 38% in CSR.",
            }
        ),
    )
    dev_router._client = dev_mock

    out_dev, tel_dev = dev_router.create_interaction(
        task_type="claim_verification",
        prompt="Verify claim: Aurenix reduced PFS risk by 48%.",
        system_instruction="Verify clinical numbers and qualifiers.",
        response_schema=schema,
        cached_content="cachedContents/mlr_csr_301",
    )
    assert tel_dev.used_live_sdk is True
    assert tel_dev.transport_mode == "generate_content"
    assert out_dev["status"] == "contradicted"
    assert out_dev["unsupported_numeric_values"] == ["48%"]

    dev_kwargs = dev_mock.models.generate_content.call_args.kwargs
    assert dev_kwargs["model"] == "gemini-4-argon"
    assert dev_kwargs["contents"] == "Verify claim: Aurenix reduced PFS risk by 48%."
    gen_cfg = dev_kwargs["config"]
    assert gen_cfg["system_instruction"] == "Verify clinical numbers and qualifiers."
    assert gen_cfg["response_mime_type"] == "application/json"
    assert gen_cfg["response_json_schema"] == schema
    assert gen_cfg["cached_content"] == "cachedContents/mlr_csr_301"
    assert len(gen_cfg["safety_settings"]) == 4

    # Mode C: Client without `interactions` attribute automatically routes to `models.generate_content`
    no_interactions_router = GeminiModelRouter(GcpAgentConfig(use_vertex_ai=True))
    no_interactions_client = SimpleNamespace(
        models=SimpleNamespace(
            generate_content=lambda **kw: SimpleNamespace(
                id="gen_no_int_003",
                text=json.dumps({"status": "partially_supported", "missing_qualifiers": ["95% CI"]}),
            )
        )
    )
    no_interactions_router._client = no_interactions_client
    out_no_int, tel_no_int = no_interactions_router.create_interaction(
        task_type="claim_verification",
        prompt="Verify claim.",
        system_instruction="Verify.",
        response_schema=schema,
    )
    assert tel_no_int.transport_mode == "generate_content"
    assert out_no_int["status"] == "partially_supported"

    # Mode D: Expired `cached_content` TTL (`404 Not Found`) automatically retries once without `cached_content`
    ttl_calls: list[dict] = []

    def _generate_with_expired_cache(**kw):
        ttl_calls.append(kw)
        if kw.get("config", {}).get("cached_content"):
            raise RuntimeError("404 NOT_FOUND: CachedContent cachedContents/expired_ttl_999 expired")
        return SimpleNamespace(
            id="gen_retry_no_cache_004",
            text=json.dumps({"status": "supported", "classification": "supported"}),
        )

    retry_router = GeminiModelRouter(GcpAgentConfig(use_vertex_ai=False, gemini_api_key="test-key"))
    retry_router._client = SimpleNamespace(
        models=SimpleNamespace(generate_content=_generate_with_expired_cache)
    )
    out_retry, tel_retry = retry_router.create_interaction(
        task_type="claim_verification",
        prompt="Verify claim after cache expiry.",
        system_instruction="Verify.",
        response_schema=schema,
        cached_content="cachedContents/expired_ttl_999",
    )
    assert len(ttl_calls) == 2
    assert ttl_calls[0]["config"].get("cached_content") == "cachedContents/expired_ttl_999"
    assert "cached_content" not in ttl_calls[1]["config"]
    assert out_retry["status"] == "supported"
    assert tel_retry.used_live_sdk is True
    assert tel_retry.cached_content is None


def test_claim_verification_verdict_schema_and_locate_and_anchor_evidence() -> None:
    """
    P0 Red-Team Remediation 2:
    Verifies `ClaimVerificationVerdict` Pydantic schema (including strict `required` array),
    ensures `locate_and_anchor_evidence` passes `response_schema=ClaimVerificationVerdict.model_json_schema()`
    and `cached_content` without duplicating the JSON schema in the prompt text, validates model
    output through `ClaimVerificationVerdict.model_validate`, and denormalizes `box_2d` coordinates.
    """
    from typing import Any
    from backend.agent.models_router import GeminiModelRouter, ModelCallTelemetry
    from backend.domain.schemas import ClaimVerificationVerdict, SupportClassification
    from backend.tools.document_tools import locate_and_anchor_evidence

    schema = ClaimVerificationVerdict.model_json_schema()
    props = schema["properties"]
    required_fields = set(schema.get("required", []))
    for required_prop in (
        "status",
        "classification",
        "population_match",
        "endpoint_match",
        "numeric_values_matched",
        "unsupported_numeric_values",
        "missing_qualifiers",
        "plain_english_rationale",
    ):
        assert required_prop in props
        assert required_prop in required_fields

    verdict = ClaimVerificationVerdict(status="supported", numeric_values_matched=["38%"])
    assert verdict.classification == "supported"
    verdict_from_cls = ClaimVerificationVerdict(classification="contradicted")
    assert verdict_from_cls.status == "contradicted"

    captured_calls: list[dict[str, Any]] = []

    class InspectingRouter(GeminiModelRouter):
        def create_interaction(self, **kwargs: Any) -> tuple[Any, ModelCallTelemetry]:
            captured_calls.append(kwargs)
            return super().create_interaction(**kwargs)

    job = JobRecord(
        job_id="job_schema_check",
        tenant_id="tenant_aurenix",
        title="Schema Inspection Job",
        sector=SectorType.PHARMA_PREAPPROVAL,
        document_type="Press Release",
        destination_platform="veeva_promomats",
        deadline_iso="2026-11-25T17:00:00Z",
        metadata={"gemini_context_cache": {"name": "cachedContents/mlr_test_cache_99"}},
    )
    ingest_working_document(
        job=job,
        body_text="In 2L+ EGFR-mutated NSCLC, Aurenix reduced progression-free survival risk by 38% (95% CI, 0.49 to 0.78).",
        actor="sarah.chen@aurenix.example",
    )
    claims = locate_and_anchor_evidence(
        job=job,
        source_pack=[
            {
                "source_id": "SRC-CSR-301",
                "version": "2.0",
                "verbatim_text": "In 2L+ EGFR-mutated NSCLC, Aurenix reduced progression-free survival risk by 38% (95% CI, 0.49 to 0.78).",
                "keywords": ["aurenix", "38%"],
                "locked_numbers": ["38%", "0.49", "0.78"],
                "required_qualifiers": ["2L+ EGFR-mutated NSCLC", "95% CI"],
                "bbox": {
                    "page": 2,
                    "box_2d": [100, 100, 500, 500],
                    "page_width": 612.0,
                    "page_height": 792.0,
                },
            }
        ],
        router=InspectingRouter(),
    )
    assert len(captured_calls) == 1
    call = captured_calls[0]
    assert call["response_schema"] == schema
    assert call["cached_content"] == "cachedContents/mlr_test_cache_99"
    assert "properties" not in call["prompt"]
    assert "plain_english_rationale" not in call["prompt"]

    # Verify `box_2d` [100, 100, 500, 500] was denormalized into PDF point coordinates on the anchor
    anchor = claims[0].anchors[0]
    assert anchor.page_number == 2
    assert anchor.x0 == 61.2
    assert anchor.y0 == 79.2
    assert anchor.x1 == 306.0
    assert anchor.y1 == 396.0

    # Verify invalid status enum from model fails ClaimVerificationVerdict validation and holds claim UNRESOLVED
    class InvalidEnumStubRouter(GeminiModelRouter):
        def create_interaction(self, **kwargs: Any) -> tuple[Any, ModelCallTelemetry]:
            return {"status": "hallucinated_invalid_enum"}, ModelCallTelemetry(
                model="gemini-4-argon",
                task_type="claim_verification",
                interaction_id="int_invalid_enum",
            )

    claims_invalid = locate_and_anchor_evidence(
        job=job,
        source_pack=[
            {
                "source_id": "SRC-CSR-301",
                "version": "2.0",
                "verbatim_text": "In 2L+ EGFR-mutated NSCLC, Aurenix reduced progression-free survival risk by 38% (95% CI, 0.49 to 0.78).",
                "keywords": ["aurenix", "38%"],
                "locked_numbers": ["38%", "0.49", "0.78"],
            }
        ],
        router=InvalidEnumStubRouter(),
    )
    assert claims_invalid[0].support_status == SupportClassification.UNRESOLVED


def test_clinical_safety_settings_calibration() -> None:
    """
    P1 Red-Team Remediation 3:
    Verifies explicit clinical-safe `safety_settings` on `GeminiModelRouter` so oncology/toxicology
    terminology in Prescribing Information is not falsely blocked (`HARM_CATEGORY_DANGEROUS_CONTENT: BLOCK_ONLY_HIGH`)
    while harassment, hate speech, and sexually explicit content remain strictly blocked (`BLOCK_LOW_AND_ABOVE`).
    """
    from backend.agent.models_router import CLINICAL_SAFETY_SETTINGS, GeminiModelRouter

    router = GeminiModelRouter()
    settings = router.get_clinical_safety_settings()
    assert settings == CLINICAL_SAFETY_SETTINGS
    by_cat = {item["category"]: item["threshold"] for item in settings}
    assert by_cat == {
        "HARM_CATEGORY_DANGEROUS_CONTENT": "BLOCK_ONLY_HIGH",
        "HARM_CATEGORY_HARASSMENT": "BLOCK_LOW_AND_ABOVE",
        "HARM_CATEGORY_HATE_SPEECH": "BLOCK_LOW_AND_ABOVE",
        "HARM_CATEGORY_SEXUALLY_EXPLICIT": "BLOCK_LOW_AND_ABOVE",
    }


def test_matryoshka_representation_learning_embed_multimodal() -> None:
    """
    P1 Red-Team Remediation 4:
    Verifies `embed_multimodal` passes `config={"output_dimensionality": dimensions, "task_type": task_type}`
    to `self._client.models.embed_content` across `RETRIEVAL_DOCUMENT`, `RETRIEVAL_QUERY`, and `FACT_VERIFICATION`,
    normalizes the output vector, and falls back deterministically on provider error including for high
    Matryoshka dimensions (e.g. 256, 768).
    """
    import math
    from types import SimpleNamespace
    from unittest.mock import MagicMock
    from backend.agent.models_router import GeminiModelRouter

    router = GeminiModelRouter()
    mock_client = MagicMock()
    mock_client.models.embed_content.return_value = SimpleNamespace(
        embeddings=[SimpleNamespace(values=[3.0, 4.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0])]
    )
    router._client = mock_client

    vec_doc = router.embed_multimodal("Aurenix reduced risk by 38%.", dimensions=8)
    assert len(vec_doc) == 8
    assert vec_doc[0] == 0.6
    assert vec_doc[1] == 0.8
    assert mock_client.models.embed_content.call_args.kwargs["config"] == {
        "output_dimensionality": 8,
        "task_type": "RETRIEVAL_DOCUMENT",
    }

    router.embed_multimodal("PFS hazard ratio", dimensions=12, task_type="retrieval_query")
    assert mock_client.models.embed_content.call_args.kwargs["config"] == {
        "output_dimensionality": 12,
        "task_type": "RETRIEVAL_QUERY",
    }

    router.embed_multimodal("Verify 38% PFS claim", dimensions=16, task_type="FACT_VERIFICATION")
    assert mock_client.models.embed_content.call_args.kwargs["config"] == {
        "output_dimensionality": 16,
        "task_type": "FACT_VERIFICATION",
    }

    # Verify fallback when embed_content raises an exception (including 256-dim and 768-dim Matryoshka vectors)
    mock_client.models.embed_content.side_effect = RuntimeError("Embedding endpoint unavailable")
    for target_dims in (16, 256, 768):
        fallback_vec = router.embed_multimodal("Aurenix clinical efficacy endpoint", dimensions=target_dims)
        assert len(fallback_vec) == target_dims
        assert abs(math.sqrt(sum(v * v for v in fallback_vec)) - 1.0) < 1e-2


@pytest.mark.usefixtures("local_prototype_security")
def test_gemini_context_caching_and_pdf_upload_registration(tmp_path) -> None:
    """
    P2 Red-Team Remediation 5:
    Verifies `GeminiModelRouter.create_reference_context_cache` (calling `self._client.caches.create`
    with `CreateCachedContentConfig`-valid `contents` for both `source_pack` dicts and binary PDF bytes
    when live SDK is active, and deterministic cache registration when offline or below token threshold),
    and verifies `upload_reference_pdf_endpoint` registers `job.metadata["gemini_context_cache"]` and
    extracts `box_2d` `[ymin, xmin, ymax, xmax]` `[0..1000]` coordinates.
    """
    from types import SimpleNamespace
    from unittest.mock import MagicMock
    from fastapi.testclient import TestClient
    from google.genai import types as genai_types
    import backend.main as main_mod
    from backend.agent.models_router import GeminiModelRouter
    from backend.document_engine import (
        create_born_digital_reference_pdf,
        denormalize_box_2d_to_pdf_rect,
        normalize_pdf_rect_to_box_2d,
    )

    # 1. Unit test live SDK `caches.create` on GeminiModelRouter with source_pack dicts AND binary PDF bytes
    router = GeminiModelRouter()
    mock_client = MagicMock()
    mock_client.caches.create.return_value = SimpleNamespace(name="cachedContents/live_sdk_cache_777")
    router._client = mock_client

    cache_info = router.create_reference_context_cache(
        cache_key="tenant_a:job_1:SRC-CSR-301:2.0",
        pdf_or_text_content=[{"verbatim_text": "Aurenix reduced risk by 38%.", "source_id": "SRC-1"}],
        ttl="3600s",
    )
    assert cache_info["used_live_sdk"] is True
    assert cache_info["name"] == "cachedContents/live_sdk_cache_777"
    assert mock_client.caches.create.call_count == 1
    assert mock_client.caches.create.call_args.kwargs["model"] == "gemini-4-argon"
    sent_cache_cfg = mock_client.caches.create.call_args.kwargs["config"]
    assert sent_cache_cfg["ttl"] == "3600s"
    # Validate config directly against real google.genai.types.CreateCachedContentConfig
    validated_cfg = genai_types.CreateCachedContentConfig.model_validate(sent_cache_cfg)
    assert validated_cfg.ttl == "3600s"

    # Second call with identical key & content reuses cached entry without re-calling caches.create
    cached_again = router.create_or_get_reference_cache(
        cache_key="tenant_a:job_1:SRC-CSR-301:2.0",
        pdf_or_text_content=[{"verbatim_text": "Aurenix reduced risk by 38%.", "source_id": "SRC-1"}],
        ttl="3600s",
    )
    assert cached_again["name"] == "cachedContents/live_sdk_cache_777"
    assert mock_client.caches.create.call_count == 1

    # Also verify binary PDF bytes normalize into a valid CreateCachedContentConfig inline_data PartDict
    pdf_bytes_cache = router.create_reference_context_cache(
        cache_key="tenant_a:job_1:SRC-RAW-PDF:1.0",
        pdf_or_text_content=b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n",
        ttl="3600s",
    )
    assert pdf_bytes_cache["used_live_sdk"] is True
    pdf_sent_cfg = mock_client.caches.create.call_args.kwargs["config"]
    genai_types.CreateCachedContentConfig.model_validate(pdf_sent_cfg)

    # Also verify graceful fallback when caches.create rejects a small PDF below min token count
    mock_client.caches.create.side_effect = RuntimeError(
        "400 INVALID_ARGUMENT: Cached content is too small (< 1024 tokens)"
    )
    small_cache = router.create_reference_context_cache(
        cache_key="tenant_a:job_1:SRC-SMALL-PDF:1.0",
        pdf_or_text_content="Short reference snippet.",
        ttl="3600s",
    )
    assert small_cache["used_live_sdk"] is False
    assert small_cache["name"].startswith("cachedContents/mlr_")
    assert "too small" in small_cache["fallback_reason"]

    # 2. Verify box_2d [ymin, xmin, ymax, xmax] 0..1000 normalization helpers
    box_2d = normalize_pdf_rect_to_box_2d(
        x0=61.2, y0=79.2, x1=306.0, y1=396.0, page_width=612.0, page_height=792.0
    )
    assert box_2d == [100, 100, 500, 500]
    denorm = denormalize_box_2d_to_pdf_rect(box_2d, page_width=612.0, page_height=792.0)
    assert denorm == {"x0": 61.2, "y0": 79.2, "x1": 306.0, "y1": 396.0}

    # 3. Verify `upload_reference_pdf_endpoint` registers `job.metadata["gemini_context_cache"]`
    client = TestClient(main_mod.app)
    job_id = "job_context_cache_pdf"
    if job_id in main_mod.JOBS_DB:
        del main_mod.JOBS_DB[job_id]

    headers = {"X-Tenant-Id": "tenant_aurenix", "X-Actor-Id": "sarah.chen@aurenix.example"}
    client.post(
        "/jobs/intake",
        headers=headers,
        json={
            "job": JobRecord(
                job_id=job_id,
                tenant_id="tenant_aurenix",
                title="Context Cache PDF Upload Job",
                sector=SectorType.PHARMA_APPROVED,
                document_type="Press Release",
                destination_platform="veeva_promomats",
                deadline_iso="2026-11-28T17:00:00Z",
            ).model_dump(),
            "initial_body": "Aurenix reduced progression-free survival risk by 38%.",
            "actor": "sarah.chen@aurenix.example",
            "source_pack": [],
        },
    )

    sample_pdf = tmp_path / "cache_ref.pdf"
    create_born_digital_reference_pdf(
        sample_pdf,
        pages_spec=[
            {
                "header": "AURORA-301 Reference — Page 1",
                "passages": ["Aurenix reduced progression-free survival risk by 38%."],
            }
        ],
    )
    with sample_pdf.open("rb") as fh:
        up_resp = client.post(
            f"/jobs/{job_id}/upload-reference-pdf",
            headers=headers,
            files={"file": ("cache_ref.pdf", fh, "application/pdf")},
            data={"source_id": "SRC-CACHE-CSR", "version": "1.0"},
        )
    assert up_resp.status_code == 200
    data = up_resp.json()
    assert "gemini_context_cache" in data
    assert data["gemini_context_cache"]["name"].startswith("cachedContents/mlr_")
    assert data["job"]["metadata"]["gemini_context_cache"]["name"] == data["gemini_context_cache"]["name"]
    first_block = data["reference_pdf"]["pages"][0]["text_blocks"][0]
    assert len(first_block["box_2d"]) == 4
    assert all(0 <= coord <= 1000 for coord in first_block["box_2d"])


@pytest.mark.usefixtures("local_prototype_security")
def test_all_27_adversarial_probes_pass() -> None:
    """Verifies all 27 adversarial probes (including Probes 25 & 26) pass."""
    from evals.audit_adversarial import run_all_27_adversarial_probes

    results = run_all_27_adversarial_probes()
    assert len(results) == 27
    for probe_key, val in results.items():
        assert val["passed"] is True, f"Adversarial probe failed: {probe_key} -> {val}"


def test_phase3_pico_clinical_entailment_and_sanitized_telemetry_receipts() -> None:
    """
    Phase 3 Gap 1:
    Verifies structured PICO clinical entailment decomposition (`population`, `endpoint`,
    `comparator`, `direction`, `magnitude`, `timeframe`) on `ClaimVerificationVerdict` and
    `ClaimAssessment`, PICO comparator conflict detection, and sanitized `telemetry_receipts`
    in `job.metadata["telemetry_receipts"]` with zero secret leakage.
    """
    import json
    from backend.agent.models_router import GcpAgentConfig, GeminiModelRouter
    from backend.domain.schemas import ClaimVerificationVerdict, SupportClassification
    from backend.tools.document_tools import locate_and_anchor_evidence

    schema = ClaimVerificationVerdict.model_json_schema()
    props = schema["properties"]
    required_fields = set(schema.get("required", []))
    for pico_field in (
        "entailment_classification",
        "population",
        "endpoint",
        "comparator",
        "direction",
        "magnitude",
        "timeframe",
        "verbatim_quote",
        "page_number",
        "box_2d",
        "missing_qualifiers",
        "numerical_mismatches",
        "rationale",
    ):
        assert pico_field in props
        assert pico_field in required_fields

    # Verify bidirectional sync on ClaimVerificationVerdict
    v_partial = ClaimVerificationVerdict(
        entailment_classification="contradicted",
        population="2L+ EGFR-mutated NSCLC",
        endpoint="progression-free survival",
        comparator="placebo",
        direction="reduced",
        magnitude="38%",
        timeframe="Week 24",
        numerical_mismatches=["48%"],
        rationale="Mismatch in PFS reduction figure.",
    )
    assert v_partial.status == "contradicted"
    assert v_partial.classification == "contradicted"
    assert v_partial.unsupported_numeric_values == ["48%"]
    assert v_partial.plain_english_rationale == "Mismatch in PFS reduction figure."

    # Verify locate_and_anchor_evidence populates PICO dimensions, detects comparator conflict,
    # and records sanitized telemetry receipts without leaking API keys
    secret_api_key = "AIzaSySecretKeyShouldNeverLeak999999"
    router = GeminiModelRouter(
        GcpAgentConfig(use_vertex_ai=False, gemini_api_key=secret_api_key)
    )
    router._client = None
    job = JobRecord(
        job_id="job_phase3_pico",
        tenant_id="tenant_aurenix",
        title="PICO Entailment & Telemetry Job",
        sector=SectorType.PHARMA_PREAPPROVAL,
        document_type="Press Release",
        destination_platform="veeva_promomats",
        deadline_iso="2026-11-29T17:00:00Z",
        metadata={"gemini_context_cache": {"name": "cachedContents/mlr_pico_cache_01"}},
    )
    ingest_working_document(
        job=job,
        body_text=(
            "In adults with 2L+ EGFR-mutated NSCLC (N=412), Aurenix reduced progression-free survival risk "
            "by 38% (95% CI, 0.49 to 0.78) at 24 weeks versus placebo. "
            "In adults with 2L+ EGFR-mutated NSCLC, Aurenix reduced progression-free survival risk "
            "by 38% (95% CI, 0.49 to 0.78) at 24 weeks versus docetaxel."
        ),
        actor="sarah.chen@aurenix.example",
    )
    claims = locate_and_anchor_evidence(
        job=job,
        source_pack=[
            {
                "source_id": "SRC-CSR-301",
                "version": "2.0",
                "verbatim_text": (
                    "In adults with 2L+ EGFR-mutated NSCLC (N=412), Aurenix reduced progression-free survival risk "
                    "by 38% (95% CI, 0.49 to 0.78) at 24 weeks versus placebo."
                ),
                "keywords": ["aurenix", "progression-free", "38%"],
                "locked_numbers": ["38%", "0.49", "0.78", "24"],
                "required_qualifiers": ["2L+ EGFR-mutated NSCLC", "95% CI"],
                "bbox": {
                    "page": 3,
                    "box_2d": [120, 90, 220, 880],
                    "page_width": 612.0,
                    "page_height": 792.0,
                },
            }
        ],
        router=router,
    )
    assert len(claims) == 2

    # Claim 1 matches PICO dimensions completely
    c1 = claims[0]
    assert c1.support_status == SupportClassification.SUPPORTED
    assert c1.entailment_classification == "supported"
    assert "NSCLC" in c1.population
    assert "progression-free survival" in c1.endpoint.lower()
    assert "placebo" in c1.comparator.lower()
    assert c1.direction == "reduced"
    assert "38%" in c1.magnitude
    assert "24 weeks" in c1.timeframe.lower()

    # Claim 2 has a PICO comparator mismatch ("versus docetaxel" vs "versus placebo" in source)
    c2 = claims[1]
    assert c2.support_status == SupportClassification.CONTRADICTED
    assert c2.entailment_classification == "contradicted"
    assert "docetaxel" in c2.comparator.lower()
    assert "PICO comparator mismatch" in c2.plain_english_rationale

    # Verify sanitized telemetry receipts in job.metadata["telemetry_receipts"]
    receipts = job.metadata.get("telemetry_receipts")
    assert isinstance(receipts, list)
    assert len(receipts) == 2
    for rcpt in receipts:
        assert set(rcpt.keys()) >= {
            "effective_model",
            "transport",
            "service_tier",
            "zero_data_retention",
            "cached_content",
            "timestamp_iso",
        }
        assert rcpt["effective_model"] == "gemini-4-argon"
        assert rcpt["zero_data_retention"] is True
        assert rcpt["cached_content"] == "cachedContents/mlr_pico_cache_01"
        serialized_rcpt = json.dumps(rcpt)
        assert secret_api_key not in serialized_rcpt


def test_phase3_docx_run_rpr_formatting_and_header_footer_preservation(tmp_path) -> None:
    """
    Phase 3 Gap 2A & 2B:
    Verifies `parse_docx_working_document`, `write_clean_revised_docx`, and
    `write_tracked_changes_redline_docx`:
    1. Traverse section headers (`section.header`) and footers (`section.footer`) in addition
       to body paragraphs and table cells.
    2. Preserve character-level `<w:rPr>` formatting (bold, italic, underline, superscript `®`, `†`, `‡`)
       across unaffected runs and replaced `<w:del>` / `<w:ins>` spans rather than flattening the
       entire paragraph into a single unformatted run.
    """
    from docx import Document
    from backend.document_engine import (
        parse_docx_working_document,
        write_clean_revised_docx,
        write_tracked_changes_redline_docx,
    )
    from backend.domain.schemas import CommentCategory, ReviewComment, WordingSpanEdit

    v1_path = tmp_path / "formatted_v1.docx"
    doc = Document()

    # Header with formatted runs
    section = doc.sections[0]
    hdr_p = section.header.paragraphs[0]
    hdr_p.text = ""
    r_hdr_bold = hdr_p.add_run("AURENIX")
    r_hdr_bold.bold = True
    r_hdr_sup = hdr_p.add_run("®")
    r_hdr_sup.font.superscript = True
    hdr_p.add_run(" Launch Press Release — Draft Header v1")

    # Body paragraph with bold, superscript ®, italic p-value, underline qualifier, and footnote †
    body_p = doc.add_paragraph()
    r_brand = body_p.add_run("AURENIX")
    r_brand.bold = True
    r_reg = body_p.add_run("®")
    r_reg.font.superscript = True
    body_p.add_run(" demonstrated ")
    r_ul = body_p.add_run("best-in-class")
    r_ul.underline = True
    body_p.add_run(" progression-free survival (")
    r_ital = body_p.add_run("p < 0.0001")
    r_ital.italic = True
    body_p.add_run(")")
    r_dag = body_p.add_run("†")
    r_dag.font.superscript = True
    body_p.add_run(" in 2L+ NSCLC.")

    # Footer with footnote symbol ‡
    ftr_p = section.footer.paragraphs[0]
    ftr_p.text = ""
    r_ddag = ftr_p.add_run("‡")
    r_ddag.font.superscript = True
    ftr_p.add_run(" Safety Notice: Preliminary unverified footer text.")

    doc.save(str(v1_path))

    # 1. Verify parse_docx_working_document traverses header, body, and footer
    parsed = parse_docx_working_document(v1_path)
    assert len(parsed["headers"]) == 1
    assert len(parsed["paragraphs"]) == 1
    assert len(parsed["footers"]) == 1
    assert "AURENIX® Launch Press Release — Draft Header v1" in parsed["body_text"]
    assert "‡ Safety Notice: Preliminary unverified footer text." in parsed["body_text"]

    # 2. Apply edits across header, body, and footer
    comments = [
        ReviewComment(
            comment_id="CMT-HDR",
            reviewer_name="MLR Reviewer",
            reviewer_role="Regulatory",
            timestamp="2026-10-29T10:00:00Z",
            anchored_passage="Draft Header v1",
            comment_text="Update header status.",
            category=CommentCategory.EDITORIAL_REQUEST,
            human_decision="accepted",
            placement_status="placed",
            proposed_spans=[
                WordingSpanEdit(
                    target_text="Draft Header v1",
                    replacement_text="Final Approved Header v2",
                )
            ],
        ),
        ReviewComment(
            comment_id="CMT-BODY-RPR",
            reviewer_name="MLR Reviewer",
            reviewer_role="Regulatory",
            timestamp="2026-10-29T10:00:00Z",
            anchored_passage="best-in-class",
            comment_text="Replace promotional superiority term.",
            category=CommentCategory.FACTUAL_CORRECTION,
            human_decision="accepted",
            placement_status="placed",
            proposed_spans=[
                WordingSpanEdit(
                    target_text="best-in-class",
                    replacement_text="statistically significant",
                )
            ],
        ),
        ReviewComment(
            comment_id="CMT-FTR",
            reviewer_name="MLR Reviewer",
            reviewer_role="Medical",
            timestamp="2026-10-29T10:00:00Z",
            anchored_passage="Preliminary unverified footer text.",
            comment_text="Lock safety footer.",
            category=CommentCategory.FACTUAL_CORRECTION,
            human_decision="accepted",
            placement_status="placed",
            proposed_spans=[
                WordingSpanEdit(
                    target_text="Preliminary unverified footer text.",
                    replacement_text="See Full Prescribing Information for Boxed Warning.",
                )
            ],
        ),
    ]

    clean_v2_path = tmp_path / "formatted_clean_v2.docx"
    redline_v2_path = tmp_path / "formatted_redline_v2.docx"

    clean_res = write_clean_revised_docx(
        source_docx_path=v1_path,
        output_docx_path=clean_v2_path,
        accepted_comments=comments,
    )
    assert clean_res["applied_span_count"] == 3

    # Inspect clean v2 DOCX runs to verify <w:rPr> formatting survived on unaffected runs AND replaced span
    doc_clean = Document(str(clean_v2_path))
    clean_body_p = doc_clean.paragraphs[0]
    assert (
        clean_body_p.text
        == "AURENIX® demonstrated statistically significant progression-free survival (p < 0.0001)† in 2L+ NSCLC."
    )
    bold_runs = [r.text for r in clean_body_p.runs if r.bold]
    sup_runs = [r.text for r in clean_body_p.runs if r.font.superscript]
    italic_runs = [r.text for r in clean_body_p.runs if r.italic]
    underline_runs = [r.text for r in clean_body_p.runs if r.underline]
    assert "AURENIX" in bold_runs
    assert "®" in sup_runs
    assert "†" in sup_runs
    assert "p < 0.0001" in italic_runs
    assert "statistically significant" in underline_runs

    # Verify header and footer in clean v2 DOCX
    clean_hdr_p = doc_clean.sections[0].header.paragraphs[0]
    assert "Final Approved Header v2" in clean_hdr_p.text
    assert any(r.font.superscript and r.text == "®" for r in clean_hdr_p.runs)
    clean_ftr_p = doc_clean.sections[0].footer.paragraphs[0]
    assert "See Full Prescribing Information for Boxed Warning." in clean_ftr_p.text
    assert any(r.font.superscript and r.text == "‡" for r in clean_ftr_p.runs)

    # 3. Verify tracked-changes redline DOCX preserves <w:rPr> alongside <w:del> and <w:ins>
    redline_res = write_tracked_changes_redline_docx(
        source_docx_path=v1_path,
        output_redline_docx_path=redline_v2_path,
        accepted_comments=comments,
    )
    assert redline_res["tracked_replacements"] == 3

    doc_redline = Document(str(redline_v2_path))
    redline_body_xml = doc_redline.paragraphs[0]._p.xml
    assert "w:del" in redline_body_xml
    assert "w:ins" in redline_body_xml
    assert "best-in-class" in redline_body_xml
    assert "statistically significant" in redline_body_xml
    assert 'w:val="superscript"' in redline_body_xml
    assert "w:b" in redline_body_xml
    assert "w:i" in redline_body_xml


def test_phase3_two_column_pdf_reading_order_and_scanned_page_ocr_fallback(tmp_path) -> None:
    """
    Phase 3 Gap 2C:
    Verifies `extract_pdf_source_pack`, `inspect_pdf_document_and_extract_source_pack`, and
    `locate_supporting_quote_in_pdf`:
    1. Sort two-column clinical reprint pages in column-first reading order (left column top-to-bottom,
       then right column top-to-bottom) rather than interleaving across columns.
    2. Detect scanned/textless pages (`ocr_fallback_required=True`) without crashing.
    """
    from backend.document_engine import (
        create_born_digital_reference_pdf,
        extract_pdf_source_pack,
        inspect_pdf_document_and_extract_source_pack,
        locate_supporting_quote_in_pdf,
    )

    reprint_pdf = tmp_path / "two_column_clinical_reprint.pdf"
    create_born_digital_reference_pdf(
        reprint_pdf,
        pages_spec=[
            {
                "header": "NEJM Clinical Reprint — AURORA-301 Two-Column Layout (Page 1)",
                "left_column": [
                    "LEFT-TOP: Baseline demographics were balanced across arms (N=412) in 2L+ NSCLC.",
                    "LEFT-BOTTOM: Median follow-up duration was 18.4 months across all study centers.",
                ],
                "right_column": [
                    "RIGHT-TOP: Primary Endpoint: Aurenix reduced progression-free survival risk by 38% (95% CI, 0.49 to 0.78).",
                    "RIGHT-BOTTOM: Secondary Endpoint: Overall response rate was 56% versus 29% in control.",
                ],
            },
            {
                "scanned_image_only": True,
            },
        ],
    )

    # 1. Verify inspect_pdf_document_and_extract_source_pack
    inspected = inspect_pdf_document_and_extract_source_pack(
        reprint_pdf,
        source_id="SRC-NEJM-REPRINT",
        version="1.0",
    )
    assert inspected["page_count"] == 2
    assert inspected["ocr_fallback_required"] is True
    assert inspected["scanned_page_numbers"] == [2]

    page1 = inspected["pages"][0]
    assert page1["two_column_layout_detected"] is True
    assert page1["ocr_fallback_required"] is False
    block_texts = [b["text"] for b in page1["text_blocks"]]
    # Find indices of the 4 column blocks and verify left column blocks precede right column blocks
    idx_lt = next(i for i, t in enumerate(block_texts) if "LEFT-TOP:" in t)
    idx_lb = next(i for i, t in enumerate(block_texts) if "LEFT-BOTTOM:" in t)
    idx_rt = next(i for i, t in enumerate(block_texts) if "RIGHT-TOP:" in t)
    idx_rb = next(i for i, t in enumerate(block_texts) if "RIGHT-BOTTOM:" in t)
    assert idx_lt < idx_lb < idx_rt < idx_rb

    page2 = inspected["pages"][1]
    assert page2["has_extractable_text"] is False
    assert page2["ocr_fallback_required"] is True

    # 2. Verify extract_pdf_source_pack includes column metadata and scanned-page fallback entry
    pack = extract_pdf_source_pack(reprint_pdf, source_id="SRC-NEJM-REPRINT", version="1.0")
    assert any(entry.get("column") == "left" for entry in pack)
    assert any(entry.get("column") == "right" for entry in pack)
    scanned_entries = [entry for entry in pack if entry["page_number"] == 2]
    assert len(scanned_entries) == 1
    assert scanned_entries[0]["ocr_fallback_required"] is True

    # 3. Verify locate_supporting_quote_in_pdf locates a right-column quote with right-column coordinates
    found_right = locate_supporting_quote_in_pdf(
        reprint_pdf,
        "Aurenix reduced progression-free survival risk by 38%",
        preferred_page=1,
    )
    assert found_right is not None
    assert found_right["found"] is True
    assert found_right["page_number"] == 1
    assert found_right["column"] == "right"
    assert found_right["x0"] >= 310.0
    assert found_right["ocr_fallback_required"] is False

    # 4. Verify locate_supporting_quote_in_pdf returns ocr_fallback_required on textless scanned page
    not_found_on_scanned = locate_supporting_quote_in_pdf(
        reprint_pdf,
        "Nonexistent quote on scanned appendix page",
        preferred_page=2,
    )
    assert not_found_on_scanned is not None
    assert not_found_on_scanned["found"] is False
    assert not_found_on_scanned["ocr_fallback_required"] is True
    assert 2 in not_found_on_scanned["scanned_page_numbers"]


@pytest.mark.usefixtures("local_prototype_security")
def test_phase3_rd_handoff_pack_importer_tool_and_endpoint() -> None:
    """
    Phase 3 Gap 3:
    Verifies `import_rd_handoff_pack` tool, `MLRGenieAgent` registration, and
    `POST /jobs/{job_id}/import-rd-handoff` endpoint:
    - Locks verified Phase 2 epidemiology & trial design statistics into `job.source_pack`
    - Generates pre-anchored "About the Disease & Epidemiology Boilerplate" (`job.rd_handoff_boilerplate`)
    - Appends draft revision with `ProvenanceEventType.ASSET_DRAFTED` (respecting Invariant 1)
    - Seeds verified `StatementReuseRecord` entries with `ProvenanceEventType.HUMAN_CONFIRMED`
    - Enforces tenant & actor authentication on the FastAPI endpoint
    """
    from fastapi.testclient import TestClient
    import backend.main as main_mod
    from backend.agent.mlr_genie_agent import MLRGenieAgent
    from backend.domain.schemas import ProvenanceEventType
    from backend.tools.document_tools import extract_human_resolved_facts

    agent = MLRGenieAgent()
    assert "import_rd_handoff_pack" in agent._tools

    client = TestClient(main_mod.app)
    job_id = "job_phase3_rd_handoff"
    if job_id in main_mod.JOBS_DB:
        del main_mod.JOBS_DB[job_id]

    headers = {
        "X-Tenant-Id": "tenant_aurenix",
        "X-Actor-Id": "vp.comms@aurenix.example",
    }
    intake_resp = client.post(
        "/jobs/intake",
        headers=headers,
        json={
            "job": JobRecord(
                job_id=job_id,
                tenant_id="tenant_aurenix",
                title="Phase 3 Launch Press Release with R&D Handoff",
                sector=SectorType.PHARMA_APPROVED,
                document_type="Press Release",
                destination_platform="veeva_promomats",
                deadline_iso="2026-12-01T17:00:00Z",
                jurisdiction="US",
                audience="HCP",
                channel="Press Release",
                product_version="Aurenix-v1.0",
            ).model_dump(),
            "initial_body": "AURENIX® Phase 3 Headline Announcement.",
            "actor": "vp.comms@aurenix.example",
            "source_pack": [],
        },
    )
    assert intake_resp.status_code == 200

    # Cross-tenant call must be rejected with 403
    wrong_tenant_resp = client.post(
        f"/jobs/{job_id}/import-rd-handoff",
        headers={"X-Tenant-Id": "tenant_other", "X-Actor-Id": "attacker@other.example"},
        json={"disease_name": "EGFR+ NSCLC", "human_confirmed": True},
    )
    assert wrong_tenant_resp.status_code == 403

    # Authorized call imports Phase 2 R&D Handoff Pack
    handoff_resp = client.post(
        f"/jobs/{job_id}/import-rd-handoff",
        headers=headers,
        json={
            "disease_name": "Moderate-to-Severe EGFR+ NSCLC",
            "trial_design_summary": "Phase 2 randomized, double-blind, placebo-controlled trial (N=280).",
            "human_confirmed": True,
        },
    )
    assert handoff_resp.status_code == 200
    payload = handoff_resp.json()
    assert "About the Disease & Epidemiology Boilerplate" in payload["rd_handoff_boilerplate"]
    assert payload["rd_handoff"]["imported_sources_count"] >= 2
    assert len(payload["reuse_library_seeded"]) >= 2

    job_after = main_mod.JOBS_DB[job_id]
    assert len(job_after.source_pack) >= 2
    assert "About the Disease & Epidemiology Boilerplate" in job_after.rd_handoff_boilerplate
    assert len(payload["seeded_claims"]) >= 2
    assert all(len(c["anchors"]) >= 1 for c in payload["seeded_claims"])
    assert len(job_after.claims) >= 2
    assert sum(1 for c in job_after.claims if len(c.anchors) >= 1) >= 2
    assert len(job_after.reuse_library) >= 2
    for rec in job_after.reuse_library:
        assert rec.provenance == ProvenanceEventType.HUMAN_CONFIRMED
        assert rec.confirmed_by == "vp.comms@aurenix.example"

    # Invariant 1 check: generated boilerplate revision is ASSET_DRAFTED, not ASSET_EDIT_RECORDED
    human_facts = extract_human_resolved_facts(
        initial_payload_body=job_after.revisions[0].body,
        revisions=job_after.revisions,
    )
    assert human_facts == []

    # Verify re-running locate_and_anchor_evidence after R&D handoff does NOT mark the boilerplate
    # section header as UNRESOLVED and automatically uses job.source_pack when source_pack is omitted
    from backend.domain.schemas import SupportClassification
    from backend.tools.document_tools import locate_and_anchor_evidence
    from backend.tools.review_tools import query_statement_reuse_library

    reverified_claims = locate_and_anchor_evidence(job=job_after)
    assert not any(
        "About the Disease & Epidemiology Boilerplate" in c.sentence_text
        for c in reverified_claims
    )
    rd_reverified = [
        c for c in reverified_claims if "Non-small cell lung cancer" in c.sentence_text or "AURORA-202" in c.sentence_text
    ]
    assert len(rd_reverified) == 2
    assert all(
        c.support_status in (SupportClassification.SUPPORTED, SupportClassification.PARTIALLY_SUPPORTED)
        and len(c.anchors) >= 1
        for c in rd_reverified
    )

    # Verify query_statement_reuse_library matches the job's explicit audience ("HCP") & product_version ("Aurenix-v1.0")
    reuse_results = query_statement_reuse_library(
        library=job_after.reuse_library,
        requesting_tenant_id="tenant_aurenix",
        query_text=job_after.reuse_library[0].exact_wording,
        target_jurisdiction="US",
        target_audience="HCP",
        target_channel="Press Release",
        target_product_version="Aurenix-v1.0",
    )
    assert len(reuse_results) >= 1
    assert reuse_results[0]["can_reuse_directly"] is True
    assert reuse_results[0]["reasons_requiring_rereview"] == []


def test_phase3_adversarial_edge_cases_pico_docx_and_cross_column_pdf(tmp_path) -> None:
    """
    Skeptical Reviewer regression tests for Phase 3 edge cases:
    1. Source pack entries with explicit PICO keys (`endpoint`, `comparator`, `timeframe`) must NOT
       mask claim drift (e.g., claiming 'overall survival' against a PFS source, or '52 weeks'
       against a '24 weeks' source).
    2. `.docx` `<w:rPr>` replacement on a span ending with a superscript footnote (`†`) or containing
       an interior bold brand token (`AURENIX`) must NOT smear superscript or bold across normal words.
    3. `locate_supporting_quote_in_pdf` must resolve quotes that wrap from the bottom of the left
       column (`LEFT-BOTTOM`) to the top of the right column (`RIGHT-TOP`) with `spans_columns=True`,
       and must prioritize direct block matches over short substring header blocks.
    """
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    from backend.document_engine import (
        create_born_digital_reference_pdf,
        locate_supporting_quote_in_pdf,
        write_clean_revised_docx,
    )
    from backend.domain.schemas import (
        CommentCategory,
        ReviewComment,
        SupportClassification,
        WordingSpanEdit,
    )
    from backend.tools.document_tools import locate_and_anchor_evidence

    # 1. PICO endpoint & timeframe conflict with explicit source PICO metadata
    pico_job = JobRecord(
        job_id="job_pico_adversarial",
        tenant_id="tenant_aurenix",
        title="PICO Endpoint & Timeframe Drift Check",
        sector=SectorType.PHARMA_APPROVED,
        document_type="Press Release",
        destination_platform="veeva_promomats",
        deadline_iso="2026-12-01T17:00:00Z",
    )
    ingest_working_document(
        job=pico_job,
        body_text=(
            "In 2L+ EGFR-mutated NSCLC, Aurenix improved overall survival by 38% (95% CI, 0.49 to 0.78) at 24 weeks. "
            "In 2L+ EGFR-mutated NSCLC, Aurenix improved progression-free survival by 38% (95% CI, 0.49 to 0.78) at 52 weeks. "
            "Important Safety Information: Contraindicated in pregnancy."
        ),
        actor="reviewer@aurenix.example",
    )
    pico_claims = locate_and_anchor_evidence(
        job=pico_job,
        source_pack=[
            {
                "source_id": "SRC-CSR-PICO-EXPLICIT",
                "version": "2.0",
                "verbatim_text": (
                    "In 2L+ EGFR-mutated NSCLC, Aurenix improved progression-free survival by 38% "
                    "(95% CI, 0.49 to 0.78) at 24 weeks versus placebo."
                ),
                "keywords": ["2l+ egfr-mutated nsclc", "aurenix", "38%", "0.49", "0.78"],
                "locked_numbers": ["38%", "0.49", "0.78", "24"],
                "required_qualifiers": ["2L+ EGFR-mutated NSCLC", "95% CI"],
                "population": "2L+ EGFR-mutated NSCLC",
                "endpoint": "progression-free survival",
                "comparator": "placebo",
                "timeframe": "24 weeks",
            }
        ],
    )
    os_claim = next(c for c in pico_claims if "overall survival" in c.sentence_text)
    assert os_claim.endpoint == "overall survival"
    assert os_claim.support_status == SupportClassification.CONTRADICTED
    assert "endpoint mismatch" in os_claim.plain_english_diagnosis.lower()

    week_claim = next(c for c in pico_claims if "52 weeks" in c.sentence_text)
    assert week_claim.support_status == SupportClassification.CONTRADICTED
    assert "52 weeks" in week_claim.numerical_mismatches

    # 2. `.docx` `<w:rPr>` non-smearing when span ends with superscript `†` or contains interior bold `AURENIX`
    src_docx = tmp_path / "rpr_no_smear_source.docx"
    clean_docx = tmp_path / "rpr_no_smear_clean.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("In Phase 3, ")
    p.add_run("preliminary ")
    r_brand = p.add_run("AURENIX")
    r_brand.bold = True
    r_reg = p.add_run("®")
    r_reg._r.get_or_add_rPr().append(parse_xml(f'<w:vertAlign {nsdecls("w")} w:val="superscript"/>'))
    p.add_run(" data showed efficacy")
    r_dag = p.add_run("†")
    r_dag._r.get_or_add_rPr().append(parse_xml(f'<w:vertAlign {nsdecls("w")} w:val="superscript"/>'))
    p.add_run(" in adults.")
    doc.save(str(src_docx))

    comment = ReviewComment(
        comment_id="cmt_no_smear_1",
        reviewer_name="Dr. Smear Check",
        reviewer_role="Medical",
        timestamp="2026-10-29T10:00:00Z",
        anchored_passage="preliminary AURENIX® data showed efficacy†",
        comment_text="Update preliminary wording.",
        category=CommentCategory.FACTUAL_CORRECTION,
        human_decision="accepted",
        placement_status="placed",
        proposed_spans=[
            WordingSpanEdit(
                target_text="preliminary AURENIX® data showed efficacy†",
                replacement_text="confirmed AURENIX® results demonstrated response†",
            )
        ],
    )
    write_clean_revised_docx(
        source_docx_path=src_docx,
        output_docx_path=clean_docx,
        accepted_comments=[comment],
    )
    reloaded = Document(str(clean_docx))
    para_After = reloaded.paragraphs[0]
    assert para_After.text == "In Phase 3, confirmed AURENIX® results demonstrated response† in adults."

    # Verify normal words ("confirmed", "results") are NOT superscript and NOT bold,
    # while "AURENIX" remains bold and "®"/"†" remain superscript
    for run in para_After.runs:
        rpr = run._r.find(qn("w:rPr"))
        is_sup = (
            rpr is not None
            and rpr.find(qn("w:vertAlign")) is not None
            and rpr.find(qn("w:vertAlign")).get(qn("w:val")) == "superscript"
        )
        is_bold = rpr is not None and rpr.find(qn("w:b")) is not None
        if "confirmed" in run.text or "results" in run.text or "demonstrated" in run.text:
            assert not is_sup, f"Normal word run '{run.text}' was smeared with superscript!"
            assert not is_bold, f"Normal word run '{run.text}' was smeared with bold!"
        if "AURENIX" in run.text:
            assert is_bold, "Interior brand token 'AURENIX' lost its bold formatting!"
        if "®" in run.text or "†" in run.text:
            assert is_sup, f"Superscript glyph '{run.text}' lost its superscript formatting!"

    # 3. Cross-column wrapped quote & short header substring priority in locate_supporting_quote_in_pdf
    wrap_pdf = tmp_path / "cross_column_wrap.pdf"
    create_born_digital_reference_pdf(
        wrap_pdf,
        pages_spec=[
            {
                "header": "2L+ EGFR-mutated NSCLC",
                "left_column": [
                    "LEFT-TOP: Study overview and patient eligibility criteria.",
                    "LEFT-BOTTOM: In 2L+ EGFR-mutated NSCLC patients treated with Aurenix,",
                ],
                "right_column": [
                    "RIGHT-TOP: median progression-free survival was 14.2 months versus 8.1 months.",
                    "RIGHT-BOTTOM: Safety profile was consistent with prior studies.",
                ],
            }
        ],
    )
    # Cross-column quote spanning LEFT-BOTTOM -> RIGHT-TOP (which also contains the short header substring)
    wrapped_match = locate_supporting_quote_in_pdf(
        wrap_pdf,
        "In 2L+ EGFR-mutated NSCLC patients treated with Aurenix, median progression-free survival was 14.2 months",
        preferred_page=1,
    )
    assert wrapped_match is not None
    assert wrapped_match["found"] is True
    assert wrapped_match.get("spans_columns") is True
    assert wrapped_match.get("continuation_bbox") is not None
    assert wrapped_match["continuation_bbox"]["column"] == "right"


@pytest.mark.usefixtures("local_prototype_security")
def test_phase3_independent_review_probes_and_auth_hardening(tmp_path) -> None:
    """
    Phase 3 Independent Review Remediation Regression Suite (`probe.py` & `extra.py`):
    1. Unseen contradiction probes (`population`, `direction`, `time_unit`, `number_role`)
       classify as `SupportClassification.CONTRADICTED` and absent PICO fields return `"unspecified"`.
    2. `extract_pdf_source_pack` populates `required_qualifiers` from clinical qualifier patterns
       when `passage_rules` is omitted.
    3. Post-edit claim staleness & automatic re-verification: `apply_accepted_changes` editing
       `'improved'` -> `'worsened'` increments `current_version` to 2, binds claims to v2 (`stale=False`),
       and re-evaluates the mutated claim as `CONTRADICTED`.
    4. Authentication & principal verification (`_require_authenticated_tenant` + HMAC-SHA256 bearer tokens):
       rejects single-header impersonation (`X-Actor-Id` alone, `X-Tenant-Id` alone), arbitrary
       `Authorization: nonsense` / `Bearer nonsense`, tampered tokens, and expired tokens with HTTP 401,
       rejects cross-tenant valid tokens with HTTP 403, and accepts valid signed tokens.
    5. Telemetry receipts honestly report `verification_mode` and `zdr_attestation`.
    """
    from fastapi import HTTPException
    from fastapi.testclient import TestClient
    from starlette.requests import Request
    import backend.main as main_mod
    from backend.document_engine import (
        create_born_digital_reference_pdf,
        extract_pdf_source_pack,
    )
    from backend.domain.schemas import (
        CommentCategory,
        ReviewComment,
        SupportClassification,
        WordingSpanEdit,
    )
    from backend.tools.document_tools import (
        apply_accepted_changes,
        ingest_working_document,
        locate_and_anchor_evidence,
    )

    # 1. Exact unseen contradiction probes from docs/independent_review_phase3_2026_10_29/probe.py
    def _run_claim_probe(sent: str, src: str):
        j = JobRecord(
            job_id="p",
            tenant_id="t",
            title="t",
            sector=SectorType.MEDTECH,
            document_type="Brief",
            destination_platform="veeva_promomats",
            deadline_iso="2026-12-01T00:00:00Z",
        )
        ingest_working_document(job=j, body_text=sent, actor="a")
        claim = locate_and_anchor_evidence(
            job=j,
            source_pack=[{"source_id": "S1", "version": "1", "verbatim_text": src}],
        )[0]
        return claim

    c_pop = _run_claim_probe(
        "Aurenix improved overall survival in children.",
        "Aurenix improved overall survival in adults.",
    )
    assert c_pop.support_status == SupportClassification.CONTRADICTED
    assert c_pop.population == "children"
    assert c_pop.timeframe == "unspecified"

    c_dir = _run_claim_probe(
        "Aurenix decreased overall survival by 38%.",
        "Aurenix increased overall survival by 38%.",
    )
    assert c_dir.support_status == SupportClassification.CONTRADICTED
    assert c_dir.direction in ("reduced", "decreased")
    assert c_dir.population == "unspecified"

    c_time = _run_claim_probe(
        "Aurenix improved progression-free survival at 24 months.",
        "Aurenix improved progression-free survival at 24 weeks.",
    )
    assert c_time.support_status == SupportClassification.CONTRADICTED
    assert "24 months" in c_time.numerical_mismatches

    c_role = _run_claim_probe(
        "Aurenix showed response in 20% and toxicity in 60% of patients.",
        "Aurenix showed response in 60% and toxicity in 20% of patients.",
    )
    assert c_role.support_status == SupportClassification.CONTRADICTED
    assert c_role.numerical_mismatches
    # Greedy population regex fix: must extract "patients", not "20% and toxicity in 60% of patients"
    assert c_role.population == "patients"

    # Intra-clause "with" connector role swap
    c_role_with = _run_claim_probe(
        "Aurenix demonstrated a 20% response rate with a 60% toxicity rate.",
        "Aurenix demonstrated a 60% response rate with a 20% toxicity rate.",
    )
    assert c_role_with.support_status == SupportClassification.CONTRADICTED
    assert "20%" in c_role_with.numerical_mismatches
    assert "60%" in c_role_with.numerical_mismatches

    # Standalone subject-position population switch ("Pediatric patients" vs "Adult patients")
    c_pop_subject = _run_claim_probe(
        "Pediatric patients treated with Aurenix had improved overall survival.",
        "Adult patients treated with Aurenix had improved overall survival.",
    )
    assert c_pop_subject.population == "Pediatric patients"
    assert c_pop_subject.support_status == SupportClassification.CONTRADICTED

    # Detrimental verb ("worsened") vs downward risk-reduction verb ("reduced")
    c_dir_down = _run_claim_probe(
        "Aurenix worsened disease progression risk by 38%.",
        "Aurenix reduced disease progression risk by 38%.",
    )
    assert c_dir_down.support_status == SupportClassification.CONTRADICTED

    # Unverified timeframe against timeframe-free source
    c_unverified_tf = _run_claim_probe(
        "Aurenix improved overall survival in adults at 24 months.",
        "Aurenix improved overall survival in adults.",
    )
    assert c_unverified_tf.support_status == SupportClassification.CONTRADICTED
    assert "24 months" in c_unverified_tf.numerical_mismatches

    # 2. Automatic required_qualifiers extraction in extract_pdf_source_pack when passage_rules omitted
    qual_pdf = tmp_path / "auto_qualifiers.pdf"
    create_born_digital_reference_pdf(
        qual_pdf,
        pages_spec=[
            {
                "header": "Clinical Efficacy Summary",
                "left_column": [
                    "In 2L+ EGFR-mutated NSCLC, Aurenix reduced progression risk by 38% (95% CI, 0.49 to 0.78; p=0.0004) in adults."
                ],
                "right_column": [],
            }
        ],
    )
    auto_pack = extract_pdf_source_pack(qual_pdf)
    assert len(auto_pack) >= 1
    assert "2L+ EGFR-mutated NSCLC" in auto_pack[0]["required_qualifiers"]
    assert "95% CI" in auto_pack[0]["required_qualifiers"]

    # 3. Exact stale / post-edit re-verification probe from probe.py + SourcePackList["entries"] compatibility
    j_stale = JobRecord(
        job_id="p_stale",
        tenant_id="t",
        title="t",
        sector=SectorType.MEDTECH,
        document_type="Brief",
        destination_platform="veeva_promomats",
        deadline_iso="2026-12-01T00:00:00Z",
    )
    ingest_working_document(job=j_stale, body_text="Aurenix improved overall survival.", actor="a")
    locate_and_anchor_evidence(
        job=j_stale,
        source_pack=[
            {
                "source_id": "S1",
                "version": "1",
                "verbatim_text": "Aurenix improved overall survival.",
            }
        ],
    )
    assert len(j_stale.source_pack["entries"]) == 1
    assert j_stale.current_version == 1
    assert j_stale.claims[0].bound_version == 1
    assert j_stale.claims[0].stale is False
    assert j_stale.claims[0].support_status == SupportClassification.SUPPORTED

    j_stale.comments.append(
        ReviewComment(
            comment_id="c1",
            reviewer_name="r",
            reviewer_role="Medical",
            timestamp="2026-10-29T00:00:00Z",
            anchored_passage="Aurenix improved overall survival.",
            comment_text="change",
            category=CommentCategory.FACTUAL_CORRECTION,
            human_decision="accepted",
            proposed_spans=[
                WordingSpanEdit(target_text="improved", replacement_text="worsened")
            ],
        )
    )
    apply_accepted_changes(job=j_stale, comments=j_stale.comments, human_actor="a")
    assert j_stale.current_version == 2
    assert [(c.sentence_text, c.support_status.value) for c in j_stale.claims] == [
        ("Aurenix worsened overall survival.", "contradicted")
    ]
    assert j_stale.claims[0].bound_version == 2
    assert j_stale.claims[0].bound_revision_sha256 == j_stale.revisions[-1].sha256_hash
    assert j_stale.claims[0].stale is False

    # Also verify that if source_pack is cleared, a subsequent edit marks existing claims stale=True
    j_stale.source_pack.clear()
    ingest_working_document(
        job=j_stale,
        body_text="Aurenix worsened overall survival significantly.",
        actor="a",
    )
    assert j_stale.current_version == 3
    assert j_stale.claims[0].stale is True

    # 3b. execute-file-roundtrip must preserve existing job.source_pack and NOT fabricate a self-supporting source from post-edit v2 text
    client = TestClient(main_mod.app)
    rt_job_id = "job_rt_preserve_contradicted_v2"
    if rt_job_id in main_mod.JOBS_DB:
        del main_mod.JOBS_DB[rt_job_id]
    headers_rt = {"X-Tenant-Id": "tenant_rt", "X-Actor-Id": "alice@tenant_rt.com"}
    intake_rt = client.post(
        "/jobs/intake",
        headers=headers_rt,
        json={
            "job": JobRecord(
                job_id=rt_job_id,
                tenant_id="tenant_rt",
                title="Roundtrip Source Integrity Check",
                sector=SectorType.MEDTECH,
                document_type="Brief",
                destination_platform="veeva_promomats",
                deadline_iso="2026-12-01T00:00:00Z",
            ).model_dump(),
            "initial_body": "Aurenix improved overall survival in adults.",
            "actor": "alice@tenant_rt.com",
            "source_pack": [
                {
                    "source_id": "S1",
                    "version": "1.0",
                    "verbatim_text": "Aurenix improved overall survival in adults.",
                }
            ],
        },
    )
    assert intake_rt.status_code == 200
    assert intake_rt.json()["job"]["claims"][0]["support_status"] == "supported"

    cmt_rt = client.post(
        f"/jobs/{rt_job_id}/comments",
        headers=headers_rt,
        json=[
            ReviewComment(
                comment_id="c_rt_1",
                reviewer_name="r",
                reviewer_role="Medical",
                timestamp="2026-10-29T00:00:00Z",
                anchored_passage="Aurenix improved overall survival in adults.",
                comment_text="change",
                category=CommentCategory.FACTUAL_CORRECTION,
                human_decision="accepted",
                proposed_spans=[
                    WordingSpanEdit(target_text="improved", replacement_text="worsened")
                ],
            ).model_dump()
        ],
    )
    assert cmt_rt.status_code == 200
    assert cmt_rt.json()["job"]["claims"][0]["support_status"] == "contradicted"

    rt_exec = client.post(f"/jobs/{rt_job_id}/execute-file-roundtrip", headers=headers_rt)
    assert rt_exec.status_code == 200
    rt_job_out = rt_exec.json()["job"]
    assert rt_job_out["current_version"] == 2
    assert rt_job_out["claims"][0]["bound_version"] == 2
    assert rt_job_out["claims"][0]["support_status"] == "contradicted"
    assert rt_exec.json()["resubmission_zip"]["evidence_bound_to_current_version"] is True

    # 4. Exact auth probes from probe.py and extra.py + HMAC-SHA256 session token & actor spoofing checks
    def _make_req(headers):
        scope = {"type": "http", "headers": headers, "query_string": b""}
        return Request(scope)

    for bad_headers in [
        [(b"x-actor-id", b"impersonated-reviewer")],
        [(b"authorization", b"nonsense")],
        [(b"authorization", b"Bearer nonsense")],
        [(b"x-tenant-id", b"victim")],
    ]:
        try:
            main_mod._require_authenticated_tenant(_make_req(bad_headers), "victim")
            raise AssertionError(f"Expected 401 HTTPException for headers {bad_headers}")
        except HTTPException as exc:
            assert exc.status_code == 401

    # Actor spoofing in POST /jobs/intake must be rejected with 403
    spoof_resp = client.post(
        "/jobs/intake",
        headers=headers_rt,
        json={
            "job": JobRecord(
                job_id="job_spoof_actor",
                tenant_id="tenant_rt",
                title="Spoof Actor Attempt",
                sector=SectorType.MEDTECH,
                document_type="Brief",
                destination_platform="veeva_promomats",
                deadline_iso="2026-12-01T00:00:00Z",
            ).model_dump(),
            "initial_body": "Aurenix improved overall survival in adults.",
            "actor": "spoofed_cmo@tenant_rt.com",
            "source_pack": [],
        },
    )
    assert spoof_resp.status_code == 403

    main_mod.JOBS_DB.save_job(
        JobRecord(
            job_id="victim-private",
            tenant_id="tenant_aurenix",
            title="Victim Private Job",
            sector=SectorType.MEDTECH,
            document_type="Brief",
            destination_platform="veeva_promomats",
            deadline_iso="2026-12-01T00:00:00Z",
        )
    )
    main_mod.JOBS_DB.save_job(
        JobRecord(
            job_id="victim-other",
            tenant_id="tenant_other_victim",
            title="Victim Other Tenant Job",
            sector=SectorType.MEDTECH,
            document_type="Brief",
            destination_platform="veeva_promomats",
            deadline_iso="2026-12-01T00:00:00Z",
        )
    )
    r_actor_only = client.get(
        "/jobs/victim-private", headers={"X-Actor-Id": "attacker"}
    )
    assert r_actor_only.status_code == 401
    r_nonsense_auth = client.get(
        "/jobs/victim-private", headers={"Authorization": "nonsense"}
    )
    assert r_nonsense_auth.status_code == 401
    r_tenant_only = client.get(
        "/jobs/victim-private", headers={"X-Tenant-Id": "tenant_aurenix"}
    )
    assert r_tenant_only.status_code == 401

    # Valid HMAC-SHA256 bearer token succeeds for matching tenant and fails with 403 for cross-tenant or conflicting X-Actor-Id
    valid_token = main_mod.issue_signed_session_token(
        tenant_id="tenant_aurenix",
        actor_id="sarah.chen@aurenix.example",
        role="reviewer",
        ttl_seconds=3600,
    )
    r_valid_bearer = client.get(
        "/jobs/victim-private",
        headers={"Authorization": f"Bearer {valid_token}"},
    )
    assert r_valid_bearer.status_code == 200

    r_conflicting_actor_header = client.get(
        "/jobs/victim-private",
        headers={
            "Authorization": f"Bearer {valid_token}",
            "X-Actor-Id": "mallory@aurenix.example",
        },
    )
    assert r_conflicting_actor_header.status_code == 403

    r_cross_tenant_bearer = client.get(
        "/jobs/victim-other",
        headers={"Authorization": f"Bearer {valid_token}"},
    )
    assert r_cross_tenant_bearer.status_code == 403

    # Tampered signature and expired token must fail with 401
    tampered_token = valid_token[:-4] + ("AAAA" if valid_token[-4:] != "AAAA" else "BBBB")
    r_tampered = client.get(
        "/jobs/victim-private",
        headers={"Authorization": f"Bearer {tampered_token}"},
    )
    assert r_tampered.status_code == 401

    from unittest.mock import patch
    with patch.object(main_mod.time, "time", return_value=1_000_000_000):
        expired_token = main_mod.issue_signed_session_token(
            tenant_id="tenant_aurenix",
            actor_id="sarah.chen@aurenix.example",
            ttl_seconds=1,
        )
    r_expired = client.get(
        "/jobs/victim-private",
        headers={"Authorization": f"Bearer {expired_token}"},
    )
    assert r_expired.status_code == 401

    # Runtime secret rotation invalidates tokens signed with old secret
    import os

    prev_secret = os.environ.get("MLR_GENIE_SESSION_SECRET")
    try:
        os.environ["MLR_GENIE_SESSION_SECRET"] = "secret-v1-test"
        tok_v1 = main_mod.issue_signed_session_token("tenant_aurenix", "sarah.chen@aurenix.example")
        os.environ["MLR_GENIE_SESSION_SECRET"] = "secret-v2-rotated"
        r_rotated = client.get(
            "/jobs/victim-private",
            headers={"Authorization": f"Bearer {tok_v1}"},
        )
        assert r_rotated.status_code == 401
    finally:
        if prev_secret is None:
            os.environ.pop("MLR_GENIE_SESSION_SECRET", None)
        else:
            os.environ["MLR_GENIE_SESSION_SECRET"] = prev_secret

    # 5. Telemetry receipt honesty in offline fallback mode
    receipts = c_pop.model_dump()  # verify receipt stored on job
    j_receipt_check = JobRecord(
        job_id="p_receipt",
        tenant_id="t",
        title="t",
        sector=SectorType.MEDTECH,
        document_type="Brief",
        destination_platform="veeva_promomats",
        deadline_iso="2026-12-01T00:00:00Z",
    )
    ingest_working_document(job=j_receipt_check, body_text="Aurenix improved overall survival in adults.", actor="a")
    locate_and_anchor_evidence(
        job=j_receipt_check,
        source_pack=[{"source_id": "S1", "version": "1", "verbatim_text": "Aurenix improved overall survival in adults."}],
    )
    rec = j_receipt_check.metadata["telemetry_receipts"][0]
    assert rec["used_live_sdk"] is False
    assert rec["verification_mode"] == "offline_deterministic_fallback"
    assert rec["zdr_attestation"] == "config_asserted_only_offline"
    assert rec is not receipts

